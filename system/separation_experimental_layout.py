"""V1 measured pagination adapted to bounded fragments, not fake full reports."""

from collections import Counter
from copy import deepcopy
from pathlib import Path
import time
import uuid

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import documents as d
import native_word
import observation_pagination as op
import optimized_pagination as v1
import result_pagination as rp
import separation_layout as base
import separation_experimental_documents as edits
from separation_experimental_word import FragmentNativePass


POLICY = "experimental-fragment-v1-zero-result-paragraph-spacing-20261025"
FOOTER_CLEARANCE_TWIPS = 80


def _fragment_metadata_tables(doc):
    tables = [t for t in d._tables(doc) if rp._metadata_caption(t) == rp.METADATA_CAPTION]
    if not tables:
        raise ValueError("试验结果缺少设备、环境和人员三行页底信息，不能跳过同页核验。")
    for table in tables:
        rows = d._rows(table)
        if (len(rows) != 3 or d._xp(table, ".//w:tbl")
                or table.tblPr.find(qn("w:tblpPr")) is None
                or not all(label in d._text(row) for label, row in zip(
                    ("所用试验设备信息见", "环境温度", "评价人员"), rows))
                or "复核人员" not in d._text(rows[-1])):
            raise ValueError("试验结果页底三行结构不明确，不能安全定位。")
        heights = d._xp(table, "./w:tr/w:trPr/w:trHeight")
        if (len(heights) != 3 or any(
                h.get(qn("w:val")) != str(d.PROJECT_METADATA_ROW_HEIGHT_TWIPS)
                or h.get(qn("w:hRule")) != "exact" for h in heights)):
            raise ValueError("试验结果页底三行必须各为固定0.8cm。")
    return tables


def _fragment_metadata_top(doc, table):
    section = None
    for block in list(doc.element.body)[doc.element.body.index(table):]:
        section = block if block.tag == qn("w:sectPr") else block.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
        if section is not None:
            break
    size = section.find(qn("w:pgSz")) if section is not None else None
    margins = section.find(qn("w:pgMar")) if section is not None else None
    if (size is None or margins is None or not size.get(qn("w:h"))
            or not margins.get(qn("w:top")) or not margins.get(qn("w:bottom"))):
        raise ValueError("试验结果缺少完整页面高度及上下边距，不能定位页底。")
    top = int(margins.get(qn("w:top")))
    printable = int(size.get(qn("w:h"))) - top - int(margins.get(qn("w:bottom")))
    position = printable - 3 * d.PROJECT_METADATA_ROW_HEIGHT_TWIPS - FOOTER_CLEARANCE_TWIPS
    if position < 0:
        raise ValueError("试验结果页面高度不足以容纳完整页底信息组。")
    return top, position


def _prepare_fragment_metadata(path):
    doc = Document(path)
    d._unpack_metadata_groups(doc)
    for table in _fragment_metadata_tables(doc):
        _top, position = _fragment_metadata_top(doc, table)
        floating = table.tblPr.find(qn("w:tblpPr"))
        table.tblPr.remove(floating)
        floating.attrib.pop(qn("w:tblpYSpec"), None)
        floating.set(qn("w:vertAnchor"), "margin")
        # Leave room for Word's mandatory trailing paragraph, not another visible row.
        floating.set(qn("w:tblpY"), str(position))
        table.tblPr.insert(1 if table.tblPr.find(qn("w:tblStyle")) is not None else 0, floating)
        for index, row in enumerate(d._rows(table)):
            for paragraph in d._xp(row, ".//w:p"):
                d._paragraph_setting(paragraph, "keepNext", index < 2)
        previous = table.getprevious()
        while previous is not None and previous.tag == qn("w:p"):
            if not any(label in d._text(previous) for label in (
                    "本试验项目的结论", "试验过程中的异常情况记录")):
                break
            d._paragraph_setting(previous, "keepNext", False)
            previous = previous.getprevious()
    doc.save(path)


def _measure_fragment_metadata(doc):
    return rp._measure_metadata_tables(doc, _fragment_metadata_tables(doc))


def _verify_fragment_metadata_pdf(doc, pdf_path, pages):
    import pymupdf

    tables = _fragment_metadata_tables(doc)
    if len(pages) != len(tables):
        raise ValueError("最终页底信息组页码数量不完整。")
    with pymupdf.open(pdf_path) as pdf:
        if not pages or pages[-1] != len(pdf):
            raise ValueError("最终页底信息未位于项目末页，或存在多余尾页。")
        for table, page_number in zip(tables, pages):
            if not 1 <= page_number <= len(pdf):
                raise ValueError("最终页底信息组页码超出PDF范围。")
            page = pdf[page_number - 1]
            text, boxes = "", []
            for block in page.get_text("rawdict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        for char in span["chars"]:
                            normalized = base.normalized_text(char["c"])
                            text += normalized
                            boxes.extend([char["bbox"]] * len(normalized))
            top, position = _fragment_metadata_top(doc, table)
            first_y = (top + position) / 20
            selected = set()
            for index, row in enumerate(d._rows(table)):
                expected = base.normalized_text(d._text(row))
                y0 = first_y + index * d.PROJECT_METADATA_ROW_HEIGHT_TWIPS / 20
                y1 = y0 + d.PROJECT_METADATA_ROW_HEIGHT_TWIPS / 20
                offset, matches = 0, []
                while expected and (offset := text.find(expected, offset)) >= 0:
                    end = offset + len(expected)
                    if all(y0 - 1 <= box[1] and box[3] <= y1 + 1 for box in boxes[offset:end]):
                        matches.append((offset, end))
                    offset += 1
                if len(matches) != 1:
                    raise ValueError(f"最终PDF页底第{index + 1}行未完整且唯一地位于预定页底，不能交付。")
                selected.update(range(*matches[0]))
            bottom = first_y + 3 * d.PROJECT_METADATA_ROW_HEIGHT_TWIPS / 20
            if any(i not in selected and box[1] < bottom and box[3] > first_y
                   for i, box in enumerate(boxes)):
                raise ValueError("最终PDF正文与页底信息组重叠，不能交付。")
            if any(pymupdf.Rect(image["bbox"]).intersects(
                    pymupdf.Rect(0, first_y, page.rect.width, bottom))
                   for image in page.get_image_info()):
                raise ValueError("最终PDF图片与页底信息组重叠，不能交付。")
    return {"groups": len(tables), "same_page": True, "bottom_position_verified": True,
            "content_complete": True, "no_overlap": True}


def observation_tables(doc, *, final=False, report=False):
    """The caller already extracted the complete observation semantic unit."""
    tables = op.report_tables(doc) if report else [
        t for t in d._tables(doc) if d._rows(t)
        and "型式评价大纲章节号及项目名称" in d._text(d._rows(t)[0])]
    if not tables:
        raise ValueError("独立观察片段缺少观察项目表，不能跳过分页。")
    for table in tables:
        op._layout(table)
        if d._xp(table, ".//w:tbl | .//w:fldChar | .//w:fldSimple | .//w:instrText | "
                       ".//w:sdt | .//w:del | .//w:ins | .//w:altChunk | "
                       ".//*[local-name()='oMath' or local-name()='oMathPara']"):
            raise ValueError("观察表含不能安全映射字符分页的嵌套表、域、修订或公式。")
        if table.tblPr.find(qn("w:tblCaption")) is None:
            caption = OxmlElement("w:tblCaption")
            caption.set(qn("w:val"), op.FLOW_PREFIX + uuid.uuid4().hex[:12])
            table.tblPr.append(caption)
    if not final:
        op._prepare_marks(tables)
    compat = doc.settings.element.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        doc.settings.element.append(compat)
    modes = [node for node in compat.findall(qn("w:compatSetting"))
             if node.get(qn("w:name")) == "compatibilityMode"]
    if len(modes) > 1:
        raise ValueError("Word兼容模式重复，不能安全分页。")
    mode = modes[0] if modes else OxmlElement("w:compatSetting")
    if not modes:
        compat.append(mode)
    for name, value in (("name", "compatibilityMode"),
                        ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")):
        mode.set(qn("w:" + name), value)
    return tables


def _project_groups(tables):
    groups = []
    for table, row, key in op._project_rows(tables):
        identity = key[1][1]
        if groups and groups[-1][0] is table and groups[-1][1] == identity:
            groups[-1][2].append(row)
        else:
            groups.append((table, identity, [row]))
    return groups


def _sample_mark_paragraphs(sample):
    paragraphs = []
    for source in d._xp(sample, "./w:p"):
        paragraph = deepcopy(source)
        lines = op._text(source).split("\v")
        d._put(paragraph, "\n".join("×" if line.strip() else "" for line in lines))
        paragraphs.append(paragraph)
    return paragraphs


def _sync_sample_marks(sample, positive):
    op._set_content(positive, _sample_mark_paragraphs(sample))
    for cell in (sample, positive):
        cell.get_or_add_tcPr().get_or_add_vAlign().set(qn("w:val"), "center")
    for paragraph in d._xp(positive, "./w:p"):
        paragraph.get_or_add_pPr().get_or_add_jc().set(qn("w:val"), "center")


def _repair_sample_boundary_lines(tables):
    identities = {row: key[1][1] for _table, row, key in op._project_rows(tables)}
    repaired = 0
    for previous, following in zip(tables, tables[1:]):
        if not _owned_continuation(previous, following):
            continue
        headers = list(op._columns(d._rows(previous)[0]))
        if len(headers) != 6:
            continue
        left_row, right_row = d._rows(previous)[-1], d._rows(following)[1]
        if identities[left_row] != identities[right_row]:
            continue
        span = headers[2][:2]
        left = next(cell for a, b, cell in op._columns(left_row) if (a, b) == span)
        right = next(cell for a, b, cell in op._columns(right_row) if (a, b) == span)
        paragraphs, suffixes = d._xp(left, "./w:p"), d._xp(right, "./w:p")
        text = op._text(paragraphs[-1])
        suffix = op._text(suffixes[0])
        if text.endswith("\v") and suffix.strip() and not suffix[0].isspace():
            # The separator between two labels must not become an empty last line.
            left.replace(paragraphs[-1], op._slice(paragraphs[-1], 0, op._units(text) - 1, final=True))
            repaired += 1
    return repaired


def _prepare_sample_projects(tables):
    """Synchronize marks without binding a project's requirements to one page."""
    _repair_sample_boundary_lines(tables)
    for table, _identity, rows in _project_groups(tables):
        headers = list(op._columns(d._rows(table)[0]))
        if len(headers) != 6 or d._norm(d._text(headers[2][2])) != "样机编号":
            continue
        sample_span = headers[2][:2]
        positive_span, negative_span = headers[-3][:2], headers[-2][:2]
        for row in rows:
            cells = {(start, stop): cell for start, stop, cell in op._columns(row)}
            sample, positive, negative = (cells[sample_span], cells[positive_span], cells[negative_span])
            if not op._repeat_mark(positive):
                raise ValueError("观察＋栏包含非叉号文字，不能自动改填。")
            if op._is_mark(negative):
                raise ValueError("观察子项目原－栏已有叉号，不能自动补＋栏造成结论冲突，请先核对。")
            _sync_sample_marks(sample, positive)


def _prepare_requirement_flow(tables):
    for table in tables:
        name_end, remarks_start = op._layout(table)
        for row in d._rows(table)[1:]:
            properties = row.get_or_add_trPr()
            for height in d._xp(properties, "./w:trHeight"):
                if height.get(qn("w:hRule")) != "exact":
                    properties.remove(height)
            constraint = properties.find(qn("w:cantSplit"))
            if constraint is None:
                constraint = OxmlElement("w:cantSplit")
                properties.append(constraint)
            constraint.set(qn("w:val"), "false")
            for start, _stop, cell in op._columns(row):
                paragraphs = d._xp(cell, "./w:p")
                if paragraphs:
                    d._paragraph_setting(paragraphs[-1], "keepNext", False)
                if start in (name_end, remarks_start):
                    for paragraph in paragraphs:
                        d._paragraph_setting(paragraph, "keepLines", False)


def _finish_project_layout(tables):
    """Coalesce equal subprojects and marks; callers must remeasure all changes."""
    before = [table.xml for table in tables]
    _repair_sample_boundary_lines(tables)
    for table in tables:
        op._remove_measurements(table)
    op.finish_name_layout(tables)
    for table, _identity, rows in _project_groups(tables):
        if len(rows) == 1:
            continue
        name_end, _ = op._layout(table)
        columns = [list(op._columns(row)) for row in rows]
        spans = [tuple((start, stop) for start, stop, _ in cells) for cells in columns]
        if len(set(spans)) != 1:
            raise ValueError("同页重复观察子项目的网格不一致，不能安全合并。")
        for index, (start, _stop, target) in enumerate(columns[0]):
            if start < name_end:
                continue
            paragraphs = [paragraph for cells in columns
                          for paragraph in d._xp(cells[index][2], "./w:p")]
            op._set_content(target, paragraphs)
            for merge in d._xp(target, "./w:tcPr/w:vMerge"):
                merge.getparent().remove(merge)
        for row in rows[1:]:
            table.remove(row)
    op.finish_name_layout(tables)
    for table in tables:
        headers = list(op._columns(d._rows(table)[0]))
        sample_span = headers[2][:2] if len(headers) == 6 else None
        positive_span, negative_span = headers[-3][:2], headers[-2][:2]
        for row in d._rows(table)[1:]:
            cells = {(start, stop): cell for start, stop, cell in op._columns(row)}
            positive, negative = cells[positive_span], cells[negative_span]
            if not op._repeat_mark(positive):
                raise ValueError("观察＋栏包含非叉号文字，不能自动改填。")
            if op._is_mark(negative):
                raise ValueError("观察子项目原－栏已有叉号，不能自动补＋栏造成结论冲突，请先核对。")
            if sample_span is None:
                d._put(positive, "×")
                op._center_mark(positive)
            else:
                sample = cells[sample_span]
                _sync_sample_marks(sample, positive)
    return any(table.xml != original for table, original in zip(tables, before))


def _project_layout_issue(tables):
    if any(len(rows) > 1 for _table, _identity, rows in _project_groups(tables)):
        return "观察同页重复子项目未完整合并，停止交付。"
    for table in tables:
        headers = list(op._columns(d._rows(table)[0]))
        sample_span = headers[2][:2] if len(headers) == 6 else None
        for row in d._rows(table)[1:]:
            cells = {(start, stop): cell for start, stop, cell in op._columns(row)}
            positive = cells[headers[-3][:2]]
            expected = ([op._text(p) for p in _sample_mark_paragraphs(cells[sample_span])]
                        if sample_span is not None else ["×"])
            if [op._text(p) for p in d._xp(positive, "./w:p")] != expected:
                return "观察＋栏叉号与样机编号不对应，停止交付。"
    return None


def _validate_project_layout(tables):
    issue = _project_layout_issue(tables)
    if issue:
        raise ValueError(issue)


def _source_paragraphs(tables, *, remarks=False):
    sources = {}
    for table, row, key in op._project_rows(tables):
        name_end, remarks_start = op._layout(table)
        column = remarks_start if remarks else name_end
        cell = next(cell for start, _stop, cell in op._columns(row) if start == column)
        texts = [op._text(p) for p in d._xp(cell, "./w:p")]
        # Existing pagination trims leading remark whitespace; provenance uses
        # canonical text only, while the actual XML is moved without trimming.
        if remarks:
            texts = [base.normalized_text(text) for text in texts]
        sources.setdefault(key[1][:2], []).extend(texts)
    return sources


def _source_requirements(tables):
    return _source_paragraphs(tables)


def _column_positions(tables, sources, *, remarks=False):
    contents, positions = {}, {}
    for table, row, key in op._project_rows(tables):
        identity = key[1][:2]
        name_end, remarks_start = op._layout(table)
        column = remarks_start if remarks else name_end
        cell = next(cell for start, _stop, cell in op._columns(row) if start == column)
        texts = [op._text(p) for p in d._xp(cell, "./w:p")]
        text = "".join(base.normalized_text(value) if remarks else value for value in texts)
        preceding = contents.get(identity, "")
        positions[row] = (identity, op._units(preceding), op._units(preceding + text), cell)
        contents[identity] = preceding + text
    if contents != {key: "".join(paragraphs) for key, paragraphs in sources.items()}:
        column_name = "备注" if remarks else "要求"
        raise ValueError(f"观察{column_name}与拆分前的段落来源不一致，不能安全回填尾段。")
    return positions


def _requirement_positions(tables, sources):
    return _column_positions(tables, sources)


def _measured_table_pages(tables, measured, locations):
    if len(tables) != len(measured):
        raise ValueError("尾段回填后的Word表格数量与测量不一致。")
    pages = []
    for table, row_marks in zip(tables, measured):
        rows = d._rows(table)[1:]
        if len(rows) != len(row_marks):
            raise ValueError("尾段回填后的Word行数与测量不一致。")
        table_pages = set()
        for row, cells_marks in zip(rows, row_marks):
            cells = d._cells(row)
            if len(cells) != len(cells_marks):
                raise ValueError("尾段回填后的Word单元格数与测量不一致。")
            for cell, marks in zip(cells, cells_marks):
                paragraphs = d._xp(cell, "./w:p")
                if marks and len(paragraphs) != len(marks):
                    raise ValueError("尾段回填后的Word段落数与测量不一致。")
                for paragraph, (mark_name, text) in zip(paragraphs, marks):
                    if op._text(paragraph) != text:
                        raise ValueError("尾段回填后的Word文字与测量不一致。")
                    table_pages.update(page for page, _offset in op._checked_segments(mark_name, text, locations))
        if not table_pages:
            raise ValueError("尾段回填后的Word表格页码缺失。")
        pages.append(table_pages)
    return pages


def _owned_continuation(previous, following):
    caption = rp._metadata_caption(previous)
    separator = previous.getnext()
    return (caption.startswith(op.FLOW_PREFIX) and rp._metadata_caption(following) == caption
            and separator is not None and separator.tag == qn("w:p")
            and separator.getnext() is following
            and bool(d._xp(separator, ".//w:br[@w:type='page']"))
            and not d._text(separator).strip()
            and not d._xp(separator, ".//w:sectPr | .//w:bookmarkStart | .//w:bookmarkEnd | "
                                     ".//w:drawing | .//w:pict | .//w:object | .//w:fldChar | "
                                     ".//w:fldSimple | .//w:commentRangeStart | ./w:pPr/w:numPr"))


def _review_observation_flow(doc, tables):
    """Recheck capacity after name/project merging, not against stale page breaks."""
    before = deepcopy(doc)
    identities = {row: key[1][1] for _table, row, key in op._project_rows(tables)}
    joined, index = 0, 0
    while index + 1 < len(tables):
        previous, following = tables[index:index + 2]
        if not _owned_continuation(previous, following):
            index += 1
            continue
        if (d._xp(previous, "./w:tblGrid/w:gridCol/@w:w")
                != d._xp(following, "./w:tblGrid/w:gridCol/@w:w")
                or d._text(d._rows(previous)[0]) != d._text(d._rows(following)[0])):
            raise ValueError("观察最终复核发现同源续表表头或列宽不同，不能安全重新接排。")
        following_rows = d._rows(following)[1:]
        if (following_rows
                and identities[d._rows(previous)[-1]] == identities[following_rows[0]]):
            index += 1
            continue
        separator = previous.getnext()
        for row in following_rows:
            previous.append(row)
        following.getparent().remove(following)
        separator.getparent().remove(separator)
        tables.pop(index + 1)
        joined += 1
    if joined:
        _finish_project_layout(tables)
        _prepare_requirement_flow(tables)
        _verify_observation_content(before, doc, finalized_marks=True)
    return joined


def _try_remark_backfill(doc, tables, sources, rejected, pages):
    positions = _column_positions(tables, sources, remarks=True)
    for index in range(len(tables) - 2, -1, -1):
        previous, following = tables[index:index + 2]
        if not _owned_continuation(previous, following):
            continue
        previous_row, next_row = d._rows(previous)[-1], d._rows(following)[1]
        identity, _start, boundary, target = positions[previous_row]
        next_identity, next_start, _next_end, source = positions[next_row]
        target_paragraphs, source_paragraphs = d._xp(target, "./w:p"), d._xp(source, "./w:p")
        if (identity != next_identity or boundary != next_start
                or not op._text(target_paragraphs[-1]).strip()
                or not source_paragraphs or not op._text(source_paragraphs[0]).strip()):
            continue
        suffix = source_paragraphs[0]
        suffix_end = boundary + op._units(base.normalized_text(op._text(suffix)))
        token = ("remark", identity, boundary, suffix_end)
        if token in rejected:
            continue
        offset = 0
        same_paragraph = False
        for text in sources[identity]:
            end = offset + op._units(text)
            if offset < boundary < end and boundary < suffix_end <= end:
                same_paragraph = True
                break
            offset = end
        if not same_paragraph:
            continue
        if [(a, b) for a, b, _ in op._columns(previous_row)] != [
                (a, b) for a, b, _ in op._columns(next_row)]:
            raise ValueError("同源备注尾段的列网格不同，不能安全回填。")
        if d._xp(target, "./w:tcPr/w:vMerge") or d._xp(source, "./w:tcPr/w:vMerge"):
            continue
        baseline = deepcopy(doc)
        for table in tables:
            op._remove_measurements(table)
        for child in list(suffix):
            if child.tag != qn("w:pPr"):
                target_paragraphs[-1].append(child)
        if len(source_paragraphs) > 1:
            source.remove(suffix)
        else:
            op._set_content(source, [deepcopy(suffix)])
        _verify_observation_content(baseline, doc, finalized_marks=True)
        _column_positions(tables, sources, remarks=True)
        return baseline, token, index, next(iter(pages[index]))
    return None


def _consume_requirement_row(doc, tables, index, target):
    previous, following = tables[index:index + 2]
    previous_row, next_row = d._rows(previous)[-1], d._rows(following)[1]
    previous_columns, next_columns = list(op._columns(previous_row)), list(op._columns(next_row))
    name_end, _ = op._layout(following)
    for (start, _stop, target_cell), (_a, _b, source_cell) in zip(previous_columns, next_columns):
        if start >= name_end and target_cell is not target:
            target_cell.extend(d._xp(source_cell, "./w:p"))
    remaining = d._rows(following)[2:]
    if remaining:
        origins = {(a, b): cell for a, b, cell in next_columns if a < name_end}
        for start, stop, cell in op._columns(remaining[0]):
            merge = cell.find("./" + qn("w:tcPr") + "/" + qn("w:vMerge"))
            if start < name_end and merge is not None and merge.get(qn("w:val"), "continue") == "continue":
                if not d._text(cell).strip():
                    origin = origins.get((start, stop))
                    if origin is None or not d._text(origin).strip():
                        raise ValueError("回填后续页的上级名称来源缺失，未猜填名称。")
                    op._copy_name(cell, origin)
                merge.set(qn("w:val"), "restart")
        following.remove(next_row)
    else:
        separator = following.getprevious()
        following.getparent().remove(following)
        separator.getparent().remove(separator)
    updated = [table for table in tables if table.getparent() is not None]
    _finish_project_layout(updated)
    return updated


def _first_requirement_group(paragraphs):
    for index, paragraph in enumerate(paragraphs):
        if op._text(paragraph).strip():
            end = index + 1
            while end < len(paragraphs):
                keep_next = d._xp(paragraphs[end - 1], "./w:pPr/w:keepNext")
                if not keep_next or keep_next[0].get(qn("w:val"), "true") in ("false", "0", "off"):
                    break
                end += 1
            return paragraphs[:end]
    return []


def _try_requirement_backfill(doc, tables, sources, excluded, pages):
    positions = _requirement_positions(tables, sources)
    for index in range(len(tables) - 2, -1, -1):
        previous, following = tables[index:index + 2]
        if not _owned_continuation(previous, following):
            continue
        previous_row, next_row = d._rows(previous)[-1], d._rows(following)[1]
        identity, _start, boundary, target = positions[previous_row]
        next_identity, next_start, next_end, source = positions[next_row]
        token = ("requirement-flow", identity, boundary, next_end)
        group = _first_requirement_group(d._xp(source, "./w:p"))
        if identity != next_identity or boundary != next_start or token in excluded or not group:
            continue
        columns, next_columns = list(op._columns(previous_row)), list(op._columns(next_row))
        if [(a, b) for a, b, _ in columns] != [(a, b) for a, b, _ in next_columns]:
            raise ValueError("同源要求回填的列网格不同，不能安全重排。")
        if any(d._xp(cell, "./w:tcPr/w:vMerge") for a, _b, cell in next_columns
               if a >= op._layout(following)[0]):
            continue
        remaining_rows = d._rows(following)[2:]
        if remaining_rows and any(d._xp(cell, "./w:tcPr/w:vMerge")
                                  for a, _b, cell in op._columns(remaining_rows[0])
                                  if a >= op._layout(following)[0]):
            continue
        offset, inside_paragraph = 0, False
        for text in sources[identity]:
            end = offset + op._units(text)
            if offset < boundary < end:
                inside_paragraph = True
                break
            offset = end
        target_paragraphs = d._xp(target, "./w:p")
        if inside_paragraph and (not op._text(group[0]).strip()
                                 or not op._text(target_paragraphs[-1]).strip()):
            continue
        baseline = deepcopy(doc)
        for table in tables:
            op._remove_measurements(table)
        if inside_paragraph:
            for child in list(group[0]):
                if child.tag != qn("w:pPr"):
                    target_paragraphs[-1].append(child)
            source.remove(group[0])
            target.extend(group[1:])
        else:
            target.extend(group)
        if not d._xp(source, "./w:p"):
            source.append(OxmlElement("w:p"))
        _verify_observation_content(baseline, doc, finalized_marks=True)
        _requirement_positions(tables, sources)
        return baseline, token, index, next(iter(pages[index]))
    return None


def _partition_requirement_backfill(doc, tables, trial, sources, measured, locations):
    baseline, token, index, expected_page = trial
    previous, following = tables[index:index + 2]
    positions = _requirement_positions(tables, sources)
    previous_row, next_row = d._rows(previous)[-1], d._rows(following)[1]
    identity, start, _end, target = positions[previous_row]
    source = positions[next_row][3]
    name_end, _ = op._layout(previous)
    column_index = next(i for i, (a, _b, _cell) in enumerate(op._columns(previous_row)) if a == name_end)
    paragraphs = d._xp(target, "./w:p")
    marks = measured[index][-1][column_index]
    if len(paragraphs) != len(marks):
        raise ValueError("Word改变了要求回填的段落数量，不能安全按页界切回。")
    left, right = [], []
    last_page = expected_page
    for paragraph, (mark_name, text) in zip(paragraphs, marks):
        if op._text(paragraph) != text:
            raise ValueError("Word改变了要求回填文字，未按旧位置切回。")
        segments = op._checked_segments(mark_name, text, locations)
        if segments[0][0] < last_page:
            if segments[0][0] < expected_page:
                return None
            raise ValueError("Word要求回填段落的页序异常。")
        last_page = segments[-1][0]
        if segments[0][0] > expected_page:
            right.append(deepcopy(paragraph))
        elif len(segments) == 1:
            left.append(deepcopy(paragraph))
        else:
            boundary = segments[1][1]
            left.append(op._slice(paragraph, 0, boundary))
            right.append(op._slice(paragraph, boundary, op._units(text), final=True))
    original_tables = [table for table in d._tables(baseline)
                       if d._rows(table) and "型式评价大纲章节号及项目名称" in d._text(d._rows(table)[0])]
    original_target = d._cells(d._rows(original_tables[index])[-1])[column_index]
    old_text = "".join(op._text(p) for p in d._xp(original_target, "./w:p"))
    new_text = "".join(op._text(p) for p in left)
    if len(base.normalized_text(new_text)) <= len(base.normalized_text(old_text)):
        return None
    if not new_text.startswith(old_text):
        raise ValueError("Word要求回填不再是原要求的连续前缀，未猜测重排。")
    original_source = d._cells(d._rows(original_tables[index + 1])[1])[column_index]
    original_paragraphs = d._xp(original_source, "./w:p")
    remaining = (d._xp(source, "./w:p")
                 if len(_first_requirement_group(original_paragraphs)) < len(original_paragraphs) else [])
    for table in tables:
        op._remove_measurements(table)
    op._set_content(target, left)
    op._set_content(source, right + remaining)
    if not right and not remaining:
        tables = _consume_requirement_row(doc, tables, index, target)
    _verify_observation_content(baseline, doc, finalized_marks=True)
    positions = _requirement_positions(tables, sources)
    boundary = positions[previous_row][2]
    saturated = ("requirement-flow", identity, boundary, token[3]) if right else None
    return boundary - token[2], saturated


def _try_tail_backfill(doc, tables, sources, rejected, pages):
    positions = _requirement_positions(tables, sources)
    for index in range(len(tables) - 2, -1, -1):
        previous, following = tables[index:index + 2]
        separator = previous.getnext()
        if not _owned_continuation(previous, following):
            continue
        previous_row, next_row = d._rows(previous)[-1], d._rows(following)[1]
        identity, _start, boundary, target = positions[previous_row]
        next_identity, next_start, next_end, source = positions[next_row]
        token = (identity, boundary, next_end)
        if identity != next_identity or boundary != next_start or token in rejected:
            continue
        offset = 0
        same_paragraph = False
        for text in sources[identity]:
            end = offset + op._units(text)
            if offset < boundary < end and boundary < next_end <= end:
                same_paragraph = True
                break
            offset = end
        target_paragraphs, source_paragraphs = d._xp(target, "./w:p"), d._xp(source, "./w:p")
        if (not same_paragraph or not op._text(target_paragraphs[-1]).strip()
                or not source_paragraphs
                or any(not op._text(p).strip() for p in source_paragraphs)):
            continue
        previous_columns, next_columns = list(op._columns(previous_row)), list(op._columns(next_row))
        if [(a, b) for a, b, _ in previous_columns] != [(a, b) for a, b, _ in next_columns]:
            raise ValueError("同源要求尾段的列网格不同，不能安全回填。")
        remaining = d._rows(following)[2:]
        name_end, _ = op._layout(following)
        if remaining and any(start >= name_end and d._xp(cell, "./w:tcPr/w:vMerge")
                             for start, _stop, cell in op._columns(remaining[0])):
            continue
        baseline = deepcopy(doc)
        for table in tables:
            op._remove_measurements(table)
        paragraph = target_paragraphs[-1]
        for suffix in source_paragraphs:
            for child in list(suffix):
                if child.tag != qn("w:pPr"):
                    paragraph.append(child)
        updated = _consume_requirement_row(doc, tables, index, target)
        _verify_observation_content(baseline, doc, finalized_marks=True)
        _requirement_positions(updated, sources)
        return baseline, token, index, next(iter(pages[index]))
    return None


def paginate(path, kind, progress=None, *, report_observe=False, report_results=False, native_pass_factory=None):
    if report_results and (kind != "result" or report_observe):
        raise ValueError("完整报告结果分页模式只能用于结果表，不能混用观察片段模式。")
    if kind == "items":
        return {"passes": 0}
    def select(doc, final=False):
        if report_observe:
            return observation_tables(doc, final=final, report=True)
        return observation_tables(doc, final=True) if final else observation_tables(doc)

    scope = "完整报告观察页" if report_observe else "V4完整报告结果" if report_results else "独立片段"
    keep_sample_groups = kind != "result" or report_results
    native_pass = native_pass_factory or FragmentNativePass
    doc = Document(path)
    if kind == "result":
        rp.mark_result_tables(doc, keep_sample_groups=keep_sample_groups)
        sources, remark_sources = {}, {}
    else:
        tables = select(doc)
        if not report_observe:
            _prepare_sample_projects(tables)
        original = deepcopy(doc)
        _finish_project_layout(tables)
        _verify_observation_content(original, doc, finalized_marks=True)
        _prepare_requirement_flow(tables)
        sources = _source_requirements(tables)
        remark_sources = _source_paragraphs(tables, remarks=True)
    review_source = deepcopy(doc) if kind == "observe" else None
    doc.save(path)
    if kind == "result":
        if report_results:
            rp.prepare_metadata_groups(path)
        else:
            _prepare_fragment_metadata(path)
    deadline, timings, presented = time.monotonic() + 300, {}, False
    pending, rejected, accepted, remark_accepted = None, set(), 0, 0
    flow_probe, flow_result, flow_accepted, flow_units, saturated = False, None, 0, 0, set()
    flow_reviewed, reviewed_boundaries, review_corrections = False, 0, 0
    iteration, pagination_rounds, finishing_rounds = -1, 0, 0
    while True:
        iteration += 1
        if presented:
            finishing_rounds += 1
            if finishing_rounds > 10:
                raise ValueError(scope + "收尾回填未在10轮Word核验内收敛，停止交付。")
        else:
            pagination_rounds += 1
            if pagination_rounds > 10:
                raise ValueError(scope + "V1分页未在10轮内收敛，停止交付，不跳过Word验证。")
        if progress:
            progress(f"{scope}V1真实分页，第{iteration + 1}轮（分页/收尾各最多10轮，共用5分钟）")
        doc = Document(path)
        tables = select(doc, final=True) if kind == "observe" and presented else (
            select(doc) if kind == "observe" else []
        )
        observations = op.measure(doc, tables)
        if not report_observe:
            rp.repair_continuation_cells(doc, keep_sample_groups=keep_sample_groups)
        measured = [] if report_observe else rp._measure_marks(doc)
        if report_observe:
            metadata = []
        elif kind == "result" and not report_results:
            metadata = _measure_fragment_metadata(doc)
        else:
            metadata = rp._measure_metadata(doc)
        doc.save(path)
        changed = False
        restored = False
        pass_instance = native_pass(path, deadline, timings)
        pass_instance.equalize_result_rows = kind == "result"
        with pass_instance as session:
            doc = Document(session.path)
            if observations:
                tables = select(doc, final=True) if presented else select(doc)
                queries = op.queries(observations)
                if flow_probe:
                    # The temporary row may overflow. Query its requirements only;
                    # the next pass verifies the entire repartitioned document.
                    index = pending[0][2]
                    name_end, _ = op._layout(tables[index])
                    column_index = next(i for i, (a, _b, _cell) in
                                        enumerate(op._columns(d._rows(tables[index])[-1])) if a == name_end)
                    queries = observations[index][-1][column_index]
                locations = session.observation_locations(queries)
                if presented:
                    if not flow_probe:
                        table_pages = _measured_table_pages(tables, observations, locations)
                    if pending is not None:
                        trial, baseline_proof = pending
                        baseline, token, target_index, expected_page = trial
                        if flow_probe:
                            flow_result = _partition_requirement_backfill(
                                doc, tables, trial, sources, observations, locations)
                            flow_probe = False
                            changed = flow_result is not None
                            failed = not changed
                            if changed and progress:
                                progress("已按Word实际页界回填要求，下一轮核验整理后的整张表及页码")
                        else:
                            failed = (any(len(pages) != 1 for pages in table_pages)
                                      or table_pages[target_index] != {expected_page})
                        if failed:
                            doc = baseline
                            tables = select(doc, final=True)
                            observations, locations, measured, metadata, bookmarks, table_pages = baseline_proof
                            rejected.add(token)
                            restored = True
                        elif not changed:
                            if token[0] == "remark":
                                remark_accepted += 1
                            elif token[0] == "requirement-flow":
                                flow_accepted += 1
                                flow_units += flow_result[0]
                                if flow_result[1] is not None:
                                    saturated.add(flow_result[1])
                                    saturated.add(flow_result[1][1:])
                            else:
                                accepted += 1
                        if not changed:
                            pending = None
                    if not changed:
                        if any(len(pages) != 1 for pages in table_pages):
                            changed = op.split(doc, tables, observations, locations, carry_nonfinal_remarks=True)
                            if changed:
                                _prepare_requirement_flow(select(doc, final=True))
                                presented = False
                                review_corrections += 1
                        elif _project_layout_issue(tables):
                            _finish_project_layout(tables)
                            _verify_observation_content(review_source, doc, finalized_marks=True)
                            changed = True
                            review_corrections += 1
                        if not changed:
                            op.validate_name_layout(tables, observations, locations)
                            _validate_project_layout(tables)
                else:
                    changed = op.split(doc, tables, observations, locations, carry_nonfinal_remarks=True)
                    if changed and not report_observe:
                        _prepare_sample_projects(select(doc, final=True))
            if not changed:
                names = v1._names(measured) + [name for group in metadata for name in group]
                if not restored:
                    bookmarks = session.pages(names, "片段结果/签名组完整定位")
                changed = rp.split_measured_tables(doc, measured, bookmarks, keep_sample_groups=keep_sample_groups)
                if not changed:
                    rp._validate_metadata_pages(metadata, bookmarks)
                    if observations and not presented:
                        op.validate_name_layout(tables, observations, locations)
                        changed = _finish_project_layout(tables)
                        reflowed = 0
                        if not flow_reviewed:
                            flow_reviewed = True
                            reflowed = _review_observation_flow(doc, tables)
                            reviewed_boundaries = reflowed
                            changed = changed or bool(reflowed)
                            if reflowed and progress:
                                progress("最终复核：解除本次自动续表的旧页界，按整理后实际容量重新接排要求")
                        presented = not bool(reflowed)
                    elif observations and presented:
                        trial = _try_tail_backfill(doc, tables, sources, rejected | saturated, table_pages)
                        if trial is None:
                            trial = _try_remark_backfill(doc, tables, remark_sources, rejected, table_pages)
                        if trial is None:
                            trial = _try_requirement_backfill(doc, tables, sources, rejected | saturated, table_pages)
                        if trial is not None:
                            # A rejected trial restores this exact, already measured snapshot.
                            pending = (trial, (observations, locations, measured, metadata, bookmarks, table_pages))
                            flow_probe = trial[1][0] == "requirement-flow"
                            changed = True
                            if progress:
                                if flow_probe:
                                    progress("收尾后由Word试排后续要求，按实际页界决定可回填的部分")
                                else:
                                    column = "备注" if trial[1][0] == "remark" else "要求"
                                    progress(f"收尾后尝试同源{column}尾段回填，下一轮Word核验能否留在原页")
                    if not changed and kind == "result":
                        session.export_pdf()
        if not changed:
            if kind == "observe":
                _verify_observation_content(review_source, doc, finalized_marks=True)
                _requirement_positions(tables, sources)
                _column_positions(tables, remark_sources, remarks=True)
            for prefix in (rp.ROW_PREFIX, rp.TABLE_PREFIX, rp.OVERFLOW_PREFIX):
                rp._remove_marks(doc, prefix)
            op.cleanup(doc)
        doc.save(path)
        if not changed:
            return {"passes": iteration + 1, "metadata_groups_checked": len(metadata),
                    **({"metadata_pages": [bookmarks[group[0]]["page"] for group in metadata]}
                       if kind == "result" and not report_results else {}),
                    "stage_seconds": timings,
                    "tail_backfill": {"accepted": accepted,
                                      "rejected": sum(token[0] not in ("remark", "requirement-flow")
                                                      for token in rejected)},
                    "remark_backfill": {"accepted": remark_accepted,
                                        "rejected": sum(token[0] == "remark" for token in rejected)},
                    "requirement_backfill": {"accepted": flow_accepted, "moved_utf16_units": flow_units,
                                            "rejected": sum(token[0] == "requirement-flow" for token in rejected)},
                    **({"final_review": {"content_verified": True, "project_layout_verified": True,
                                         "page_boundaries_verified": True,
                                         "reflowed_boundaries": reviewed_boundaries,
                                         "automatic_corrections": review_corrections + reviewed_boundaries}}
                       if kind == "observe" else {})}


def _characters(doc):
    return Counter(base.normalized_text("".join(d._text(p) for p in d._xp(doc.element.body, ".//w:p"))))


def _observation_content(doc):
    """Compare semantic rows, resolving V1's same-page parent-name merges."""
    outside = _characters(doc)
    projects = {}
    tables = [t for t in d._tables(doc) if d._rows(t)
              and "型式评价大纲章节号及项目名称" in d._text(d._rows(t)[0])]
    for table in tables:
        outside.subtract(_characters_in(table))
        name_end, _ = op._layout(table)
        mark_columns = op._mark_columns(table)
        owners = {}
        for row in d._rows(table)[1:]:
            names, values, marks = [], {}, {}
            next_owners = {}
            for start, stop, cell in op._columns(row):
                span = (start, stop)
                text = base.normalized_text(d._text(cell))
                if start < name_end:
                    merge = cell.find("./" + qn("w:tcPr") + "/" + qn("w:vMerge"))
                    if merge is not None and merge.get(qn("w:val"), "continue") == "continue":
                        if span not in owners or text and text != owners[span]:
                            raise ValueError("观察名称合并缺少对应上级名称或内容冲突，停止交付。")
                        text = owners[span]
                    if merge is not None:
                        next_owners[span] = text
                    names.append((start, stop, text))
                elif span in mark_columns and op._repeat_mark(cell):
                    marks[span] = Counter(text)
                else:
                    values[span] = text
            owners = next_owners
            key = tuple(names)
            if not any(text for _, _, text in names):
                raise ValueError("观察项目缺少可核对的名称，停止交付。")
            entry = projects.setdefault(key, {
                "values": {}, "marks": Counter(),
                "positive_span": list(op._columns(d._rows(table)[0]))[-3][:2],
            })
            for span, text in values.items():
                entry["values"][span] = entry["values"].get(span, "") + text
            for span, characters in marks.items():
                for character, count in characters.items():
                    entry["marks"][(span, character)] += count
    return outside, projects


def _characters_in(element):
    return Counter(base.normalized_text("".join(d._text(p) for p in d._xp(element, ".//w:p"))))


def _verify_observation_content(before, after, *, finalized_marks=False):
    old_outside, old_projects = _observation_content(before)
    new_outside, new_projects = _observation_content(after)
    if old_outside != new_outside or old_projects.keys() != new_projects.keys():
        raise ValueError("观察分页改变了基本信息或项目名称对应关系，停止交付。")
    for key, original in old_projects.items():
        current = new_projects[key]
        removed = original["marks"] - current["marks"]
        if finalized_marks:
            removed = Counter({mark: count for mark, count in removed.items()
                               if mark[0] != original["positive_span"]})
        if original["values"] != current["values"] or removed:
            raise ValueError("观察分页改变了同项目要求、备注或原栏标记，停止交付。")


def _verify_observation_page_content(doc, pdf):
    _outside, projects = _observation_content(doc)
    names = {next(text for _a, _b, text in reversed(key) if text) for key in projects}
    header = base.normalized_text("型式评价大纲章节号及项目名称")
    for index, page in enumerate(pdf):
        text = base.normalized_text(page.get_text())
        if header in text and not any(name in text for name in names):
            raise ValueError(f"观察第{index + 1}页只有续表表头或缺少对应项目内容，停止交付。")


def render_verify(path, kind, structure, progress=None):
    before = Document(path)
    characters = _characters(before)
    pagination = paginate(path, kind, progress)
    after = Document(path)
    if kind == "observe":
        _verify_observation_content(before, after, finalized_marks=True)
        _validate_project_layout(observation_tables(after, final=True))
    elif characters - _characters(after):
        raise ValueError("片段分页丢失原文字，停止交付。")
    if kind == "result":
        if ([base.normalized_text(d._text(p)) for p in d._xp(before.element.body, "./w:p") if d._text(p)]
                != [base.normalized_text(d._text(p)) for p in d._xp(after.element.body, "./w:p") if d._text(p)]):
            raise ValueError("试验结果分页改变了表格外正文，停止交付。")
        blocks = edits.result_blocks(after)
        if ([list(b["key"][1]) for b in blocks] != structure["requirements"]
                or any(b["source_labels"] != structure["selected"] for b in blocks)
                or any(d._text(r["_value_cell"]) for r in d._result_schema(after))):
            raise ValueError("分页后技术条件/参试编号覆盖改变或实测结果非空，停止交付。")
    signature = base.content_signature(after)
    native_word.render(path, Path(path).with_suffix(".pdf"), timeout_seconds=180, progress=progress,
                       verify_editing_pages=kind == "result")
    checks = base.verify_rendered(path, Path(path).with_suffix(".pdf"), signature, kind, 1, 1)
    if kind == "result":
        checks["metadata_layout"] = _verify_fragment_metadata_pdf(
            Document(path), Path(path).with_suffix(".pdf"), pagination["metadata_pages"])
    from separation_experimental_format import (
        verify_report_headers, verify_result_closing, verify_result_intro, verify_result_titles,
    )

    checks["header_format"] = verify_report_headers(Document(path))
    if "header_format" in structure and checks["header_format"] != structure["header_format"]:
        raise ValueError("试验拆分输出页眉数量或格式与已确认计划不一致，停止交付。")
    if kind != "observe":
        import pymupdf
        from separation_experimental_format import verify_header_lines

        with pymupdf.open(Path(path).with_suffix(".pdf")) as pdf:
            verify_header_lines(Document(path), pdf)
    if kind == "result":
        checks["title_format"] = verify_result_titles(Document(path))
        if "title_format" in structure and checks["title_format"] != structure["title_format"]:
            raise ValueError("试验拆分输出结果标题数量与已确认计划不一致，停止交付。")
        checks["intro_format"] = verify_result_intro(Document(path))
        if "intro_format" in structure and checks["intro_format"] != structure["intro_format"]:
            raise ValueError("试验拆分输出结果时间及数据记录段数量与已确认计划不一致，停止交付。")
        checks["closing_format"] = verify_result_closing(Document(path))
        if "closing_format" in structure and checks["closing_format"] != structure["closing_format"]:
            raise ValueError("试验拆分输出结果合格判定、结论及异常记录段数量与已确认计划不一致，停止交付。")
    if kind == "observe":
        import pymupdf
        from gas_observation_flow import describe as describe_gas_flow
        from separation_experimental_format import verify_observe_lines

        if "gas_flow_ranges" in structure:
            checks["gas_flow_ranges"] = describe_gas_flow(Document(path))
            if checks["gas_flow_ranges"] != structure["gas_flow_ranges"]:
                raise ValueError("燃气表流量范围的型号分段或参数在分页后改变，停止交付。")
        with pymupdf.open(Path(path).with_suffix(".pdf")) as pdf:
            verify_observe_lines(Document(path), pdf)
            _verify_observation_page_content(Document(path), pdf)
        rendered = Document(path)
        _verify_observation_content(before, rendered, finalized_marks=True)
        _validate_project_layout(observation_tables(rendered, final=True))
        checks["final_review"] = {**pagination["final_review"], "pdf_verified": True,
                                 "no_header_only_pages": True}
    checks.update(layout_policy=POLICY, pagination=pagination,
                  review_note=("观察内容完整性、续页、编号/叉号及Word/PDF最终复核通过；"
                               "可安全修正的版式已重新排版核验，原技术要求不改写，仍请人工审版。"
                               if kind == "observe" else
                               "独立片段V1及Word检查通过；实测格已清空，其他原文保留，请人工审版；报告端兼容尚待适配。"))
    return checks
