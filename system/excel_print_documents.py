"""Convert the linked workbook's print sheet, not its calculation worksheets."""

from decimal import Decimal, ROUND_HALF_UP
from copy import deepcopy
from math import ceil
from pathlib import Path
import re
import uuid

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import get_column_letter, range_boundaries
from word_fonts import text_width


PAGE_LABEL = re.compile(r"第\s*页\s*[/／]\s*共\s*页")
PRINT_BODY_CAPTION = "excel-print-body-"
PRINT_SIGNATURE_CAPTION = "excel-print-signature-"
BORDER_STYLES = {
    "thin": ("single", 4), "medium": ("single", 8), "thick": ("single", 12),
    "hair": ("single", 2), "double": ("double", 4), "dotted": ("dotted", 4),
    "dashed": ("dashed", 4), "dashDot": ("dotDash", 4), "dashDotDot": ("dotDotDash", 4),
    "mediumDashed": ("dashed", 8), "mediumDashDot": ("dotDash", 8),
    "mediumDashDotDot": ("dotDotDash", 8), "slantDashDot": ("dotDash", 8),
}


def display_value(value, number_format):
    if value is None:
        return ""
    if isinstance(value, (str, CellRichText)):
        return str(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if not isinstance(value, (int, float)):
        raise ValueError("尚不支持打印页的数据类型：" + type(value).__name__)
    if number_format in ("General", "@"):
        return str(value) if isinstance(value, int) else format(value, ".15g")
    match = re.fullmatch(r"0(?:\.(0+))?(%)?(_ )?", number_format)
    if not match:
        raise ValueError("尚不支持打印页数字格式：" + number_format)
    places = len(match[1] or "")
    number = Decimal(str(value)) * (100 if match[2] else 1)
    rounded = number.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    return format(rounded, f".{places}f") + ("%" if match[2] else "") + (" " if match[3] else "")


def _color(color):
    if color is None or color.type == "auto" or (color.type == "indexed" and color.indexed == 64):
        return "000000"
    if color.type == "rgb":
        return color.rgb[-6:]
    if color.type == "indexed" and 0 <= color.indexed < len(COLOR_INDEX):
        return COLOR_INDEX[color.indexed][-6:]
    if color.type == "theme" and color.theme in (0, 1) and not color.tint:
        return "FFFFFF" if color.theme == 0 else "000000"
    raise ValueError("打印页包含尚不支持的主题颜色，不能静默改变原格式。")


def _font(run, base, inline, scale):
    def value(name):
        selected = getattr(inline, name, None) if inline is not None else None
        return selected if selected is not None else getattr(base, name)

    run.font.name = "Arial"
    run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt((value("sz") or 11) * scale)
    run.bold, run.italic = bool(value("b")), bool(value("i"))
    run.font.strike = bool(value("strike"))
    underline = value("u")
    if underline not in (None, "single", "double"):
        raise ValueError("打印页包含尚不支持的下划线：" + underline)
    run.underline = WD_UNDERLINE.DOUBLE if underline == "double" else bool(underline)
    run.font.subscript = value("vertAlign") == "subscript"
    run.font.superscript = value("vertAlign") == "superscript"
    run.font.color.rgb = RGBColor.from_string(_color(value("color")))


def _cell_borders(properties, edges, scale):
    for old in list(properties.findall(qn("w:tcBorders"))):
        properties.remove(old)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        original = edges[side]
        edge = OxmlElement("w:" + side)
        if original is not None and original.style:
            style, size = BORDER_STYLES[original.style]
            edge.set(qn("w:val"), style)
            edge.set(qn("w:sz"), str(max(2, round(size * scale))))
            edge.set(qn("w:color"), _color(original.color))
        else:
            edge.set(qn("w:val"), "nil")
        borders.append(edge)
    properties.append(borders)


def _border_edges(sheet, row, left, right, merged, repairs):
    area = merged.get((row, left))
    top, bottom = (area.min_row, area.max_row) if area else (row, row)
    first, last = (area.min_col, area.max_col) if area else (left, right)
    anchor = sheet.cell(top, first)
    edges = {
        "top": anchor.border.top, "left": anchor.border.left,
        "bottom": sheet.cell(bottom, first).border.bottom,
        "right": sheet.cell(top, last).border.right,
    }
    present = [side for side, edge in edges.items() if edge is not None and edge.style]
    if len(present) == 3:
        missing = next(side for side in edges if side not in present)
        opposite = {"top": "bottom", "bottom": "top", "left": "right", "right": "left"}
        edges[missing] = edges[opposite[missing]]
        repairs.add(f"{area if area else anchor.coordinate}:{missing}")
    if row != top:
        edges["top"] = None
    if row != bottom:
        edges["bottom"] = None
    return edges


def _layout_text(value):
    if not isinstance(value, (str, CellRichText)):
        return value
    text = str(value)
    compact = re.sub(r"\s+", "", text)
    approved = {
        "检测流量点", "重复性误差", "流体扰动：参比条件", "流体扰动：90°单弯头",
        "机械环境：振动跌落前", "机械环境：振动跌落后", "降流量", "升流量",
    }
    gaps = list(re.finditer(r"(?<=\S)[ \t\u3000]+(?=\S)", text))
    if compact not in approved or not gaps:
        return value
    edits = {}
    for index, gap in enumerate(gaps):
        for position in range(gap.start(), gap.end()):
            edits[position] = "\n" if index == 0 and position == gap.start() else ""
    if isinstance(value, str):
        return "".join(edits.get(index, character) for index, character in enumerate(text))
    output, offset = CellRichText(), 0
    for part in value:
        original = part.text if isinstance(part, TextBlock) else part
        changed = "".join(edits.get(offset + index, character) for index, character in enumerate(original))
        offset += len(original)
        if changed:
            output.append(TextBlock(part.font, changed) if isinstance(part, TextBlock) else changed)
    return output


def _cell_text_width(cell):
    lines = [0.0]
    for run in cell.iter(qn("w:r")):
        size = run.find("./" + qn("w:rPr") + "/" + qn("w:sz"))
        if size is None:
            raise ValueError("打印页单元格缺少明确字号，不能可靠调整标签宽度。")
        for node in run:
            if node.tag == qn("w:t"):
                lines[-1] += text_width(node.text or "", int(size.get(qn("w:val"))) / 2)
            elif node.tag == qn("w:br"):
                lines.append(0.0)
    return ceil(max(lines) * 20) + 44


def _widen_labels(table):
    grid = table._tbl.tblGrid
    widths = [int(column.get(qn("w:w"))) for column in grid]
    positions = [0]
    for width in widths:
        positions.append(positions[-1] + width)
    rows, boundaries, adjusted = [], set(positions), []
    for row in table._tbl.tr_lst:
        cells, edges, column = list(row.tc_lst), [0], 0
        for cell in cells:
            column += cell.grid_span
            edges.append(positions[column])
        for index, cell in enumerate(cells[:-1]):
            text = "".join(node.text or "" for node in cell.iter(qn("w:t")))
            if not re.fullmatch(r"[\u4e00-\u9fff]{2,8}[：:]", text):
                continue
            following = cells[index + 1]
            if cell.tcPr.vMerge is not None or following.tcPr.vMerge is not None:
                continue
            required = _cell_text_width(cell)
            extra = max(0, required - (edges[index + 1] - edges[index]))
            if extra:
                remaining = extra
                for donor_index in range(index + 1, len(cells)):
                    donor = cells[donor_index]
                    if donor.tcPr.vMerge is not None:
                        break
                    spare = max(0, edges[donor_index + 1] - edges[donor_index] - _cell_text_width(donor))
                    taken = min(remaining, spare)
                    for boundary in range(index + 1, donor_index + 1):
                        edges[boundary] += taken
                    remaining -= taken
                    if not remaining:
                        break
                if remaining:
                    raise ValueError(f"打印页“{text}”及右侧内容无法同时单行显示，未缩小字号或截断内容。")
                adjusted.append(text)
            no_wrap = OxmlElement("w:noWrap")
            no_wrap.set(qn("w:val"), "true")
            cell.tcPr.append(no_wrap)
        boundaries.update(edges)
        rows.append((cells, edges))
    if adjusted:
        ordered = sorted(boundaries)
        for column in list(grid):
            grid.remove(column)
        for left, right in zip(ordered, ordered[1:]):
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(right - left))
            grid.append(column)
        indices = {position: index for index, position in enumerate(ordered)}
        for cells, edges in rows:
            for cell, left, right in zip(cells, edges, edges[1:]):
                cell.get_or_add_tcPr().get_or_add_tcW().set(qn("w:w"), str(right - left))
                cell.get_or_add_tcPr().get_or_add_gridSpan().set(qn("w:val"), str(indices[right] - indices[left]))
    return list(dict.fromkeys(adjusted))


def _minimal_paragraph(keep_next=False):
    paragraph = OxmlElement("w:p")
    properties = paragraph.get_or_add_pPr()
    spacing = properties.get_or_add_spacing()
    for name, value in (("before", "0"), ("after", "0"), ("line", "1"), ("lineRule", "exact")):
        spacing.set(qn("w:" + name), value)
    for name, enabled in (("keepNext", keep_next), ("pageBreakBefore", False), ("snapToGrid", False)):
        setting = OxmlElement("w:" + name)
        setting.set(qn("w:val"), "true" if enabled else "false")
        properties.append(setting)
    return paragraph


def _anchor_signature(table, required):
    rows = list(table._tbl.tr_lst)
    labels = (("试验员：", "核验员：", "测量日期："),
              ("检测人员：", "核验人员：", "检测日期："))
    signatures = [index for index, row in enumerate(rows)
                  if any(first in re.sub(r"\s+", "", "".join(node.text or "" for node in row.iter(qn("w:t"))))
                         for first, _second, _date in labels)]
    if not signatures and not required:
        return False
    if signatures != [len(rows) - 2]:
        raise ValueError("打印页末尾未找到唯一的试验员行及横线，不能猜测移动内容。")
    signature, underline = rows[-2:]
    text = "".join(node.text or "" for node in signature.iter(qn("w:t")))
    if (not any(all(label in re.sub(r"\s+", "", text) for label in group) for group in labels)
            or any((node.text or "").strip() for node in underline.iter(qn("w:t")))
            or any(cell.tcPr.vMerge is not None for row in (signature, underline) for cell in row.tc_lst)):
        raise ValueError("打印页签名/横线结构不支持整体定位，未删除或重写内容。")
    heights = [row.get_or_add_trPr().find(qn("w:trHeight")) for row in (signature, underline)]
    if any(height is None for height in heights):
        raise ValueError("打印页签名/横线缺少行高，无法预留页底空间。")
    signature_height, underline_height = [int(height.get(qn("w:val"))) for height in heights]
    edges, column = [], 0
    for cell in underline.tc_lst:
        edge = cell.tcPr.find("./" + qn("w:tcBorders") + "/" + qn("w:bottom"))
        if edge is None:
            raise ValueError("打印页下方横线的边框定义不完整。")
        edges.append((column, column + cell.grid_span, edge))
        column += cell.grid_span
    if not any(edge.get(qn("w:val")) not in ("nil", "none") for _, _, edge in edges):
        raise ValueError("打印页末行没有可定位的横线。")
    column = 0
    for cell in signature.tc_lst:
        covered = [edge for start, stop, edge in edges if start < column + cell.grid_span and stop > column]
        if not covered or len({tuple(sorted(edge.attrib.items())) for edge in covered}) != 1:
            raise ValueError("打印页横线在签名格内不连续，不能改变原横线范围。")
        properties = cell.get_or_add_tcPr()
        borders = properties.find(qn("w:tcBorders"))
        old = borders.find(qn("w:bottom"))
        borders.replace(old, deepcopy(covered[0]))
        margins = properties.find(qn("w:tcMar"))
        margins.find(qn("w:bottom")).set(qn("w:w"), str(underline_height))
        column += cell.grid_span
    floating = deepcopy(table._tbl)
    for row in list(floating.tr_lst):
        floating.remove(row)
    table._tbl.remove(signature)
    table._tbl.remove(underline)
    floating.append(signature)
    height = signature.get_or_add_trPr().get_or_add_trHeight()
    height.set(qn("w:val"), str(signature_height + underline_height))
    height.set(qn("w:hRule"), "atLeast")
    position = OxmlElement("w:tblpPr")
    for name, value in (("vertAnchor", "margin"), ("horzAnchor", "margin"), ("tblpYSpec", "bottom"),
                        ("tblpX", "0"), ("topFromText", "120"), ("bottomFromText", "0"),
                        ("leftFromText", "0"), ("rightFromText", "0")):
        position.set(qn("w:" + name), value)
    properties = floating.tblPr
    properties.insert(1, position)
    overlap = OxmlElement("w:tblOverlap")
    overlap.set(qn("w:val"), "never")
    properties.insert(2, overlap)
    width = properties.find(qn("w:tblW"))
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(int(node.get(qn("w:w"))) for node in floating.tblGrid)))
    identity = uuid.uuid4().hex
    for element, prefix in ((table._tbl, PRINT_BODY_CAPTION), (floating, PRINT_SIGNATURE_CAPTION)):
        caption = OxmlElement("w:tblCaption")
        caption.set(qn("w:val"), prefix + identity)
        element.tblPr.append(caption)
    anchor = _minimal_paragraph(keep_next=True)
    table._tbl.addnext(anchor)
    anchor.addnext(floating)
    floating.addnext(_minimal_paragraph())
    return True


def _header_text(value):
    if not isinstance(value, CellRichText):
        return value.lstrip()
    parts, leading = CellRichText(), True
    for part in value:
        text = part.text if isinstance(part, TextBlock) else part
        if leading:
            text = text.lstrip()
        if text:
            leading = False
            parts.append(TextBlock(part.font, text) if isinstance(part, TextBlock) else text)
    return parts


def _fill_cell(target, source, value, scale, line_height, print_header=False):
    value = _layout_text(value)
    properties = target._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement("w:" + side)
        margin.set(qn("w:w"), "0" if side in ("top", "bottom") else "12")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)
    if source.fill.patternType not in (None, "solid"):
        raise ValueError("打印页包含尚不支持的图案填充：" + str(source.fill.patternType))
    if source.fill.patternType == "solid":
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), _color(source.fill.fgColor))
        properties.append(shading)
    if source.alignment.textRotation:
        raise ValueError("打印页包含旋转单元格，须先确认转换版式。")
    target.vertical_alignment = {
        "top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
        "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM, None: WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
    }[source.alignment.vertical]
    paragraph = target.paragraphs[0]
    alignment = source.alignment.horizontal
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
        "centerContinuous": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "distributed": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
        None: WD_ALIGN_PARAGRAPH.RIGHT if isinstance(value, (int, float)) and not isinstance(value, bool)
        else WD_ALIGN_PARAGRAPH.LEFT,
    }[alignment]
    if print_header:
        value = _header_text(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    sizes = [source.font.sz or 11]
    if isinstance(value, CellRichText):
        sizes.extend(part.font.sz for part in value if isinstance(part, TextBlock) and part.font.sz is not None)
    fmt.line_spacing = Pt(max(line_height, max(sizes) * scale + 0.5))
    fmt.keep_together = True
    grid = OxmlElement("w:snapToGrid")
    grid.set(qn("w:val"), "0")
    paragraph._p.get_or_add_pPr().append(grid)
    if isinstance(value, CellRichText):
        for part in value:
            run = paragraph.add_run(part.text if isinstance(part, TextBlock) else part)
            _font(run, source.font, part.font if isinstance(part, TextBlock) else None, scale)
    else:
        run = paragraph.add_run(display_value(value, source.number_format))
        _font(run, source.font, None, scale)


def _has_border(cell):
    return any(edge is not None and edge.style for edge in (
        cell.border.top, cell.border.bottom, cell.border.left, cell.border.right))


def _page_ranges(sheet, bounds):
    left, top, right, bottom = bounds
    footers = set()
    for row in range(top, bottom + 1):
        values = [str(sheet.cell(row, col).value).strip() for col in range(left, right + 1)
                  if sheet.cell(row, col).value is not None and str(sheet.cell(row, col).value).strip()]
        if len(values) == 1 and PAGE_LABEL.fullmatch(values[0]):
            footers.add(row)
    boundaries = {item.id + 1 for item in sheet.row_breaks.brk if top <= item.id < bottom}
    boundaries.update(row + 1 for row in footers if row < bottom)
    pages = []
    for start, stop in zip([top, *sorted(boundaries)], [*sorted(boundaries), bottom + 1]):
        rows = [row for row in range(start, stop) if row not in footers and not sheet.row_dimensions[row].hidden]
        rows = [row for row in rows if any(
            (sheet.cell(row, col).value is not None and str(sheet.cell(row, col).value).strip())
            or _has_border(sheet.cell(row, col))
            for col in range(left, right + 1))]
        if any(sheet.cell(row, col).value is not None and str(sheet.cell(row, col).value).strip()
               for row in rows for col in range(left, right + 1)):
            pages.append(rows)
    return pages, footers


def _select_report_pages(sheet, pages, projects, report_type):
    from gas_records import PRINT_PROJECTS, compact, project_kind

    if report_type == "流量计（气体）":
        from gas_flowmeter_excel import PRINT_PROJECTS, project_kind

    candidates = []
    for project_index, project in enumerate(projects):
        samples = list(dict.fromkeys(row.get("sample_id", "").strip() for row in project["rows"]
                                     if row.get("sample_id", "").strip()))
        eligible = {
            row["sample_id"].strip() for row in project["rows"]
            if row.get("sample_id", "").strip() and row.get("value") is not None
            and str(row.get("value", "")).strip()
            and (not row.get("matched_sample_id") or row["matched_sample_id"] == row["sample_id"])
        }
        candidates.append((project_index, project, project_kind(project), samples, eligible))
    included, excluded = [], []
    for page_index, rows in enumerate(pages):
        header_label = "编号" if report_type == "流量计（气体）" else "出厂编号："
        headers = [row for row in rows if compact(sheet.cell(row, 2).value) == header_label]
        if len(headers) != 1 or headers[0] not in PRINT_PROJECTS:
            raise ValueError(f"打印页第{page_index + 1}组未定位支持的项目及出厂编号，不能猜测排序或删除。")
        header = headers[0]
        sample = display_value(sheet.cell(header, 3).value, sheet.cell(header, 3).number_format).strip()
        if not sample:
            raise ValueError(f"打印页第{page_index + 1}组缺少出厂编号，不能猜测与报告的对应关系。")
        kinds = PRINT_PROJECTS[header]
        info = {"page_index": page_index, "source_rows": [rows[0], rows[-1]], "sample_id": sample,
                "project_kinds": list(kinds)}
        matches = [(index, project, samples) for index, project, kind, samples, eligible in candidates
                   if kind in kinds and sample in eligible]
        if matches:
            index, project, samples = matches[0]
            info.update(project_index=index, project_id=project["id"], project_name=project["name"],
                        sample_index=samples.index(sample))
            included.append(info)
        else:
            info["reason"] = "报告未选择对应项目，或该项目下此编号没有本次实测结果。"
            excluded.append(info)
    return included, excluded


def print_sheet_document(path, *, projects=None, report_type="燃气表"):
    path = Path(path)
    formulas = load_workbook(path, data_only=False, keep_links=False, rich_text=True)
    cached = load_workbook(path, data_only=True, keep_links=False, rich_text=True)
    try:
        if "打印页" not in formulas.sheetnames:
            raise ValueError(path.name + "：未找到“打印页”工作表，未改用计算表代替。")
        sheet, values = formulas["打印页"], cached["打印页"]
        if sheet._images or sheet._charts:
            raise ValueError(path.name + "：打印页含图片或图表，当前不能保证完整可编辑转换。")
        if sheet.col_breaks.brk or sheet.print_title_cols or sheet.print_title_rows:
            raise ValueError(path.name + "：打印页含分栏分页或重复打印标题，须先确认转换范围。")
        populated = [cell for row in sheet for cell in row
                     if cell.value is not None and str(cell.value).strip()]
        if not populated:
            raise ValueError(path.name + "：打印页为空。")
        if sheet.print_area:
            area = str(sheet.print_area)
            if "," in area:
                raise ValueError(path.name + "：打印页有多个打印区域，须先确认转换顺序。")
            bounds = range_boundaries(area.rsplit("!", 1)[-1])
        else:
            printed = [*populated, *(cell for row in sheet for cell in row if _has_border(cell))]
            bounds = (1, 1, max(cell.column for cell in printed), max(cell.row for cell in printed))
        pages, footers = _page_ranges(sheet, bounds)
        if not pages:
            raise ValueError(path.name + "：打印区域没有可输出的记录。")
        source_page_count = len(pages)
        if projects is None:
            included = [{"page_index": index, "source_rows": [rows[0], rows[-1]]}
                        for index, rows in enumerate(pages)]
            excluded = []
        else:
            included, excluded = _select_report_pages(values, pages, projects, report_type)
        pages = [pages[info["page_index"]] for info in included]
        selected_rows = {row for rows in pages for row in rows}
        for cell in populated:
            dimension = sheet.column_dimensions.get(get_column_letter(cell.column))
            if (cell.data_type == "f" and values[cell.coordinate].value is None and cell.row in selected_rows
                    and bounds[0] <= cell.column <= bounds[2] and not (dimension is not None and dimension.hidden)):
                raise ValueError(f"{path.name}：打印页{cell.coordinate}的公式没有缓存结果，请在Excel计算并保存后重新上传。")
        scale = (sheet.page_setup.scale or 100) / 100
        if not 0.1 <= scale <= 4:
            raise ValueError(path.name + "：打印缩放比例无效。")
        if sheet.page_setup.paperSize is not None and str(sheet.page_setup.paperSize) != "9":
            raise ValueError(path.name + "：当前打印页转换仅支持A4纸张。")
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt((formulas._fonts[0].sz or 11) * scale)
        normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
        line_height = sheet.sheet_format.defaultRowHeight * scale - 0.5
        normal.paragraph_format.line_spacing = Pt(line_height)
        section = doc.sections[0]
        width, height = (297, 210) if sheet.page_setup.orientation == "landscape" else (210, 297)
        section.page_width, section.page_height = Mm(width), Mm(height)
        for side in ("left", "right", "top", "bottom"):
            setattr(section, side + "_margin", Inches(getattr(sheet.page_margins, side)))
        section.footer_distance = min(Mm(5), section.bottom_margin // 2)
        section.header_distance = min(Mm(5), section.top_margin // 2)
        for header in (section.header, section.first_page_header, section.even_page_header):
            paragraph = header.paragraphs[0]
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.add_run().font.size = Pt(1)
        available = section.page_width - section.left_margin - section.right_margin
        columns = [col for col in range(bounds[0], bounds[2] + 1)
                   if not (sheet.column_dimensions.get(get_column_letter(col)) is not None
                           and sheet.column_dimensions[get_column_letter(col)].hidden)]
        if available <= 0 or not columns:
            raise ValueError(path.name + "：打印页宽度或可见列无效。")
        widths = [(sheet.column_dimensions[get_column_letter(col)].width
                   if get_column_letter(col) in sheet.column_dimensions else
                   sheet.sheet_format.defaultColWidth or 8.43) * 7 + 5 for col in columns]
        widths = [round(available * width / sum(widths)) for width in widths]
        column_map = {col: index for index, col in enumerate(columns)}
        warnings, repairs, widened = [], set(), []
        merge_map = {(row, col): area for area in sheet.merged_cells.ranges
                     for row in range(area.min_row, area.max_row + 1)
                     for col in range(area.min_col, area.max_col + 1)}
        for index, rows in enumerate(pages):
            if index:
                separator = doc.add_paragraph()
                separator.paragraph_format.page_break_before = True
                separator.paragraph_format.space_before = separator.paragraph_format.space_after = Pt(0)
                separator.paragraph_format.line_spacing = Pt(1)
                separator.paragraph_format.keep_with_next = True
            included[index]["body_start"] = len(doc.element.body) - 1
            table = doc.add_table(rows=len(rows), cols=len(columns))
            table.autofit = False
            table._tbl.tblPr.get_or_add_tblStyle().val = doc.styles["Normal Table"].style_id
            for column, width in zip(table.columns, widths):
                column.width = width
            row_map = {row: index for index, row in enumerate(rows)}
            for merged in sheet.merged_cells.ranges:
                if merged.min_row in footers:
                    continue
                covered_rows = [row for row in rows if merged.min_row <= row <= merged.max_row]
                covered_cols = [col for col in columns if merged.min_col <= col <= merged.max_col]
                if not covered_rows or not covered_cols:
                    continue
                if (merged.min_row not in row_map or merged.min_col not in column_map
                        or merged.max_row > rows[-1] or merged.max_col > bounds[2]):
                    raise ValueError(f"{path.name}：合并区域{merged}被打印边界截断，未丢弃内容。")
                table.cell(row_map[covered_rows[0]], column_map[covered_cols[0]]).merge(
                    table.cell(row_map[covered_rows[-1]], column_map[covered_cols[-1]]))
            occupied = {(row, col) for area in sheet.merged_cells.ranges
                        for row in rows if area.min_row <= row <= area.max_row
                        for col in columns if area.min_col <= col <= area.max_col}
            overflow = set()
            for row in rows:
                for col in columns:
                    source = sheet.cell(row, col)
                    if ((row, col) in occupied or (row, col) in overflow or source.alignment.wrapText
                            or not isinstance(values.cell(row, col).value, (str, CellRichText))
                            or not str(values.cell(row, col).value).strip()):
                        continue
                    end = column_map[col]
                    while end + 1 < len(columns):
                        following = columns[end + 1]
                        other = sheet.cell(row, following)
                        if ((row, following) in occupied or other.value is not None
                                or any(edge is not None and edge.style for edge in
                                       (source.border.right, other.border.left, other.border.right))
                                or source.border.top != other.border.top or source.border.bottom != other.border.bottom):
                            break
                        end += 1
                    if end > column_map[col]:
                        table.cell(row_map[row], column_map[col]).merge(table.cell(row_map[row], end))
                        overflow.update((row, following) for following in columns[column_map[col] + 1:end + 1])
            targets = table._cells
            for row in rows:
                target_row = table.rows[row_map[row]]
                target_row.height = Pt((sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight) * scale)
                empty_rule = all(not isinstance(sheet.cell(row, col), MergedCell)
                                 and (values.cell(row, col).value is None
                                      or not str(values.cell(row, col).value).strip()) for col in columns)
                target_row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if empty_rule else WD_ROW_HEIGHT_RULE.AT_LEAST
                target_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for col in columns:
                    source = sheet.cell(row, col)
                    if isinstance(source, MergedCell) or (row, col) in overflow:
                        continue
                    value = values.cell(row, col).value
                    if values.cell(row, col).data_type == "e":
                        warnings.append(f"{path.name}：打印页{source.coordinate}保留原Excel错误值{value}，请核对。")
                    try:
                        target = targets[row_map[row] * len(columns) + column_map[col]]
                        target.width = sum(widths[column_map[col]:column_map[col] + target._tc.grid_span])
                        print_header = row == rows[0] and bool(re.match(r"^\s+文件编号[：:]", str(value or "")))
                        _fill_cell(target, source, value, scale, line_height, print_header)
                    except (ValueError, KeyError) as exc:
                        raise ValueError(f"{path.name}：打印页{source.coordinate}转换失败：{exc}") from exc
                grid_column = 0
                for cell in target_row._tr.tc_lst:
                    edges = _border_edges(sheet, row, columns[grid_column],
                                          columns[grid_column + cell.grid_span - 1], merge_map, repairs)
                    _cell_borders(cell.get_or_add_tcPr(), edges, scale)
                    grid_column += cell.grid_span
            widened.extend(_widen_labels(table))
            included[index]["signature_at_bottom"] = _anchor_signature(table, projects is not None)
            included[index]["body_stop"] = len(doc.element.body) - 1
        if excluded:
            warnings.append(f"{path.name}：按报告项目及编号筛选，追加{len(included)}组Excel记录，"
                            f"排除{len(excluded)}组未选项目或无本次实测结果的记录；原Excel未删除。")
        return doc, {"source": path.name, "sheet": "打印页", "page_groups": len(pages),
                     "source_page_groups": source_page_count, "included_pages": included,
                     "excluded_pages": excluded, "border_repairs": sorted(repairs),
                     "widened_labels": list(dict.fromkeys(widened)), "warnings": warnings}
    finally:
        formulas.close()
        cached.close()
