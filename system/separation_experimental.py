"""Isolated experimental jobs; the legacy splitter is used only as stable I/O."""

from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import logging
import os
from pathlib import Path
import shutil
import threading
import unicodedata
import uuid
from typing import List

from docx import Document
from fastapi import APIRouter, BackgroundTasks, Depends, File, UploadFile
from fastapi.responses import Response
from download_transport import DownloadOptions, download_options, download_response
from pydantic import BaseModel, ConfigDict, Field, StrictBool
from starlette.concurrency import run_in_threadpool
import pymupdf

import documents as d
from docx_images import remove_unused_images
import separation as legacy
import separation_documents as sd
import separation_layout as original_layout
import separation_experimental_documents as edits
import separation_experimental_layout as layout
import separation_experimental_format as formatting
from storage import ConflictError


API = "/api/separations-experimental"
NOTICE = ("试验素材：AAA/BBB/CCC等仅为占位型号；观察顶部及项目记录显式样机编号列使用本方案A001/B001等编号。仅清空所有识别出的试验结果数据格；"
          "日期、人员、设备、环境、结论及其他原文保留旧拆分行为，须人工核对。"
          "多行观察及多型号结果与报告生成端的兼容尚待后续适配；不自动入生产素材库。")


class Part(legacy.SplitPart):
    source_id: str = ""
    manual_filename: bool = False
    groups: List[edits.ModelGroup] = Field(default_factory=list, max_length=6)
    selected: List[str] = Field(default_factory=list, max_length=120)


class Plan(BaseModel):
    model_config = ConfigDict(extra="forbid")
    root_directory_token: str
    normalize_result_widths: StrictBool = False
    parts: List[Part] = Field(min_length=1, max_length=100)


class Export(legacy.ExportRequest):
    configuration: Plan


def fingerprint(request):
    return hashlib.sha256(json.dumps(request.model_dump(), sort_keys=True, ensure_ascii=False).encode()).hexdigest()


class ExperimentalSeparations(legacy.Separations):
    def __init__(self, directory, check_docx):
        self.directory = Path(directory) / "separations-experimental"
        self.directory.mkdir(exist_ok=True)
        self.lock = threading.RLock()
        self.check_docx = check_docx
        for path in self.directory.glob("*/job.json"):
            job = json.loads(path.read_text(encoding="utf-8"))
            if job["status"] in ("analyzing", "exporting"):
                job.update(status="failed", error="试验服务处理中断；未自动重试或写入，请核对后显式重试。")
                self.save(job)

    def update(self, identity, **values):
        if "_units" in values and values.get("parts"):
            values["_units"] = self.named_units(self.get(identity), values["_units"])
            for part in values["parts"]:
                part["source_id"] = part["id"]
                part["groups"] = [] if part["kind"] == "items" else [{"model": "AAA", "count": 3}]
                part["selected"] = [] if part["kind"] == "items" else ["A001", "A002", "A003"]
                part["filename"] = edits.suggested_filename(part["id"], part["kind"],
                    [edits.ModelGroup(**g) for g in part["groups"]], part["selected"],
                    next(u for u in values["_units"] if u["id"] == part["source_id"]))
            for page in values.get("pages", []):
                page["image_url"] = page["image_url"].replace("/api/separations/", API + "/")
        return super().update(identity, **values)

    def named_units(self, job, units=None):
        units = deepcopy(job.get("_units", []) if units is None else units)
        if any(unit["kind"] == "result" for unit in units):
            doc = Document(self.source(job))
            for unit in units:
                if unit["kind"] == "result":
                    unit["subprojects"] = edits.result_subprojects(doc, unit)
                    unit["display_project"] = edits.result_project_title(unit)
                    unit["folder_name"] = edits.result_project_folder(unit) if unit.get("sequence") else ""
        return units

    def public(self, job):
        result = super().public(job)
        result["normalize_result_widths"] = job.get("normalize_result_widths", False)
        result["warnings"] = list(dict.fromkeys([NOTICE, *result.get("warnings", [])]))
        result["source_units"] = [
            {key: unit.get(key, "") for key in
             ("id", "kind", "project", "sequence", "start_page", "end_page",
              "subprojects", "display_project", "folder_name")}
            for unit in self.named_units(job)
        ]
        return result

    def output_directory(self, job, root_token, part, units):
        base = self.destination(job, root_token, part)
        if part.kind != "result":
            return base
        if base.exists() and not base.is_dir():
            raise ValueError("结果类别保存位置不是文件夹，请核对保存目录。")
        if len(units) != 1:
            raise ValueError("每份results必须对应一个明确的试验项目，请分别建立卡片。")
        target = base / edits.result_project_folder(units[0])
        if base.exists():
            matches = [path for path in base.iterdir() if
                       unicodedata.normalize("NFC", path.name).casefold() ==
                       unicodedata.normalize("NFC", target.name).casefold()]
            if len(matches) > 1:
                raise ValueError("项目文件夹名称存在大小写或Unicode重名，请核对保存目录。")
            if matches:
                target = matches[0]
        if (target.is_symlink() or not target.resolve().is_relative_to(base.resolve())
                or (target.exists() and not target.is_dir())):
            raise ValueError("项目保存位置不是安全的实际文件夹，请核对目录，不能使用文件或符号链接。")
        return target

    def plan(self, identity, request):
        with self.lock:
            job = self.get(identity)
            if job["status"] not in ("ready", "completed", "partial", "failed") or not job.get("_units"):
                raise ConflictError("请等待原生分析完成。")
            job.update(plan=None, _plan=None)
            self.save(job)
            source = self.source(job)
            job["_units"] = self.named_units(job)
            if len({part.id for part in request.parts}) != len(request.parts):
                raise ValueError("方案标识重复。")
            outputs, adjustments, names, items_pages, project_folders = [], [], set(), set(), {}
            for part in request.parts:
                if part.kind == "items":
                    if part.groups or part.selected:
                        raise ValueError("items不支持型号改写，必须保留原内容。")
                else:
                    edits.selection(part.groups, part.selected, part.kind)
                if not part.enabled:
                    continue
                if not part.start_page <= part.end_page <= job["total_pages"]:
                    raise ValueError("所选页码必须在原报告物理页范围内。")
                units = sorted([u for u in job["_units"] if u["kind"] == part.kind
                    and u["start_page"] <= part.end_page and u["end_page"] >= part.start_page],
                    key=lambda u: u["start_block"])
                if part.source_id:
                    source_unit = next((u for u in job["_units"] if u["id"] == part.source_id
                                        and u["kind"] == part.kind), None)
                    if source_unit is None:
                        raise ValueError("源项目标识无效或类别已改变，请重新选择源范围。")
                    if source_unit["start_page"] <= part.start_page <= part.end_page <= source_unit["end_page"]:
                        # Two complete projects can share a physical page.
                        # A copied card retains its exact semantic source unit.
                        units = [source_unit]
                if not units:
                    raise ValueError("所选范围未包含此类别的完整章节/项目。")
                first, last = units[0], units[-1]
                # Never include another section merely because it sits between
                # two selected units of the same kind.
                if any(u["kind"] != part.kind and first["start_block"] < u["start_block"] < last["end_block"]
                       for u in job["_units"]):
                    raise ValueError("所选范围跨越其他类别，不能合并输出。")
                start, end = first["start_page"], last["end_page"]
                if part.kind == "items":
                    pages = set(range(start, end + 1))
                    if pages & items_pages:
                        raise ValueError("items沿用原拆分安全行为，不能重复导出同一范围。")
                    items_pages |= pages
                if (start, end) != (part.start_page, part.end_page):
                    adjustments.append({"id": part.id, "requested_start": part.start_page,
                        "requested_end": part.end_page, "start_page": start, "end_page": end,
                        "message": "扩展至完整语义章节/项目；不切断跨页表格。"})
                filename = legacy.safe_filename(part.filename)
                base_directory = self.destination(job, request.root_directory_token, part)
                directory = self.output_directory(job, request.root_directory_token, part, units)
                # Experimental files must never silently become production assets.
                production = Path(__file__).resolve().parents[1] / "report model"
                if directory.resolve().is_relative_to(production.resolve()):
                    raise ValueError("试验输出禁止写入生产report model素材目录，请选择独立试验目录。")
                target = directory / filename
                if part.kind == "result":
                    folder_key = unicodedata.normalize("NFC", str(directory)).casefold()
                    if folder_key in project_folders and project_folders[folder_key] != first["id"]:
                        raise ValueError("不同源项目对应同名项目文件夹，请核对原报告序号和名称。")
                    project_folders[folder_key] = first["id"]
                key = unicodedata.normalize("NFC", str(target)).casefold()
                if key in names or (directory.exists() and any(
                        unicodedata.normalize("NFC", p.name).casefold() == unicodedata.normalize("NFC", filename).casefold()
                        for p in directory.iterdir())):
                    raise ValueError("文件重名，不能覆盖：" + str(target))
                names.add(key)
                original = Document(source)
                fragment = edits.planning_fragment(original, first["start_block"], last["end_block"])
                try:
                    checks = edits.transform(fragment, part.kind, part.groups, part.selected)
                    widths = (formatting.result_widths(fragment, enabled=request.normalize_result_widths,
                              project=first["project"]) if part.kind == "result" else None)
                except ValueError as exc:
                    raise ValueError(f"{filename}（P{start}～P{end}）：{exc}") from exc
                labels = edits.selection(part.groups, part.selected, part.kind) if part.kind != "items" else []
                outputs.append({"id": part.id, "kind": part.kind, "filename": filename,
                    "directory": str(directory), "base_directory": str(base_directory),
                    "full_path": str(target), "start_page": start, "end_page": end,
                    "models": edits.roster(part.groups), "selected": labels, "structure": checks,
                    "result_widths": widths,
                    "source_projects": [{"id": u["id"], "name": u.get("project", u["id"]),
                                         "display_name": u.get("display_project", u.get("project", u["id"])),
                                         "start_block": u["start_block"], "end_block_exclusive": u["end_block"]}
                                        for u in units],
                    "format_adjustments": formatting.HEADER_DESCRIPTION + "items正文保持原拆分；观察顶部及项目记录显式样机编号列使用本方案编号，有样机编号列的项目优先整项同页，当前页放不下则整体移至新页，本身超过一页仍续写，＋栏逐行对应编号；按真实Word分页续名称/叉号；结果表头及数据行最小0.8cm，序号、项目名称及技术要求各列分别按同表上下相邻文字相同纵向合并，不以格式、序号或上级名称变化阻断其他列合并，合并格沿用首格格式，样机编号和试验结果不合并，跨页独立续名，按V1续表头及签名模板排页底；页码重算。"
                                          + ("燃气表流量范围按型号数量分段，型号依次为AAA、BBB、CCC等；合成时按总览流量范围替换，分界流量=0.1×最大流量。"
                                             if "gas_flow_ranges" in checks else "")
                                          + (formatting.RESULT_TITLE_DESCRIPTION + formatting.RESULT_INTRO_DESCRIPTION
                                             + formatting.RESULT_CLOSING_DESCRIPTION
                                             if part.kind == "result" else "")
                                          + (widths["description"] if widths else ""),
                    "_start_block": first["start_block"], "_end_block": last["end_block"]})
            if not outputs:
                raise ValueError("请至少启用一个方案。")
            revision = job.get("_plan_revision", 0) + 1
            plan = {"revision": revision, "outputs": outputs, "adjustments": adjustments,
                    "normalize_result_widths": request.normalize_result_widths, "layout_policy": layout.POLICY}
            public = deepcopy(plan)
            public["outputs"] = [{k: v for k, v in o.items() if not k.startswith("_")} for o in outputs]
            job.update(parts=[p.model_dump() for p in request.parts], root_directory_token=request.root_directory_token,
                normalize_result_widths=request.normalize_result_widths,
                root_directory_label=job.get("_directories", {}).get(request.root_directory_token, ""),
                _configuration=fingerprint(request), _plan=plan, plan=public, _plan_revision=revision,
                status="ready", error="")
            self.save(job)
            return job

    def begin_export(self, identity, request):
        with self.lock:
            job = self.get(identity)
            if (job["status"] != "ready" or not job.get("_plan") or not request.confirm_adjustments
                    or request.plan_revision != job["_plan"]["revision"]
                    or fingerprint(request.configuration) != job.get("_configuration")
                    or job["_plan"]["layout_policy"] != layout.POLICY):
                raise ConflictError("方案已改变、尚未确认或旧计划失效，请重新保存检查。")
            return self.update(identity, status="exporting", error="", exports=[], _exports=[],
                export_failures=[], stage="逐份核验试验素材；失败文件跳过，核验结束后统一保存合格文件")

    def export(self, identity):
        current = ""
        exports, failures = [], []

        def record_failure(output, phase, exc):
            logging.getLogger("report_workbench").exception(
                "Experimental split file failed (%s): %s", phase, current)
            failures.append({"id": output["id"], "filename": output["filename"],
                "path": output["full_path"], "error": str(exc), "phase": phase})
            self.update(identity, export_failures=list(failures))

        try:
            job = self.get(identity)
            source, folder = self.source(job), self.folder(identity)
            if legacy.file_hash(folder / "source.pdf") != job["_pdf_hash"]:
                raise ValueError("源分页基准已改变，请重新分析。")
            staging = folder / ("export-" + uuid.uuid4().hex)
            staging.mkdir()
            verified = []
            outputs = job["_plan"]["outputs"]
            for index, output in enumerate(outputs, 1):
                current = (f"{Path(output['directory']).name}/{output['filename']}"
                           if output["kind"] == "result" else output["filename"])
                try:
                    part = next(Part.model_validate(p) for p in job["parts"] if p["id"] == output["id"])
                    path = staging / (part.id + ".docx")
                    sd.extract(source, path, output["_start_block"], output["_end_block"], 1, job["total_pages"],
                               job["_numbered"], job["_bookmarks"], freeze_total_pages=False,
                               report_number=sd.EXPORT_REPORT_NUMBER)
                    ranges = [(u["start_block"] - output["_start_block"], u["end_block"] - output["_start_block"])
                        for u in job["_units"] if u["kind"] == "result"
                        and output["_start_block"] <= u["start_block"] < output["_end_block"]]
                    original_layout.normalize(path, part.kind, ranges, result_format=part.kind == "result",
                                              split_result_tables=False)
                    doc = Document(path)
                    structure = edits.transform(doc, part.kind, part.groups, part.selected)
                    widths = (formatting.result_widths(doc, enabled=job.get("normalize_result_widths", False),
                              project=output["source_projects"][0]["name"]) if part.kind == "result" else None)
                    if widths != output.get("result_widths"):
                        raise ValueError("结果列宽选项与已确认计划不一致，请重新保存检查。")
                    doc.save(path)
                    checks = layout.render_verify(path, part.kind, structure,
                        progress=lambda stage: self.progress(identity,
                            f"核验 {index}/{len(outputs)} · 已跳过{len(failures)}份 · {current} · {stage}"))
                    checks.update(source_start_page=output["start_page"], source_end_page=output["end_page"])
                    if widths is not None:
                        if widths["applied"]:
                            formatting.verify_result_widths(Document(path))
                        checks["result_widths"] = widths
                    self.progress(identity, "清理未使用图片 · " + current)
                    checks["image_cleanup"] = remove_unused_images(path)
                    verified.append((output, path, checks, part))
                except (OSError, ValueError, RuntimeError) as exc:
                    record_failure(output, "核验", exc)
                self.progress(identity, f"已核验{index}/{len(outputs)}份；合格{len(verified)}份，跳过{len(failures)}份")
            current = "源文件复核"
            self.source(job)
            if legacy.file_hash(folder / "source.pdf") != job["_pdf_hash"]:
                raise ValueError("源分页基准已改变，请重新分析。")
            for output, path, checks, part in verified:
                current = (f"{Path(output['directory']).name}/{output['filename']}"
                           if output["kind"] == "result" else output["filename"])
                created = None
                try:
                    units = [u for u in job["_units"] if any(
                        source["id"] == u["id"] for source in output["source_projects"])]
                    directory = self.output_directory(job, job["root_directory_token"], part, units)
                    target = Path(output["full_path"])
                    if directory != target.parent or target.name != legacy.safe_filename(part.filename):
                        raise ValueError("目录授权已改变。")
                    base = self.destination(job, job["root_directory_token"], part)
                    base.mkdir(exist_ok=True)
                    directory.mkdir(exist_ok=True)
                    if self.output_directory(job, job["root_directory_token"], part, units) != directory:
                        raise ValueError("项目目录已改变，停止写入。")
                    if any(unicodedata.normalize("NFC", p.name).casefold() ==
                           unicodedata.normalize("NFC", target.name).casefold() for p in directory.iterdir()):
                        raise FileExistsError("文件已存在，不覆盖：" + str(target))
                    with target.open("xb") as handle:
                        stat = os.fstat(handle.fileno())
                        created = (target, stat.st_dev, stat.st_ino)
                        with path.open("rb") as data:
                            shutil.copyfileobj(data, handle)
                        handle.flush()
                        os.fsync(handle.fileno())
                    if legacy.file_hash(target) != legacy.file_hash(path):
                        raise ValueError("写入后文件校验失败。")
                    exports.append({"id": part.id, "filename": target.name, "path": str(target),
                        "checks": checks, "url": f"{API}/{identity}/files/{part.id}",
                        "_copy": str(path.relative_to(folder))})
                except (OSError, ValueError, RuntimeError) as exc:
                    if created is not None:
                        target, device, inode = created
                        if target.exists() and (target.stat().st_dev, target.stat().st_ino) == (device, inode):
                            target.unlink()
                    record_failure(output, "保存", exc)
                self.update(identity,
                    exports=[{k: v for k, v in x.items() if not k.startswith("_")} for x in exports],
                    _exports=exports, stage=f"统一保存中；已保存{len(exports)}份，失败{len(failures)}份")
            summary = f"处理结束：成功保存{len(exports)}份，失败{len(failures)}份"
            error = "\n".join(f"{item['path']}：{item['error']}" for item in failures)
            self.update(identity, status=("partial" if exports else "failed") if failures else "completed",
                error=error, stage=summary + "；失败文件未交付，报告端兼容尚待适配",
                export_failures=failures)
        except Exception as exc:
            logging.getLogger("report_workbench").exception("Experimental split failed")
            self.update(identity, status="failed", error=current + "：" + str(exc),
                stage=f"任务中断；已保存{len(exports)}份保留，其余未交付；未跳过Word核验",
                exports=[{k: v for k, v in x.items() if not k.startswith("_")} for x in exports],
                _exports=exports, export_failures=failures)


def routes(directory, check_docx):
    service = ExperimentalSeparations(directory, check_docx)
    router = APIRouter(prefix=API)

    @router.get("")
    def listing():
        return sorted([{key: job[key] for key in ("id", "name", "status", "created_at")}
            for path in service.directory.glob("*/job.json") for job in [service.get(path.parent.name)]],
            key=lambda j: j["created_at"], reverse=True)

    @router.post("/upload")
    async def upload(background: BackgroundTasks, file: UploadFile = File(...)):
        if not file.filename or not file.filename.lower().endswith(".docx"):
            raise ValueError("请选择可编辑DOCX源报告。")
        identity = uuid.uuid4().hex
        folder = service.folder(identity)
        folder.mkdir()
        path = folder / "source.docx"
        try:
            size = 0
            with path.open("xb") as handle:
                while True:
                    chunk = await file.read(1024 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    if size > legacy.MAX_UPLOAD:
                        raise ValueError("报告最大支持32MB。")
                    handle.write(chunk)
            check_docx(path)
        except (ValueError, OSError):
            path.unlink(missing_ok=True)
            folder.rmdir()
            raise
        finally:
            await file.close()
        job = {"id": identity, "name": Path(file.filename).name, "status": "uploaded", "stage": "", "error": "",
            "created_at": datetime.now(timezone.utc).isoformat(), "_source_hash": legacy.file_hash(path),
            "pages": [], "parts": [], "warnings": [], "exports": [], "plan": None,
            "total_pages": 0, "model_count": None, "root_directory_token": "", "root_directory_label": ""}
        service.save(job)
        job = service.begin_analysis(identity)
        background.add_task(service.analyze, identity)
        return service.public(job)

    @router.get("/{identity}")
    def detail(identity: str):
        return service.public(service.get(identity))

    @router.post("/{identity}/analyze")
    def analyze(identity: str, background: BackgroundTasks):
        job = service.begin_analysis(identity)
        background.add_task(service.analyze, identity)
        return service.public(job)

    @router.post("/{identity}/folder")
    async def folder(identity: str, request: legacy.FolderRequest):
        return await run_in_threadpool(service.choose_folder, identity)

    @router.post("/{identity}/plan")
    def plan(identity: str, request: Plan):
        return service.public(service.plan(identity, request))

    @router.post("/{identity}/invalidate")
    def invalidate(identity: str):
        with service.lock:
            job = service.get(identity)
            if job["status"] == "exporting":
                raise ConflictError("正在导出，不能修改。")
            return service.public(service.update(identity, plan=None, _plan=None))

    @router.post("/{identity}/export")
    def export(identity: str, request: Export, background: BackgroundTasks):
        job = service.begin_export(identity, request)
        background.add_task(service.export, identity)
        return service.public(job)

    @router.get("/{identity}/source.pdf")
    def pdf(identity: str, options: DownloadOptions = Depends(download_options)):
        if not service.get(identity).get("_pdf_hash"):
            raise ValueError("源报告尚未完成Word分页。")
        return download_response(service.folder(identity) / "source.pdf", options, media_type="application/pdf")

    @router.get("/{identity}/pages/{page}.png")
    def page(identity: str, page: int):
        if not 1 <= page <= service.get(identity)["total_pages"]:
            raise ValueError("预览页超出范围。")
        with pymupdf.open(service.folder(identity) / "source.pdf") as doc:
            return Response(doc[page - 1].get_pixmap(matrix=pymupdf.Matrix(.9, .9), alpha=False).tobytes("png"),
                            media_type="image/png")

    @router.get("/{identity}/files/{file_id}")
    def output(identity: str, file_id: str, options: DownloadOptions = Depends(download_options)):
        item = next((x for x in service.get(identity).get("_exports", []) if x["id"] == file_id), None)
        if not item:
            raise KeyError(file_id)
        return download_response(service.folder(identity) / item["_copy"], options, filename=item["filename"])

    return router
