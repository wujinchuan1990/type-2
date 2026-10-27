"""Gas-meter charts from current records and selected result requirements."""

from copy import deepcopy
from functools import lru_cache
from io import BytesIO
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw

import documents
from curve_results import _manual_result
from gas_records import FLOWS, _key, compact, decimal, format_result, project_kind
from word_fonts import font as word_font


CURVE_HEADINGS = {"误差曲线": "误差曲线图：", "耐久性": "耐久性试验后误差曲线图："}
NAVY, MUTED, AMBER = "#162d49", "#728196", "#b37b33"


@lru_cache(maxsize=24)
def _font(size, chinese=False, italic=False):
    return word_font("宋体" if chinese else "Arial Italic" if italic else "Arial").font_variant(size=size)


def _flow_label(draw, x, y, flow):
    multiplier, suffix = flow.split("q")
    regular, italic, small = _font(26), _font(26, italic=True), _font(18)
    multiplier_width = draw.textlength(multiplier, font=regular)
    q_width = draw.textlength("q", font=italic)
    width = multiplier_width + q_width + draw.textlength(suffix, font=small)
    left, baseline = x - width / 2, y + 26
    draw.text((left, baseline), multiplier, font=regular, fill=NAVY, anchor="ls")
    draw.text((left + multiplier_width, baseline), "q", font=italic, fill=NAVY, anchor="ls")
    draw.text((left + multiplier_width + q_width, baseline + 7), suffix,
              font=small, fill=NAVY, anchor="ls")


def _dashed(draw, start, end, width=3):
    x1, y1 = start
    x2, y2 = end
    length = ((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5
    for distance in range(0, int(length), 19):
        stop = min(distance + 11, length)
        draw.line((x1 + (x2 - x1) * distance / length, y1 + (y2 - y1) * distance / length,
                   x1 + (x2 - x1) * stop / length, y1 + (y2 - y1) * stop / length),
                  fill=AMBER, width=width)


def curve_source(project, projects, sample):
    name = project_kind(project)
    source_name = "准确度等级和最大允许误差" if name == "误差曲线" else "耐久性"
    candidates = [p for p in projects if project_kind(p) == source_name
                  and any(row.get("sample_id") == sample for row in p["rows"])]
    if len(candidates) != 1:
        raise ValueError("缺少唯一对应的“" + source_name + "”所选项目，无法核对本样机误差限")
    return candidates[0]


def curve_limits(source, sample):
    limits = {"low": set(), "high": set()}
    for row in source["rows"]:
        if row.get("sample_id") != sample:
            continue
        requirement = compact(row["requirement"]).lower().replace("<=", "≤")
        if "最大允许误差" not in requirement:
            continue
        matches = re.findall(r"最大允许误差[：:]?±(\d+(?:\.\d+)?)%", requirement)
        band = "low" if "qmin≤q<qt" in requirement else "high" if "qt≤q≤qmax" in requirement else ""
        if len(matches) != 1 or not band:
            raise ValueError("最大允许误差的流量区间或正负限值无法唯一识别")
        limits[band].add(decimal(matches[0]))
    if any(len(values) != 1 for values in limits.values()):
        raise ValueError("低／高流量段最大允许误差缺失或冲突")
    low, high = next(iter(limits["low"])), next(iter(limits["high"]))
    if min(low, high) <= 0:
        raise ValueError("最大允许误差幅值须为正数")
    return low, high


def curve_with_manual_values(curve, source):
    result = deepcopy(curve)
    sample = curve["sample_id"]
    for point in result["points"]:
        rows = [row for row in source["rows"] if row.get("sample_id") == sample
                and _key(project_kind(source), row, source["rows"]) == "mean:" + point["flow"]]
        manual = [row for row in rows if _manual_result(row)]
        if not manual:
            continue
        if len(rows) != 1:
            raise ValueError("同一流量点存在多条人工结果，不能选择其中一条作图")
        row = manual[0]
        if row.get("matched_sample_id") and row["matched_sample_id"] != sample:
            raise ValueError("人工误差结果属于旧样机编号，请先核对")
        match = re.fullmatch(r"([+\-]?\d+(?:\.\d+)?)%?", compact(row["value"]))
        if not match:
            raise ValueError("人工误差结果为空或不是有效百分数，不能使用旧记录替代")
        point["error"] = str(decimal(match[1]))
    return result


def plot_curve(curve, kind, limits):
    points = curve["points"]
    if [point["flow"] for point in points] != list(reversed(FLOWS)):
        raise ValueError("燃气表误差曲线须包含按流量顺序排列的8个完整测量点。")
    if kind not in CURVE_HEADINGS:
        raise ValueError("不支持的燃气表曲线类型：" + kind)
    if not curve.get("sample_id", "").strip() or not curve.get("model", "").strip():
        raise ValueError("曲线缺少本次样机编号或完整型号，请重新读取原始记录")
    values = [decimal(point["error"]) for point in points]
    low, high = map(decimal, limits)
    if min(low, high) <= 0:
        raise ValueError("曲线误差限须为正数")
    bound, step = (4.0, 1.0) if kind == "误差曲线" else (7.5, 1.5)
    if max(abs(value) for value in values) >= bound or max(low, high) >= bound:
        raise ValueError("实测误差或限值超出曲线纵轴±%s%%，请核对数据与图轴范围，未裁掉超范围点" % bound)
    intervals = round(2 * bound / step)
    color = "#2863b0" if kind == "误差曲线" else "#16847f"
    heading = "初始误差曲线" if kind == "误差曲线" else "耐久性试验后误差曲线"
    image = Image.new("RGB", (1480, 670), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((0, 0, 1479, 669), radius=16, outline="#e0e7ef", width=2)
    draw.rounded_rectangle((38, 38, 44, 74), radius=3, fill=color)
    draw.text((64, 32), heading, font=_font(33, chinese=True), fill=NAVY, stroke_width=1)
    identity = curve["sample_id"] + "  /  " + curve["model"]
    if draw.textlength(identity, font=_font(23)) > 880:
        raise ValueError("样机编号及型号过长，无法完整放入曲线标题区，请核对")
    draw.text((1435, 50), identity, font=_font(23), fill=MUTED, anchor="rt")
    draw.text((66, 125), "平均示值误差（%）", font=_font(23, chinese=True), fill=MUTED)
    left, right, top, bottom = 130, 1400, 192, 540

    def y_for(value):
        return top + (bound - float(value)) / (2 * bound) * (bottom - top)

    for index in range(intervals + 1):
        value = bound - index * step
        y = y_for(value)
        draw.line((left, y, right, y), fill="#bbc5d1" if value == 0 else "#e9eef4", width=1)
        draw.text((left - 24, y), f"{value:.1f}", font=_font(23), fill=MUTED, anchor="rm")
    positions = [(left + 42 + index * (right - left - 84) / 7, y_for(value))
                 for index, value in enumerate(values)]
    boundary = positions[3][0]
    occupied = []
    for sign in (1, -1):
        low_y, high_y = y_for(sign * low), y_for(sign * high)
        _dashed(draw, (left, low_y), (boundary, low_y))
        _dashed(draw, (boundary, low_y), (boundary, high_y), width=2)
        _dashed(draw, (boundary, high_y), (right, high_y))
        draw.ellipse((boundary - 5, low_y - 5, boundary + 5, low_y + 5),
                     fill="white", outline=AMBER, width=2)
        draw.ellipse((boundary - 5, high_y - 5, boundary + 5, high_y + 5), fill=AMBER)
        for x, y, limit in [((left + boundary) / 2, low_y, sign * low),
                            ((boundary + right) / 2, high_y, sign * high)]:
            text, text_y = f"{limit:+}%", y - 18 if sign > 0 else y + 18
            box = draw.textbbox((x, text_y), text, font=_font(22), anchor="mm")
            occupied.append(box)
            draw.rectangle((box[0] - 5, box[1] - 2, box[2] + 5, box[3] + 2), fill="white")
            draw.text((x, text_y), text, font=_font(22), fill=AMBER, anchor="mm")
    draw.line(positions, fill=color, width=4)
    for index, ((x, y), value) in enumerate(zip(positions, values)):
        draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill="white", outline=color, width=4)
        text = format_result(value)
        limit = low if index < 3 else high
        for offset in (21, -24, 42, -45):
            text_y = y + offset
            box = draw.textbbox((x, text_y), text, font=_font(21), anchor="mm")
            clashes = any(box[0] < old[2] + 4 and box[2] > old[0] - 4
                          and box[1] < old[3] + 4 and box[3] > old[1] - 4 for old in occupied)
            if (top < box[1] - 2 and box[3] + 2 < bottom and not clashes
                    and not any(box[1] - 4 <= y_for(sign * limit) <= box[3] + 4 for sign in (1, -1))):
                break
        else:
            raise ValueError("流量点%s的实测值标注无法避开限值线，请核对图轴范围" % points[index]["flow"])
        occupied.append(box)
        draw.rounded_rectangle((box[0] - 4, box[1] - 2, box[2] + 4, box[3] + 2),
                               radius=3, fill="white")
        draw.text((x, text_y), text, font=_font(21), fill=color, anchor="mm")
    for (x, _), point in zip(positions, points):
        _flow_label(draw, x, 563, point["flow"])
    draw.line((64, 634, 101, 634), fill=color, width=4)
    draw.text((112, 634), "实测误差", font=_font(19, chinese=True), fill=MUTED, anchor="lm")
    _dashed(draw, (255, 634), (306, 634))
    draw.text((318, 634), "最大允许误差上 / 下限", font=_font(19, chinese=True), fill=MUTED, anchor="lm")
    stream = BytesIO()
    image.save(stream, format="PNG", dpi=(240, 240))
    stream.seek(0)
    return stream


def clear_curve_pictures(doc, heading):
    body = doc.element.body
    old = []
    for block in list(body)[body.index(heading) + 1:]:
        text = documents._text(block).strip()
        if text and not (text.startswith("编号") and ("误差曲线图" in text or "未生成曲线图" in text)):
            break
        if block.tag != documents.qn("w:p"):
            break
        old.append(block)
    for block in old:
        body.remove(block)
    removed_ids = {identity for block in old for identity in
                   documents._xp(block, ".//@r:embed | .//@r:id | .//@r:link")}
    retained_ids = set(documents._xp(doc.element, ".//@r:embed | .//@r:id | .//@r:link"))
    for identity in removed_ids - retained_ids:
        doc.part.drop_rel(identity)


def _chart_anchor(doc, kind):
    body = doc.element.body
    judges = [p for p in documents._xp(body, "./w:p") if "本试验项目合格判定要求" in documents._text(p)]
    headings = [p for p in documents._xp(body, "./w:p") if documents._text(p).strip() == CURVE_HEADINGS[kind]]
    if len(judges) != 1 or len(headings) > 1:
        raise ValueError("燃气表曲线图合格判定要求／曲线标题缺失或不唯一，请核对结果素材。")
    tables = [block for block in list(body)[:body.index(judges[0])]
              if block.tag == documents.qn("w:tbl") and "试验结果" in documents._text(block)]
    if not tables:
        raise ValueError("燃气表曲线图前未找到试验数据记录表格。")
    if headings:
        heading = headings[0]
        clear_curve_pictures(doc, heading)
    else:
        heading = doc.add_paragraph(CURVE_HEADINGS[kind])._p
    tables[-1].addnext(heading)
    if not headings:
        clear_curve_pictures(doc, heading)
    paragraph = documents.Paragraph(heading, doc)
    paragraph.paragraph_format.page_break_before = False
    paragraph.paragraph_format.keep_with_next = True
    return heading


def project_curve_images(project, projects):
    kind = project_kind(project)
    if kind not in CURVE_HEADINGS:
        return []
    result = []
    samples = list(dict.fromkeys(row.get("sample_id") or row["sample_label"] for row in project["rows"]))
    for sample in samples:
        curves = [curve for curve in project.get("gas_curves", []) if curve["sample_id"] == sample]
        item = {"sample_id": sample, "model": "", "source": "", "limit_source": "", "image": None, "warning": ""}
        try:
            if len(curves) != 1:
                raise ValueError("缺少唯一完整的本次Excel%s误差数据" % ("初始" if kind == "误差曲线" else "耐久性复测"))
            source = curve_source(project, projects, sample)
            limits = curve_limits(source, sample)
            curve = curve_with_manual_values(curves[0], source)
            item.update(image=plot_curve(curve, kind, limits), model=curve["model"],
                        source=curve["source"], limit_source=source["name"])
        except ValueError as exc:
            item["warning"] = "编号%s：%s，未生成曲线图。" % (sample, exc)
        result.append(item)
    return result


def fill_curves(doc, project, projects):
    kind = project_kind(project)
    if kind not in CURVE_HEADINGS:
        return []
    cursor = _chart_anchor(doc, kind)
    warnings = []
    for item in project_curve_images(project, projects):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(6)
        if item["image"] is not None:
            shape = paragraph.add_run().add_picture(item["image"], width=Cm(12))
            shape._inline.docPr.set("descr", "%s；样机%s；型号%s；限值来源：%s；数据来源：%s" % (
                CURVE_HEADINGS[kind], item["sample_id"], item["model"], item["limit_source"], item["source"]))
        else:
            warnings.append(item["warning"])
            paragraph.add_run(item["warning"]).font.size = Pt(12)
        cursor.addnext(paragraph._p)
        cursor = paragraph._p
    return warnings
