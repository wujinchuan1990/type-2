"""Compose the companion, unfilled original-record form from this export."""

import re
import json
from io import BytesIO
from copy import deepcopy
from math import floor
from pathlib import Path

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.parts.hdrftr import HeaderPart
from docx.shared import Pt
from docxcompose.composer import Composer

import documents as d
from parse_cache import ParseCache

_EXCEL_CACHE = ParseCache(limit=8)
from environment import aggregate_project_environment
from validation import parse_date
from word_fonts import font as _cover_font, text_width as _cover_text_width


TEMPLATE_PATH = Path("cover of original record/原始记录封面.docx")


def _one(elements, description):
    if len(elements) != 1:
        raise ValueError(f"原始记录{description}缺失或不唯一（{len(elements)}处）。")
    return elements[0]


def _cover_value(outline, label):
    blocks = list(outline.element.body)
    boundary = next((i for i, block in enumerate(blocks)
                     if d._norm(d._text(block)).startswith("批准日期")), None)
    if boundary is None:
        raise ValueError("原始记录封面：无法定位大纲封面的批准日期边界。")
    pattern = r"^\s*" + r"\s*".join(label) + r"\s*[：:]?\s*(\S.*?)\s*$"
    matches = [re.match(pattern, d.Paragraph(p, outline).text)
               for p in blocks[:boundary] if p.tag == qn("w:p")]
    values = [m.group(1).strip() for m in matches if m]
    return _one(values, f"封面来源「{label}」")


def cover_information(task, outline):
    fields = {f["key"]: f.get("value", "").strip() for f in task.get("overview", [])}
    values = {
        "report_number": fields.get("report_number", ""),
        "instrument_name": fields.get("instrument_name", ""),
        "manufacturer": fields.get("manufacturer", ""),
        "receipt_location": "张衡路1500号",
        "conclusion": "",
        "samples": "见后页",
        "evaluators": "", "reviewers": "",
    }
    warnings = []
    for key, reader in (("outline_number", lambda: _cover_value(outline, "编号")),
                        ("model", lambda: d.outline_cover_model(outline, "本次大纲")["value"])):
        try:
            values[key] = reader()
            if key == "outline_number":
                values[key] = re.sub(r"^D\s+(?=\S)", "D", values[key])
        except ValueError as exc:
            values[key] = ""
            warnings.append(str(exc))
    projects = task.get("projects", [])
    for key, choice in (("start", min), ("end", max)):
        days = []
        for project in projects:
            try:
                days.append(parse_date(project.get(key, "")))
            except ValueError as exc:
                warnings.append(f"原始记录封面：{project['name']}的{key}日期无法汇总（{exc}）。")
        values[key] = choice(days).isoformat() if days and len(days) == len(projects) else ""
    environment = aggregate_project_environment(projects)
    values.update(environment["values"])
    warnings.extend("原始记录封面：" + issue for issue in environment["issues"])
    for key, label in (("report_number", "报告编号"), ("instrument_name", "计量器具名称"),
                       ("manufacturer", "制造厂"), ("start", "受样日期/试验开始日期"),
                       ("end", "试验结束日期"), ("receipt_location", "受样地点"),
                       ("samples", "样机编号")):
        if not values.get(key):
            warnings.append(f"原始记录封面：{label}未填写，保留空白。")
    return values, warnings


def _cover_paragraph(doc, label, cover_end):
    pattern = re.compile(r"^\s*" + r"\s*".join(label))
    return _one([p for p in list(doc.element.body)[:doc.element.body.index(cover_end)]
                 if p.tag == qn("w:p") and pattern.match(d.Paragraph(p, doc).text)], f"封面「{label}」")


def _source_run_widths(runs):
    width, fitted = 0, set()
    for run in runs:
        text = "".join(d._xp(run, "./w:t/text()"))
        fit = run.rPr.find(qn("w:fitText")) if run.rPr is not None else None
        if fit is not None:
            identity = fit.get(qn("w:id"))
            if identity not in fitted:
                width += int(fit.get(qn("w:val"))) / 20
                fitted.add(identity)
            continue
        size = int(_one(d._xp(run, "./w:rPr/w:sz/@w:val"), "封面原字号")) / 2
        spacing = d._xp(run, "./w:rPr/w:spacing/@w:val")
        # The source cover uses the East Asian theme (SimSun), including its half-em spaces.
        width += _cover_font("宋体").getlength(text) * size / 1000
        width += len(text) * (int(spacing[0]) / 20 if spacing else 0)
    return width


def _cover_line(doc, paragraph, value, value_start=None, aligned_start=None, font_size=None):
    runs = d._xp(paragraph, "./w:r")
    if value_start is None:
        first = _one([r for r in runs if d._xp(r, "./w:rPr/w:u[not(@w:val='none')]")][:1],
                     "封面下划线起点")
        value_start = sum(len("".join(d._xp(r, "./w:t/text()"))) for r in runs[:runs.index(first)])
    prefix, offset = [], 0
    prototype = None
    for run in runs:
        text = "".join(d._xp(run, "./w:t/text()"))
        if offset < value_start:
            copy = deepcopy(run)
            nodes = d._xp(copy, "./w:t")
            for node in nodes:
                node.text = ""
            if nodes:
                d._write_slot_text(nodes[0], text[:value_start - offset])
            prefix.append(copy)
        if offset + len(text) > value_start and prototype is None:
            prototype = run
        offset += len(text)
    if prototype is None:
        raise ValueError("原始记录封面未找到填写区：" + d._text(paragraph))
    ppr = paragraph.get_or_add_pPr()
    indent = ppr.get_or_add_ind()
    left = int(indent.get(qn("w:left"), "0")) + int(indent.get(qn("w:firstLine"), "0"))
    section = doc.sections[0]
    right = round((section.page_width - section.left_margin - section.right_margin) / 635)
    right -= int(indent.get(qn("w:right"), "0"))
    start = left + round(_source_run_widths(prefix) * 20)
    if aligned_start is not None:
        # Trim only the label's trailing spacer, never its text or character sizing.
        for run in reversed(prefix):
            nodes = d._xp(run, "./w:t")
            for node in reversed(nodes):
                node.text = (node.text or "").rstrip()
            if d._text(run):
                break
        prefix_width = _source_run_widths(prefix) * 20
        spacer_width = aligned_start - left - prefix_width
        if spacer_width < 0:
            raise ValueError("原始记录封面标签超出统一下划线起点。")
        spacer = OxmlElement("w:r")
        props = deepcopy(prototype.rPr)
        for prop in d._xp(props, "./w:fitText | ./w:spacing | ./w:u"):
            props.remove(prop)
        fit = OxmlElement("w:fitText")
        fit.set(qn("w:val"), str(round(spacer_width)))
        fit.set(qn("w:id"), "999")
        props.append(fit)
        spacer.append(props)
        node = OxmlElement("w:t")
        spacer.append(node)
        d._write_slot_text(node, " ")
        prefix.append(spacer)
        start = aligned_start
    width = (right - start) / 20
    if width <= 2:
        raise ValueError("原始记录封面填写区宽度不足：" + d._text(paragraph))
    value = re.sub(r"[\r\n\t]+", " ", value).strip()
    original_size = int(_one(d._xp(prototype, "./w:rPr/w:sz/@w:val"), "封面填写字号")) / 2
    if font_size is not None:
        original_size = font_size
    measured = _cover_text_width(value, original_size)
    if font_size is not None and measured > width - 2:
        raise ValueError(f"原始记录封面试验日期无法以{font_size}磅完整单行显示，请检查模板横线宽度。")
    size = original_size if measured <= width - 2 else floor(original_size * (width - 2) / measured * 2) / 2
    if size < 0.5:
        raise ValueError("原始记录封面内容过长，最小字号仍无法单行容纳，请缩短填写内容。")
    for child in list(paragraph):
        if child is not ppr:
            paragraph.remove(child)
    for run in prefix:
        paragraph.append(run)
    tabs = ppr.get_or_add_tabs()
    for tab in list(tabs):
        tabs.remove(tab)
    for alignment, position in (("center", (start + right) // 2), ("right", right)):
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), alignment)
        tab.set(qn("w:pos"), str(position))
        tabs.append(tab)
    for tag in ("autoSpaceDE", "autoSpaceDN"):
        setting = ppr.find(qn("w:" + tag))
        if setting is None:
            setting = OxmlElement("w:" + tag)
            ppr.insert_element_before(setting, "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind")
        setting.set(qn("w:val"), "false")
    # Explicit stops pin the rule independently of content length and font size.
    for text, is_tab, underline, points in (
        ("", True, True, original_size),
        (value, False, True, size), ("", True, True, original_size),
    ):
        run = d.Paragraph(paragraph, doc).add_run()
        run._r.append(deepcopy(prototype.rPr))
        for prop in d._xp(run._r, "./w:rPr/w:spacing | ./w:rPr/w:fitText | ./w:rPr/w:w"):
            prop.getparent().remove(prop)
        run.font.size = Pt(points)
        complex_size = run._r.rPr.find(qn("w:szCs"))
        if complex_size is None:
            complex_size = OxmlElement("w:szCs")
            run._r.rPr.append(complex_size)
        complex_size.set(qn("w:val"), str(round(points * 2)))
        run.font.underline = underline
        if is_tab:
            run.add_tab()
        else:
            run.text = text


def _chinese_date(value):
    day = parse_date(value)
    return f"{day.year}年{day.month:02d}月{day.day:02d}日"


def _fill_cover(doc, boundary, values):
    reference = _cover_paragraph(doc, "型式评价报告编号", boundary)
    prefix = []
    for run in d._xp(reference, "./w:r"):
        if d._xp(run, "./w:rPr/w:u"):
            break
        prefix.append(run)
    aligned_start = int(reference.pPr.ind.get(qn("w:left"), "0")) + round(_source_run_widths(prefix) * 20)
    labels = {
        "型式评价报告编号": "report_number", "大纲编号": "outline_number",
        "计量器具名称": "instrument_name", "型号": "model", "制造厂": "manufacturer",
        "样机编号": "samples", "受样地点": "receipt_location", "受样日期": "start",
        "试验结论": "conclusion", "评价人员": "evaluators", "复核人员": "reviewers",
    }
    for label, key in labels.items():
        value = values.get(key, "")
        if key == "start" and value:
            day = parse_date(value)
            value = f"{day.year:04d}.{day.month:02d}.{day.day:02d}"
        _cover_line(doc, _cover_paragraph(doc, label, boundary), value)
    paragraph = _cover_paragraph(doc, "试验日期", boundary)
    raw = d.Paragraph(paragraph, doc).text
    label = re.match(r"^\s*试\s*验\s*日\s*期\s*", raw)
    dates = (_chinese_date(values["start"]) + "到" + _chinese_date(values["end"])
             if values["start"] and values["end"] else "")
    _cover_line(doc, paragraph, dates, label.end(), aligned_start, font_size=13)
    for key, label in (
        ("environment_temperature", "环境温度"),
        ("environment_humidity", "环境相对湿度"),
    ):
        paragraph = _cover_paragraph(doc, label, boundary)
        raw = d.Paragraph(paragraph, doc).text
        _cover_line(doc, paragraph, values.get(key, "").rstrip("；"), raw.index("从"), aligned_start)


def _number_ranges(numbers):
    return d._number_ranges(numbers, separator="~")


def _fill_locations(table, equipment):
    rows = d._rows(table)
    if len(rows) != 2 or any(len(d._cells(row)) != 2 for row in rows):
        raise ValueError("原始记录检测地点表不是预期的两列模板。")
    prototype = deepcopy(rows[1])
    table.remove(rows[1])
    groups, warnings = {}, []
    for index, item in enumerate(equipment, 1):
        location = item.get("storage_location", "张衡路1500号").strip()
        groups.setdefault(location, []).append(index)
        if not location:
            warnings.append(f"原始记录设备序号{index}未填写存放地点，检测地点留空。")
    if not groups:
        groups[""] = []
        warnings.append("原始记录未选择本次设备，检测地点表留空；设备清单沿用报告内容待核对。")
    for location, numbers in groups.items():
        row = deepcopy(prototype)
        for cell, value in zip(d._cells(row), (_number_ranges(numbers) if numbers else "", location)):
            d._put(cell, value)
        table.append(row)
    return warnings


def _align_location_table(table, reference):
    columns = d._xp(table, "./w:tblGrid/w:gridCol")
    source_width = sum(int(value) for value in d._xp(reference, "./w:tblGrid/w:gridCol/@w:w"))
    if len(columns) != 2 or source_width <= 0 or any(len(d._cells(row)) != 2 for row in d._rows(table)):
        raise ValueError("设备表/检测地点表列宽结构无法用于左右对齐。")
    old_widths = [int(column.get(qn("w:w"))) for column in columns]
    if sum(old_widths) <= 0:
        raise ValueError("检测地点表缺少有效列宽。")
    first = round(source_width * old_widths[0] / sum(old_widths))
    widths = [first, source_width - first]
    for column, width in zip(columns, widths):
        column.set(qn("w:w"), str(width))
    for row in d._rows(table):
        for cell, width in zip(d._cells(row), widths):
            cell.get_or_add_tcPr().get_or_add_tcW().set(qn("w:w"), str(width))
            cell.tcPr.tcW.set(qn("w:type"), "dxa")
    table_width = table.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table.tblPr.insert_element_before(table_width, "w:jc", "w:tblInd", "w:tblBorders", "w:tblLayout")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(source_width))
    table.tblPr.get_or_add_tblLayout().set(qn("w:type"), "fixed")
    alignment = d._xp(reference, "./w:tblPr/w:jc/@w:val") or ["left"]
    for row in d._rows(table):
        for setting in d._xp(row, "./w:trPr/w:jc"):
            setting.set(qn("w:val"), alignment[0])
    for name, defaults in (("jc", {"val": "left"}), ("tblInd", {"w": "0", "type": "dxa"}),
                           ("tblCellMar", {})):
        old = table.tblPr.find(qn("w:" + name))
        if old is not None:
            table.tblPr.remove(old)
        source = reference.tblPr.find(qn("w:" + name))
        element = deepcopy(source) if source is not None else OxmlElement("w:" + name)
        if source is None:
            for key, value in defaults.items():
                element.set(qn("w:" + key), value)
        table.tblPr.insert_element_before(element, *{
            "jc": ("w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:tblLayout", "w:tblCellMar", "w:tblLook"),
            "tblInd": ("w:tblBorders", "w:tblLayout", "w:tblCellMar", "w:tblLook"),
            "tblCellMar": ("w:tblLook",),
        }[name])


def _blank_metadata_value(value):
    # Keep inline writing space; removing names/numbers would pull later labels to the left.
    import unicodedata

    return "".join(character if character.isspace() else
                   "\u3000" if unicodedata.east_asian_width(character) in ("W", "F") else
                   "\u2007" if character.isdigit() else "\u2008" for character in value)


def _blank_report_metadata(table):
    labels = r"环境温度|相对湿度|大气压力|评价人员|复核人员"
    units = {"环境温度": r"℃|°C", "相对湿度": r"%RH|%|％", "大气压力": r"kPa|Pa"}
    for row in d._rows(table):
        active_label = ""
        for cell in d._cells(row):
            text = "".join(d._xp(cell, ".//w:t/text()"))
            if "所用试验设备信息见" in text:
                cleared, count = re.subn(r"(序号\s*[：:])(.*)",
                                        lambda match: match[1] + _blank_metadata_value(match[2]), text, flags=re.S)
                if count != 1:
                    raise ValueError("项目页底设备引用未找到唯一序号位置，无法安全清空。")
                d._replace_result_field_text(cell, cleared)
                continue
            matches = list(re.finditer(r"(" + labels + r")\s*[：:]", text))
            if matches:
                active_label = matches[-1][1]
                for index in reversed(range(len(matches))):
                    match = matches[index]
                    end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
                    value = text[match.end():end]
                    unit = re.search(r"(?:" + units[match[1]] + r")\s*$", value) if match[1] in units else None
                    start = len(value) - len(value.lstrip())
                    stop = unit.start() if unit else len(value.rstrip())
                    if stop > start:
                        text = (text[:match.end() + start] + _blank_metadata_value(value[start:stop]) +
                                text[match.end() + stop:])
                d._replace_result_field_text(cell, text)
            elif active_label and text.strip():
                if active_label not in units or not re.fullmatch(units[active_label], text.strip()):
                    unit = re.search(r"(?:" + units[active_label] + r")\s*$", text) if active_label in units else None
                    d._replace_result_field_text(cell, unit[0] if unit else "")
            elif text.strip():
                raise ValueError("项目页底含无法识别的填写栏，未猜测清空：" + text)


def _standardize_gas_metadata(table, *, include_pressure=True):
    rows = d._rows(table)
    labels = (("所用试验设备信息见", "序号"),
              ("环境温度", "相对湿度", "大气压力") if include_pressure else ("环境温度", "相对湿度"),
              ("评价人员", "复核人员"))
    if (len(rows) != 3 or any(len(d._cells(row)) != 1 for row in rows)
            or any(any(d._text(row).count(label) != 1 for label in expected)
                   for row, expected in zip(rows, labels))):
        raise ValueError("燃气表原始记录页底信息组须为设备、环境、人员三行，标签必须唯一。")
    # Match the existing gas form width; explicit stops replace value-sized blank spaces.
    width = 8548
    grid = table.tblGrid
    for column in list(grid):
        grid.remove(column)
    column = OxmlElement("w:gridCol")
    column.set(qn("w:w"), str(width))
    grid.append(column)
    properties = table.tblPr
    table_width = _one(d._xp(properties, "./w:tblW"), "燃气表页底表宽")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(width))
    properties.get_or_add_tblLayout().set(qn("w:type"), "fixed")
    properties.get_or_add_jc().set(qn("w:val"), "left")
    for old in d._xp(properties, "./w:tblInd | ./w:tblBorders | ./w:tblCellSpacing"):
        properties.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    properties.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    content = (
        ("所用试验设备信息见《所用试验设备的信息一览表》序号：", ()),
        (("环境温度：\t℃\t相对湿度：\t%\t大气压力：\tkPa", (27, 33, 60, 66, 94))
         if include_pressure else ("环境温度：\t℃\t相对湿度：\t%", (27, 33, 60))),
        ("评价人员：\t复核人员：", (45,)),
    )
    for index, (row, (text, stops)) in enumerate(zip(rows, content)):
        cell = d._cells(row)[0]
        prototype = _one(d._xp(cell, "./w:p"), "燃气表页底行段落")
        runs = d._xp(prototype, ".//w:r[w:t]")
        if not runs or runs[0].rPr is None:
            raise ValueError("燃气表原始记录页底行缺少文字格式，无法保留原字号。")
        style = deepcopy(runs[0].rPr)
        for old in d._xp(style, "./w:spacing | ./w:fitText | ./w:w | ./w:position"):
            style.remove(old)
        for child in list(cell):
            cell.remove(child)
        cell.get_or_add_tcPr().get_or_add_tcW().set(qn("w:w"), str(width))
        cell.tcPr.tcW.set(qn("w:type"), "dxa")
        d._fit_label_cell(cell)
        for side in ("top", "bottom"):
            margin = OxmlElement("w:" + side)
            margin.set(qn("w:w"), "0")
            margin.set(qn("w:type"), "dxa")
            cell.tcPr.find(qn("w:tcMar")).append(margin)
        cell.tcPr.get_or_add_vAlign().set(qn("w:val"), "center")
        paragraph = OxmlElement("w:p")
        cell.append(paragraph)
        d._left_metadata_cell(cell)
        ppr = paragraph.get_or_add_pPr()
        spacing = ppr.get_or_add_spacing()
        for key, value in (("before", "0"), ("after", "0"), ("line", "360"), ("lineRule", "auto")):
            spacing.set(qn("w:" + key), value)
        for name in ("autoSpaceDE", "autoSpaceDN", "snapToGrid"):
            setting = OxmlElement("w:" + name)
            setting.set(qn("w:val"), "false")
            ppr.insert_element_before(setting, "w:spacing", "w:ind", "w:jc")
        d._paragraph_setting(paragraph, "pageBreakBefore", False)
        d._paragraph_setting(paragraph, "keepNext", index < 2)
        for percent in stops:
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "left")
            tab.set(qn("w:pos"), str(round(width * percent / 100)))
            ppr.get_or_add_tabs().append(tab)
        run = d.Paragraph(paragraph, None).add_run(text)
        run._r.insert(0, style)
        for old in d._xp(row, "./w:tblPrEx"):
            row.remove(old)
        row_properties = row.get_or_add_trPr()
        row_properties.get_or_add_trHeight().set(qn("w:val"), str(d.PROJECT_METADATA_ROW_HEIGHT_TWIPS))
        row_properties.trHeight.set(qn("w:hRule"), "exact")


def _blank_result_fragment(report, template, blocks, *, report_type="", include_pressure=False, root=None,
                           reuse_layout=False, signature_width_twips=None):
    fragment = d._source_fragment(report, blocks)
    if reuse_layout:
        from optimized_reuse import unpack_metadata, release_result_breaks

        unpack_metadata(fragment)
        release_result_breaks(fragment)
    body = fragment.element.body
    for row in d._result_schema(fragment, allow_empty_samples=True):
        d._put(row["_value_cell"], "")
    source_body = template.element.body
    blank_start = d._unique_anchor(source_body, "试验开始日期")
    blank_end = d._unique_anchor(source_body, "试验结束日期")
    blank_judge = d._unique_anchor(source_body, "本试验项目合格判定要求")
    blank_abnormal = d._unique_anchor(source_body, "试验过程中的异常情况记录")
    headings = 0
    for paragraph in list(d._xp(body, "./w:p")):
        text = d._text(paragraph)
        if re.match(r"^型式评价结果\s*[-—－–]", text):
            replacement = re.sub(r"^型式评价结果\s*[-—－–]\s*", "型式评价原始记录-", text)
            d._replace_result_field_text(paragraph, replacement)
            headings += 1
        elif "试验的开始时间" in text or "试验开始日期" in text:
            d._replace_result_field_text(paragraph, d.Paragraph(blank_start, template).text)
        elif "试验的结束时间" in text or "试验结束日期" in text:
            d._replace_result_field_text(paragraph, d.Paragraph(blank_end, template).text)
        elif "本试验项目合格判定要求" in text:
            d._replace_result_field_text(paragraph, d.Paragraph(blank_judge, template).text)
        elif "本试验项目的结论" in text:
            body.remove(paragraph)
        elif "试验过程中的异常情况记录" in text:
            d._replace_result_field_text(paragraph, d.Paragraph(blank_abnormal, template).text)
        elif "所用试验设备信息见" in text:
            body.remove(paragraph)
    if not headings:
        raise ValueError("报告结果区域未找到可替换的型式评价结果标题。")
    tables = [t for t in d._tables(fragment)
              if "环境温度" in d._text(t) and "评价人员" in d._text(t) and "复核人员" in d._text(t)]
    if not tables:
        raise ValueError("报告结果区域缺少项目页底信息表。")
    from signature_templates import apply_signature_templates, is_signature_table

    for table in tables:
        _blank_report_metadata(table)
        if root is None and not is_signature_table(table) and report_type == "燃气表":
            _standardize_gas_metadata(table, include_pressure=include_pressure)
    if root is not None:
        apply_signature_templates(fragment, root, include_pressure, blank=True,
                                  width_twips=signature_width_twips)
    elif include_pressure and report_type != "燃气表" and not any(is_signature_table(t) for t in tables):
        from atmospheric_pressure import set_project_pressure

        set_project_pressure(fragment, True, blank=True)
    if root is None and report_type in ("水表", "流量计", "流量计（气体）", "流量计（液体）") and not any(is_signature_table(t) for t in tables):
        d._format_liquid_metadata(fragment)
    first = next(p for p in d._xp(body, "./w:p") if d._text(p))
    d._page_start(first)
    if not reuse_layout:
        d._split_repeated_project_headers(fragment)
    return fragment


def _field(paragraph, code, nested=None):
    for kind in ("begin", "instruction", "separate", "end"):
        run = paragraph.add_run()
        run.font.size = Pt(10.5)
        if kind == "instruction":
            node = OxmlElement("w:instrText")
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            node.text = code
            run._r.append(node)
            if nested:
                _field(paragraph, nested)
                tail = paragraph.add_run()
                tail.font.size = Pt(10.5)
                instruction = OxmlElement("w:instrText")
                instruction.text = " - 1 "
                instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                tail._r.append(instruction)
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                node.set(qn("w:dirty"), "true")
            run._r.append(node)


def _page_numbers(doc):
    for index, section in enumerate(doc.sections):
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            footer.is_linked_to_previous = False
            for child in list(footer._element):
                footer._element.remove(child)
            paragraph = footer.add_paragraph()
            paragraph.alignment = 1
            if index == 0:
                continue
            paragraph.add_run("第 ").font.size = Pt(10.5)
            _field(paragraph, " PAGE ")
            paragraph.add_run(" 页/共 ").font.size = Pt(10.5)
            _field(paragraph, " = ", " NUMPAGES ")
            paragraph.add_run(" 页").font.size = Pt(10.5)
        numbering = section._sectPr.find(qn("w:pgNumType"))
        if numbering is None:
            numbering = OxmlElement("w:pgNumType")
            section._sectPr.insert_element_before(numbering, "w:cols", "w:formProt", "w:vAlign", "w:docGrid")
        numbering.attrib.pop(qn("w:start"), None)
        if index == 1:
            numbering.set(qn("w:start"), "1")
    update = doc.settings.element.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        doc.settings.element.append(update)
    update.set(qn("w:val"), "true")


def _section_geometry(source, header_source):
    section = deepcopy(source)
    for ref in d._xp(section, "./w:headerReference | ./w:footerReference"):
        section.remove(ref)
    for ref in d._xp(header_source, "./w:headerReference"):
        section.insert(0, deepcopy(ref))
    return section


def _observe_header(composer, observe, section, template):
    header = observe.sections[-1].header
    text_paragraph = _one([p for p in header.paragraphs if p.text.strip()], "观察素材页眉文字段落")
    d._put(text_paragraph._p, d._text(template.sections[-1].header._element))
    part = HeaderPart.new(composer.doc.part.package)
    part._element = deepcopy(header._element)
    composer.add_styles(observe, part.element)
    for ref in d._xp(section, "./w:headerReference"):
        section.remove(ref)
    ref = OxmlElement("w:headerReference")
    ref.set(qn("w:type"), "default")
    ref.set(qn("r:id"), composer.doc.part.relate_to(part, RT.HEADER))
    section.insert(0, ref)


def _end_section(body, section):
    paragraph = OxmlElement("w:p")
    props = paragraph.get_or_add_pPr()
    props.append(deepcopy(section))
    spacing = props.get_or_add_spacing()
    for key, value in (("before", "0"), ("after", "0"), ("line", "1"), ("lineRule", "exact")):
        spacing.set(qn("w:" + key), value)
    body.insert(len(body) - 1, paragraph)


def _fit_uploaded_flow_ranges(source):
    number = r"[+\-−‑]?\d+(?:\.\d+)?(?:[eE][+\-]?\d+)?"
    unit = r"[（(]?(?:m[3³]|[lL])/(?:h|min|s)[）)]?"
    value = rf"[（(]?{number}(?:[~～至–—\-]{number})?[）)]?(?:{unit})?"
    for table in d._tables(source):
        for row in d._rows(table):
            cells = d._cells(row)
            for index, label in enumerate(cells[:-1]):
                if d._norm(d._text(label)).rstrip("：:") != "流量范围":
                    continue
                if not re.fullmatch(value, d._norm(d._text(cells[index + 1]))):
                    continue
                fields = [cells[index + 1]]
                if index + 2 < len(cells) and re.fullmatch(
                        unit, d._norm(d._text(cells[index + 2]))):
                    fields.append(cells[index + 2])
                for cell in fields:
                    # Preserve intentional source line breaks; only reclaim padding in single-line fields.
                    if len(d._xp(cell, "./w:p")) != 1 or d._xp(cell, ".//w:br | .//w:cr | .//w:tab | .//w:tbl"):
                        continue
                    width = d._xp(cell, "./w:tcPr/w:tcW[@w:type='dxa']/@w:w")
                    if len(width) != 1 or int(width[0]) <= 0:
                        raise ValueError("追加原始记录的流量范围单元格缺少有效固定宽度，无法核对单行显示。")
                    if d._sample_header_min_width(source, table, cell) <= int(width[0]):
                        continue
                    adjusted = deepcopy(cell)
                    d._fit_label_cell(adjusted, no_wrap=False)
                    for edge in d._xp(adjusted, "./w:tcPr/w:tcMar/w:start | ./w:tcPr/w:tcMar/w:end"):
                        edge.set(qn("w:w"), "0")
                        edge.set(qn("w:type"), "dxa")
                    if d._sample_header_min_width(source, table, adjusted) > int(width[0]):
                        raise ValueError("追加原始记录的流量范围或单位在原列宽、原字号下仍无法单行显示；"
                                         "未缩字或截断内容，请核对原始记录版式。")
                    cell.replace(cell.tcPr, adjusted.tcPr)


def _append_upload(composer, source, namespace):
    source = d._isolated_styles(source, namespace)
    body = composer.doc.element.body
    _end_section(body, body[-1])
    start = len(body) - 1
    section_map = [(i, deepcopy(d._xp(block, "./w:pPr/w:sectPr")[0]))
                   for i, block in enumerate(source.element.body) if d._xp(block, "./w:pPr/w:sectPr")]
    d._insert(composer, start, source)
    header_cache = {}

    def imported_section(section):
        section = deepcopy(section)
        for ref in d._xp(section, "./w:footerReference"):
            section.remove(ref)
        for ref in d._xp(section, "./w:headerReference"):
            old = ref.get(qn("r:id"))
            if old not in header_cache:
                original = source.part.related_parts[old]
                part = HeaderPart.new(composer.doc.part.package)
                part._element = deepcopy(original.element)
                relationships = {}
                for node in part.element.iter():
                    for attribute, rid in list(node.attrib.items()):
                        if attribute.startswith("{" + nsmap["r"] + "}"):
                            if rid not in relationships:
                                relationships[rid] = composer.add_relationship(original, part, original.rels[rid]).rId
                            node.set(attribute, relationships[rid])
                composer.add_styles(source, part.element)
                header_cache[old] = composer.doc.part.relate_to(part, RT.HEADER)
            ref.set(qn("r:id"), header_cache[old])
        return section

    for offset, section in section_map:
        body[start + offset].get_or_add_pPr().append(imported_section(section))
    final = imported_section(source.element.body[-1])
    body.replace(body[-1], final)
    first = body[start + section_map[0][0]].pPr.sectPr if section_map else final
    first.get_or_add_type().set(qn("w:val"), "nextPage")
    existing = {ref.get(qn("w:type")) for ref in d._xp(first, "./w:headerReference")}
    for kind in ("default", "first", "even"):
        if kind not in existing:
            part = HeaderPart.new(composer.doc.part.package)
            ref = OxmlElement("w:headerReference")
            ref.set(qn("w:type"), kind)
            ref.set(qn("r:id"), composer.doc.part.relate_to(part, RT.HEADER))
            first.insert(0, ref)


def _selected_excel_pages(root, task, record_paths):
    from excel_print_documents import print_sheet_document

    if not any(Path(path).suffix.lower() == ".xlsx" for path in record_paths):
        return [], [], [], []
    projects = d.order_projects(root, task["items_asset_id"], task.get("projects", []))
    pages, details, warnings = [], [], []
    for index, path in enumerate(record_paths):
        if Path(path).suffix.lower() != ".xlsx":
            continue
        key = json.dumps([
            {"id": project["id"], "name": project["name"],
             "rows": [{field: row.get(field) for field in
                       ("sample_id", "value", "matched_sample_id", "requirement")}
                      for row in project["rows"]]} for project in projects
        ], sort_keys=True, ensure_ascii=False)

        def convert():
            source, info = print_sheet_document(path, projects=projects, report_type=task.get("report_type", "水表"))
            content = BytesIO()
            source.save(content)
            return content.getvalue(), info

        content, info = _EXCEL_CACHE.read(Path(path), ("excel-print-v2", task.get("report_type", "水表"), key), convert)
        source = d.Document(BytesIO(content))
        warnings.extend(info.pop("warnings"))
        details.append(info)
        for page in info["included_pages"]:
            blocks = list(source.element.body)[page["body_start"]:page["body_stop"]]
            title = "流量计测量记录单" if task.get("report_type") == "流量计（气体）" else "燃气表测量记录单"
            if not blocks or title not in d._text(blocks[0]):
                raise ValueError("Excel打印页转换内容与项目筛选明细不一致，未发布不完整记录。")
            pages.append((page, index, source, blocks, info["source"]))
    pages.sort(key=lambda item: (item[0]["project_index"], item[0]["sample_index"], item[1], item[0]["page_index"]))
    order = [{**page, "source": filename} for page, _index, _source, _table, filename in pages]
    identities = [(page["project_id"], page["sample_id"]) for page, *_rest in pages]
    if len(set(identities)) != len(identities):
        warnings.append("同项目同编号有多份Excel打印记录，已按关联顺序全部保留，请核对批次；未擅自删除其中一份。")
    return pages, details, order, warnings


def compose_original_record(root, task, outline_path, report_path, output, record_paths=(), *,
                            reuse_layout=False, observation_builder=None, signature_width_twips=None):
    root = d._root(root)
    path = root / TEMPLATE_PATH
    output = Path(output).resolve()
    if output in {path.resolve(), Path(outline_path).resolve(), Path(report_path).resolve(),
                  *(Path(p).resolve() for p in record_paths)}:
        raise ValueError("不能覆盖原始记录模板、大纲或已生成的型评报告。")
    template, report, outline = d._load(path), d._load(report_path), d._load(outline_path)
    doc = deepcopy(template)
    body = doc.element.body
    first_section = _one(d._xp(body, "./w:p/w:pPr/w:sectPr"), "封面分节边界")
    cover_end = first_section.getparent().getparent()
    cover_index = body.index(cover_end)
    values, warnings = cover_information(task, outline)
    # The last cover paragraph carries the section boundary and must also be filled.
    marker = OxmlElement("w:p")
    cover_end.addnext(marker)
    _fill_cover(doc, marker, values)
    body.remove(marker)
    final_section = deepcopy(body.find(qn("w:sectPr")))
    for block in list(body)[cover_index + 1:]:
        body.remove(block)
    body.append(final_section)
    composer = Composer(doc)
    observe = d._load(d._asset(root, task["observe_asset_id"], "observe")["_path"])
    observe_header_source = deepcopy(observe) if reuse_layout else observe
    from observation_documents import (fill_observation_identities, fill_observation_remarks,
                                       prepare_observation_export)
    from observation_certificates import fill_certificates

    if reuse_layout:
        from optimized_reuse import observation_fragment

        observe = observation_fragment(report, observe)
    else:
        if observation_builder is None:
            warnings.extend(fill_observation_remarks(observe, outline, task)["issues"])
            if task.get("report_type", "水表") != "水表":
                warnings.extend(fill_observation_identities(observe, task))
            prepare_observation_export(observe)
        else:
            warnings.extend(observation_builder(observe, outline, task)["warnings"])
        warnings.extend(fill_certificates(observe, outline, task))
    compat = doc.settings.element.find(qn("w:compat"))
    observe_compat = observe.settings.element.find(qn("w:compat"))
    if compat is not None and observe_compat is not None:
        doc.settings.element.replace(compat, deepcopy(observe_compat))
    # Old observe compatibility can prevent Word from opening the nested footer groups.
    d._modern_word_compatibility(doc)
    observe = d._isolated_styles(observe, "recordObserve_", preserve_table_defaults=True)
    heading = d._unique_anchor(observe.element.body, "样机的基本信息")
    for block in list(observe.element.body)[:observe.element.body.index(heading)]:
        if not d._outline_spacing_paragraph(block):
            raise ValueError("观察素材在样机基本信息标题之前含有正文，无法保证该标题为第二页第一行。")
        observe.element.body.remove(block)
    d._paragraph_setting(heading, "keepNext", True)
    spacing = heading.get_or_add_pPr().get_or_add_spacing()
    for name in ("before", "beforeLines", "beforeAutospacing"):
        spacing.set(qn("w:" + name), "0")
    cursor = d._insert(composer, len(body) - 1, observe)
    observe_section = _section_geometry(observe.sections[-1]._sectPr, final_section)
    observe_section.get_or_add_type().set(qn("w:val"), "nextPage")
    _observe_header(composer, observe_header_source if reuse_layout else observe, observe_section, template)
    _end_section(body, observe_section)
    cursor += 1
    report_body = report.element.body
    equipment_heading = d._unique_anchor(report_body, "三、试验项目记录")
    result_heading = d._unique_anchor(report_body, "四、型式评价结果摘要")
    equipment_tables = [b for b in list(report_body)[report_body.index(equipment_heading):report_body.index(result_heading)]
                        if b.tag == qn("w:tbl") and "证书有效期" in d._text(b)]
    equipment_table = _one(equipment_tables, "报告附件设备信息表")
    record_heading = d._unique_anchor(template.element.body, "三、试验项目记录")
    # The equipment-reference line also mentions this caption; locate its exact paragraph.
    record_caption = _one([p for p in d._xp(template.element.body, "./w:p")
                           if d._norm(d._text(p)) == d._norm("所用试验设备的信息一览表")], "设备表标题")
    equipment_fragment = d._source_fragment(template, [record_heading, record_caption])
    d._page_start(equipment_fragment.element.body[0])
    cursor = d._insert(composer, cursor, equipment_fragment)
    cursor = d._insert(composer, cursor, d._source_fragment(report, [equipment_table]))
    location = deepcopy(_one([t for t in d._tables(template)
                             if "所用试验设备序号" in d._text(t) and "检测地点" in d._text(t)], "检测地点表"))
    warnings.extend(_fill_locations(location, task.get("equipment", [])))
    _align_location_table(location, equipment_table)
    spacer = OxmlElement("w:p")
    spacing = spacer.get_or_add_pPr().get_or_add_spacing()
    for name, value in (("before", "0"), ("after", "0"), ("line", "240"), ("lineRule", "exact")):
        spacing.set(qn("w:" + name), value)
    cursor = d._insert(composer, cursor, d._source_fragment(template, [spacer, location]))
    end = d._unique_anchor(report_body, "样机的整体外形照片")
    result_blocks = list(report_body)[report_body.index(result_heading) + 1:report_body.index(end)]
    if task.get("projects"):
        cursor = d._insert(composer, cursor, _blank_result_fragment(
            report, template, result_blocks, report_type=task.get("report_type", ""),
            include_pressure=task.get("include_atmospheric_pressure", False), root=root,
            reuse_layout=reuse_layout, signature_width_twips=signature_width_twips))
        if task.get("report_type") == "燃气表":
            warnings.append("配套原始记录保留本次误差曲线和耐久性曲线及缺图提醒，位于对应数据表下方、合格判定要求上方；其他填写项仍留空，须补充核对。")
    else:
        warnings.append("原始记录未选择试验项目，不复制模板中的历史示例项目。")
    # Keep the original cover/header; subsequent pages use the report result-page geometry.
    result_section = next((sect for b in result_blocks for sect in d._xp(b, ".//w:sectPr")), None)
    if result_section is not None:
        for name in ("pgSz", "pgMar", "cols", "docGrid"):
            source = result_section.find(qn("w:" + name))
            target = final_section.find(qn("w:" + name))
            if source is not None and target is not None:
                final_section.replace(target, deepcopy(source))
    d._editable_page_breaks(doc)
    excel_pages, excel_print_sheets, excel_print_order, excel_warnings = _selected_excel_pages(root, task, record_paths)
    warnings.extend(excel_warnings)
    excel_appended = False
    for index, record_path in enumerate(record_paths):
        if Path(record_path).suffix.lower() == ".xlsx":
            if not excel_appended:
                for _page, source_index, source, blocks, _filename in excel_pages:
                    fragment = d._source_fragment(source, blocks)
                    fragment.element.body.replace(fragment.element.body[-1], deepcopy(source.element.body[-1]))
                    _append_upload(composer, fragment, f"recordExcel{source_index}_")
                excel_appended = True
        else:
            source = d._load(record_path)
            _fit_uploaded_flow_ranges(source)
            _append_upload(composer, source, f"recordUpload{index}_")
    if not record_paths and not task.get("record_ids"):
        warnings.append("本次未关联上传原始记录，合成文件末尾没有附加实测记录。")
    _page_numbers(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    d._normalize_report_fonts(output)
    d._load(output)
    return {"warnings": warnings + [
        "配套原始记录中项目表单的实测结果、日期、人员等填写项已清空，须补充核对；末尾附加的上传实测记录保留原内容，不清空。",
        "原始记录封面试验结论留空，不代表系统核实通过或完成签字。",
        "DOCX原始记录保持关联顺序；Excel打印记录按报告项目及编号顺序追加，共用页只放在最靠前且有该台实测结果的项目下，未选项目或该编号无本次实测结果的页不追加。生成结果不再单独输出Excel，上传原件保留。封面不编号，第二页起为第1页，总页数不含封面。",
    ], "cover": values, "excel_print_sheets": excel_print_sheets, "excel_print_order": excel_print_order}
