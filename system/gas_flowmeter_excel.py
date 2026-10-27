"""Gas-flowmeter XLSX records in the JDJ/E31T002.2 print-sheet layout."""

from copy import deepcopy
from decimal import Decimal
import re
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from openpyxl.utils.exceptions import InvalidFileException

from gas_records import Workbook, check_xlsx, compact, format_result, parse_measurement_period


PARSER = "gas-flowmeter-excel-v3"
PREFIX = "气体流量计Excel："
PRINT_PROJECTS = {6: ("最大允许误差", "重复性", "耐压强度"), 57: ("计量性能复测",)}
PRINT_TITLE = "流量计测量记录单"
RETEST_NAME = "耐压强度、贮存环境性能、电磁兼容试验后计量性能复测"
PRESSURE = re.compile(r"(最大工作压力\s*[:：]\s*)([+-]?\d+(?:\.\d+)?)(\s*MPa)", re.I)


def pressure_requirement(requirement, value):
    from gas_records import decimal

    pressure = decimal(value)
    matches = list(PRESSURE.finditer(requirement))
    if pressure <= 0 or len(matches) != 1:
        raise ValueError("技术要求须有唯一的最大工作压力，记录压力须为有效正数。")
    number = format(pressure, "f")
    if "." in number:
        number = number.rstrip("0").rstrip(".")
    match = matches[0]
    return requirement[:match.start(2)] + number + requirement[match.end(2):]


def fill_pressure_requirements(doc, project):
    """Change only source-backed pressure numbers, splitting shared cells if necessary."""
    import documents as d
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    submitted = {row["id"]: row for row in project.get("rows", [])}
    if not any(row.get("working_pressure") for row in submitted.values()):
        return
    if project["name"] != "耐压强度":
        raise ValueError("最大工作压力只能应用于气体流量计耐压强度技术要求。")
    schema = d._result_schema(doc)
    tables = d._tables(doc)
    effective = {index: list(d._effective_rows(table)) for index, table in enumerate(tables)}
    groups = {}
    for item in schema:
        row = submitted[item["id"]]
        candidate = row.get("working_pressure")
        if candidate:
            expected = pressure_requirement(row["requirement"], candidate["value"])
            if (not candidate["source"] or expected != row.get("resolved_requirement")
                    or item["requirement"] not in (row["requirement"], expected)):
                raise ValueError("最大工作压力来源或技术要求与本次记录不一致，请重新识别。")
        ti, ri, _ci = [int(value) for value in re.fullmatch(r"t(\d+):r(\d+):c(\d+)", item["id"]).groups()]
        physical, owners = effective[ti][ri]
        cells = [(physical[index], owners[index]) for index in range(2, len(physical) - 2)
                 if PRESSURE.search(d._text(owners[index]))]
        if candidate and len(cells) != 1:
            raise ValueError("耐压技术要求中的最大工作压力单元格不唯一，不能自动改写。")
        if len(cells) == 1:
            cell, owner = cells[0]
            groups.setdefault(owner, []).append((cell, candidate))
    for owner, entries in groups.items():
        if not any(candidate for _cell, candidate in entries):
            continue
        protected = (".//w:bookmarkStart | .//w:bookmarkEnd | .//w:drawing | .//w:pict | "
                     ".//w:fldChar | .//w:instrText | .//w:hyperlink | .//w:object | .//w:sectPr")
        if any(d._xp(cell, protected) or (cell is not owner and d._text(cell).strip())
               for cell, _candidate in entries):
            raise ValueError("耐压技术要求包含非普通文字结构，不能拆分共享单元格。")
        content = [deepcopy(child) for child in owner if child.tag != qn("w:tcPr")]
        source = "".join(d._xp(owner, ".//w:t/text()"))
        desired = [pressure_requirement(source, candidate["value"]) if candidate else source
                   for _cell, candidate in entries]
        for index, (cell, _candidate) in enumerate(entries):
            first = index == 0 or desired[index] != desired[index - 1]
            more = index + 1 < len(entries) and desired[index] == desired[index + 1]
            for child in list(cell):
                if child.tag != qn("w:tcPr"):
                    cell.remove(child)
            for merge in d._xp(cell, "./w:tcPr/w:vMerge"):
                merge.getparent().remove(merge)
            if first:
                for child in content:
                    cell.append(deepcopy(child))
                d._replace_result_field_text(cell, desired[index])
            else:
                cell.append(OxmlElement("w:p"))
            if more or not first:
                cell.get_or_add_tcPr().get_or_add_vMerge().set(qn("w:val"), "restart" if first else "continue")


def project_kind(project):
    name = re.sub(r"^\d+[.．、]\s*", "", project["name"].strip())
    return "计量性能复测" if name == RETEST_NAME else name


def _source_number(book, address, *, positive=False):
    try:
        return book.number(address, positive=positive)
    except (KeyError, IndexError) as exc:
        raise ValueError(book.source(address) + "引用不存在的工作表或单元格，请重算保存后上传。") from exc


def _verify(book, address, expected):
    value = _source_number(book, address)
    if abs(value - expected) > Decimal("0.00000001"):
        raise ValueError(book.source(address) + "与测量数据计算不一致，请在Excel重算保存后上传。")
    return value


def _context(book, header):
    from atmospheric_pressure import extract_pressure

    if PRINT_TITLE not in compact(book.text("B%d" % (header - 4))):
        raise ValueError(book.source("B%d" % (header - 4)) + "不是支持的气体流量计测量记录打印页。")
    offset = header - 6
    for address, label in (("B4", "器具名称"), ("B5", "型号/规格"), ("B6", "编号"),
                           ("I4", "证书编号"), ("B7", "大气压力"), ("G7", "流量范围")):
        column, row = re.fullmatch(r"([A-Z]+)(\d+)", address).groups()
        book.anchor(column + str(int(row) + offset), label)
    sample, model = book.text("C%d" % header), book.text("C%d" % (header - 1))
    instrument = book.text("C%d" % (header - 2))
    if not sample or not model or "流量计" not in compact(instrument):
        raise ValueError(book.source("C%d" % header) + "缺少明确的流量计名称、型号或出厂编号。")
    raw_sheet = "初始" if header == 6 else "贮存环境电磁兼容后复测"
    if raw_sheet not in book.values.sheetnames:
        raise ValueError("气体流量计Excel缺少原始测量工作表：" + raw_sheet)
    if (book.text("H3", raw_sheet) != sample
            or compact(book.text("H2", raw_sheet)) != compact(model)):
        raise ValueError(book.source("H3", raw_sheet) + "与打印页样机编号或型号不一致，不能跨样机取值。")
    if "气" not in compact(book.text("B7", raw_sheet)):
        raise ValueError(book.source("B7", raw_sheet) + "未明确为气体测量介质，不能用于流量计（气体）。")
    parameter_row = 14 if header == 6 else 61
    for column, label in (("C", "qmax"), ("G", "qmin"), ("K", "qt")):
        book.anchor(column + str(parameter_row), label)
    parameters = {}
    provenance = {}
    for key, column in (("qmax", "E"), ("qmin", "I"), ("qt", "L")):
        address = column + str(parameter_row)
        parameters[key] = str(_source_number(book, address, positive=True))
        provenance[key] = book.source(address)
    if not Decimal(parameters["qmin"]) < Decimal(parameters["qt"]) < Decimal(parameters["qmax"]):
        raise ValueError(book.source("E%d" % parameter_row) + "流量参数须满足qmin < qt < qmax。")
    for column in ("F", "J", "M"):
        if not re.fullmatch(r"m(?:³|3|\^3)/h", compact(book.text(column + str(parameter_row))), re.I):
            raise ValueError(book.source(column + str(parameter_row)) + "流量参数缺少明确的m³/h单位。")
    flow_range = re.fullmatch(r"\(?(\d+(?:\.\d+)?)[～~至-](\d+(?:\.\d+)?)\)?m(?:³|3|\^3)/h",
                              compact(book.text("I%d" % (header + 1))), re.I)
    if (not flow_range or Decimal(flow_range[1]) != Decimal(parameters["qmin"])
            or Decimal(flow_range[2]) != Decimal(parameters["qmax"])):
        raise ValueError(book.source("I%d" % (header + 1)) + "流量范围与qmin/qmax参数不一致。")
    date_row = 47 if header == 6 else 90
    book.anchor("J%d" % date_row, "检测日期：")
    warnings = []
    try:
        start, end = parse_measurement_period(book.text("K%d" % date_row))
    except ValueError as exc:
        start = end = ""
        warnings.append(PREFIX + book.source("K%d" % date_row) + "检测日期未采用：" + str(exc))
    provenance.update({
        "sample_id": book.source("C%d" % header), "model": book.source("C%d" % (header - 1)),
        "sample_name": book.source("C%d" % (header - 2)),
        "certificate": book.source("K%d" % (header - 2)),
        "manufacturer": book.source("K%d" % (header - 1)), "time": book.source("K%d" % date_row),
    })
    pressure_address = "C%d" % (header + 1)
    pressure_value = book.text(pressure_address)
    return {
        "sample_id": sample, "model": model, "sample_name": instrument,
        "manufacturer": book.text("K%d" % (header - 1)),
        "certificate": book.text("K%d" % (header - 2)), "parser": PARSER,
        "source": book.source("C%d" % header), "start": start, "end": end,
        "start_source": provenance["time"] if start else "", "end_source": provenance["time"] if end else "",
        "temperature": book.text("K%d" % header).replace("℃", "").replace("°C", "").strip(),
        "temperature_unit": "℃",
        "humidity": book.text("M%d" % header).replace("%RH", "").replace("%", "").strip(),
        "humidity_unit": "%RH", "provenance": provenance, **parameters,
        "parameters": {key: [{"value": value, "unit": "m3/h", "source": provenance[key]}]
                       for key, value in parameters.items()},
        "atmospheric_pressure_candidates": extract_pressure(
            "大气压力：" + pressure_value, book.source(pressure_address)) if pressure_value else [],
        "issues": warnings,
    }


def _measurements(book, context, header):
    initial = header == 6
    top, unit_row = (15, 18) if initial else (62, 65)
    for column, label in (("B", "检测流量点"), ("C", "序号"), ("E", "流量计读数"),
                          ("G", "标准器读数"), ("I", "示值误差"), ("J", "平均示值误差")):
        book.anchor(column + str(top), label)
    for column in ("E", "G"):
        book.anchor(column + str(unit_row), "（L）")
    for column in ("I", "J") + (("K",) if initial else ()):
        book.anchor(column + str(unit_row), "（%）")
    if initial:
        book.anchor("K15", "重复性")
    blocks = [(19, 6, "qmax"), (25, 6, "0.4qmax"), (31, 6, "qt"), (37, 6, "qmin")] if initial else [
        (66, 3, "qmin")]
    metrics, raw_rows, issues = [], [], []
    for row, count, flow_label in blocks:
        try:
            book.anchor("B%d" % row, flow_label)
            flow = (Decimal(context["qmax"]) * Decimal("0.4") if flow_label == "0.4qmax"
                    else Decimal(context[flow_label]))
            errors, raw = [], []
            for index in range(row, row + count):
                ordinal = _source_number(book, "C%d" % index)
                if ordinal != index - row + 1:
                    raise ValueError(book.source("C%d" % index) + "测量次数不连续，不能混合不同批次。")
                meter = _source_number(book, "E%d" % index, positive=True)
                standard = _source_number(book, "G%d" % index, positive=True)
                errors.append(_verify(book, "I%d" % index, (meter - standard) / standard * 100))
                raw.append({"source": book.source("E%d:G%d" % (index, index)),
                            "cells": [{"raw": book.text(column + str(index)),
                                       "source": book.source(column + str(index))}
                                      for column in ("C", "E", "G", "I")]})
            mean = _verify(book, "J%d" % row, sum(errors) / len(errors))
            values = [("平均示值误差", "J", mean)]
            if initial:
                deviation = (sum((error - mean) ** 2 for error in errors) / (len(errors) - 1)).sqrt()
                values.append(("重复性", "K", _verify(book, "K%d" % row, deviation)))
            for metric, column, value in values:
                address = column + str(row)
                metrics.append({
                    "metric": metric, "label": flow_label + " / " + metric,
                    "value": format_result(value), "raw_value": str(value), "unit": "%",
                    "flow": str(flow), "flow_unit": "m3/h", "flow_label": flow_label,
                    "flow_source": book.source("B%d" % row), "source": book.source(address),
                    "formula": book.formulas["打印页"][address].value
                    if book.formulas["打印页"][address].data_type == "f" else "",
                })
            raw_rows.extend(raw)
        except ValueError as exc:
            issues.append(PREFIX + flow_label + "未采用：" + str(exc))
    return metrics, raw_rows, issues


def parse_upload(path):
    check_xlsx(path)
    try:
        book = Workbook(path)
    except (KeyError, BadZipFile, XMLSyntaxError, InvalidFileException) as exc:
        raise ValueError("气体流量计Excel内部结构无法读取，请另存为有效XLSX后重新上传。") from exc
    try:
        if "打印页" not in book.values.sheetnames or PRINT_TITLE not in compact(book.text("B2")):
            raise ValueError("不支持的气体流量计Excel布局，须有JDJ/E31T002.2流量计测量记录打印页。")
        result = {"parser": PARSER, "report_type": "流量计（气体）", "fields": {}, "records": [],
                  "warnings": [], "outline": {}}
        for header, projects in PRINT_PROJECTS.items():
            if header != 6 and not book.text("B%d" % (header - 4)):
                continue
            context = _context(book, header)
            metrics, raw, issues = _measurements(book, context, header)
            result["warnings"].extend(context["issues"] + issues)
            for project in projects:
                selected = [metric for metric in metrics if metric["metric"] ==
                            ("重复性" if project == "重复性" else "平均示值误差")]
                if project == "耐压强度":
                    book.anchor("B12", "耐压强度：")
                    book.anchor("F12", "（最大工作压力")
                    value = book.text("C12")
                    if not value:
                        result["warnings"].append(PREFIX + book.source("C12") + "耐压检查为空，未采用。")
                        continue
                    book.anchor("J12", "MPa）")
                    try:
                        pressure = _source_number(book, "I12", positive=True)
                    except ValueError as exc:
                        result["warnings"].append(PREFIX + "耐压检查未采用：" + str(exc))
                        continue
                    selected = [{"metric": "耐压强度", "label": "耐压强度", "value": value, "raw": value,
                                 "source": book.source("C12") + "；最大工作压力：" + book.source("I12"),
                                 "maximum_working_pressure": str(pressure), "unit": ""}]
                if selected:
                    result["records"].append(dict(deepcopy(context), project=project,
                                                  rows=deepcopy(selected), raw_rows=deepcopy(raw)))
            for key, field in (("certificate", "report_number"), ("sample_name", "instrument_name"),
                               ("manufacturer", "manufacturer"), ("model", "model")):
                if context[key]:
                    candidate = {"value": context[key], "source": context["provenance"][key]}
                    if candidate not in result["fields"].setdefault(field, []):
                        result["fields"][field].append(candidate)
        if not result["records"]:
            raise ValueError("未找到可采用的气体流量计Excel实测记录：" + "；".join(result["warnings"]))
        result["warnings"].extend([
            "气体流量计Excel按打印页已保存的平均示值误差、重复性及单元格来源识别；结果保留两位小数、负零归零，原始精度保留。",
            "贮存、振动和电磁兼容试验本身未提供独立记录，不由复测数据推断符合要求；耐压最大工作压力按对应样机的唯一原始记录更新。",
        ])
        result["warnings"] = list(dict.fromkeys(result["warnings"]))
        return result
    finally:
        book.values.close()
        book.formulas.close()
