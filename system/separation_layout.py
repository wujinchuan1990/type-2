"""Content-preserving report layout for editable split materials."""

from collections import Counter
from copy import deepcopy
import hashlib
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf

import documents as d
from separation_documents import effective_sections


POLICY = "report-layout-v1-results-text-merge-20261024"
TEMPLATE = Path(__file__).resolve().parents[1] / "report model/report.docx"


def normalized_text(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))


def content_signature(doc):
    """Ignore layout, but retain text order, cell content/merges and image bytes."""
    body = doc.element.body
    texts = tuple(normalized_text(d._text(p)) for p in d._xp(body, ".//w:p")
                  if normalized_text(d._text(p)))
    cells = []
    for cell in d._xp(body, ".//w:tc"):
        # Layout may wrap existing metadata paragraphs in a new borderless cell.
        table = cell.getparent().getparent()
        if "环境温度" in d._text(table) and "复核人员" in d._text(table):
            continue
        cells.append((normalized_text(d._text(cell)),
                      tuple(d._xp(cell, "./w:tcPr/w:gridSpan/@w:val")),
                      tuple(d._xp(cell, "./w:tcPr/w:vMerge/@w:val")),
                      bool(d._xp(cell, "./w:tcPr/w:vMerge"))))
    pictures = []
    for node in body.iter():
        if node.tag in (qn("w:headerReference"), qn("w:footerReference")):
            continue
        for key in (qn("r:embed"), qn("r:link"), qn("r:id")):
            identity = node.get(key)
            if not identity:
                continue
            relation = doc.part.rels[identity]
            pictures.append((relation.reltype, relation.target_ref if relation.is_external
                             else hashlib.sha256(relation.target_part.blob).hexdigest()))
    headers = tuple(normalized_text(d._text(relation.target_part.element))
                    for relation in doc.part.rels.values() if relation.reltype.endswith("/header"))
    return {"paragraphs": texts, "cells": tuple(cells), "relationships": tuple(pictures), "headers": headers}


def template_properties(kind):
    doc = Document(TEMPLATE)
    phrase = "型式评价项目及评价结果一览表" if kind == "items" else "样机的基本信息"
    block = d._unique_anchor(doc.element.body, phrase)
    index = list(doc.element.body).index(block)
    return next(props for start, end, props in effective_sections(doc.element.body) if start <= index <= end)


def zero_heading_spacing(paragraph):
    spacing = paragraph.get_or_add_pPr().get_or_add_spacing()
    for key in ("before", "beforeLines", "beforeAutospacing"):
        spacing.set(qn("w:" + key), "0")


def first_heading(body):
    for block in list(body):
        if block.tag == qn("w:p") and not d._text(block) and d._outline_spacing_paragraph(block):
            body.remove(block)
        elif block.tag != qn("w:sectPr"):
            if block.tag != qn("w:p") or not d._text(block):
                raise ValueError("拆分部分缺少可识别的首段标题，无法应用报告版式。")
            d._page_start(block, enabled=False)
            zero_heading_spacing(block)
            return block
    raise ValueError("拆分部分没有正文。")


def paragraph_metadata(doc):
    """Support reports whose final equipment/environment/staff are paragraphs."""
    body = doc.element.body
    candidates = [p for p in d._xp(body, "./w:p") if "所用试验设备信息见" in d._text(p)]
    if not candidates:
        return
    equipment = candidates[-1]
    following = [p for p in list(body)[list(body).index(equipment) + 1:]
                 if p.tag != qn("w:sectPr") and (d._text(p) or d._xp(p, ".//w:drawing | .//w:pict"))]
    if not following or following[0].tag != qn("w:p"):
        return
    if (len(following) != 2 or not all(p.tag == qn("w:p") for p in following)
            or "环境温度" not in d._text(following[0])
            or not all(label in d._text(following[1]) for label in ("评价人员", "复核人员"))):
        raise ValueError("项目末尾环境/人员段落结构不明确，不能移动或删除原文。")
    section = doc.sections[-1]
    width = (section.page_width - section.left_margin - section.right_margin) // 635
    table = OxmlElement("w:tbl")
    props, grid = OxmlElement("w:tblPr"), OxmlElement("w:tblGrid")
    col = OxmlElement("w:gridCol")
    col.set(qn("w:w"), str(width))
    grid.append(col)
    table.extend([props, grid])
    following[0].addprevious(table)
    for paragraph in following:
        row, cell = OxmlElement("w:tr"), OxmlElement("w:tc")
        cell.append(paragraph)
        row.append(cell)
        table.append(row)


def page_numbers(doc):
    seen = set()
    for section in doc.sections:
        title = section._sectPr.find(qn("w:titlePg"))
        if title is not None:
            section._sectPr.remove(title)
        footers = [section.footer]
        if doc.settings.odd_and_even_pages_header_footer:
            footers.append(section.even_page_footer)
        for footer in footers:
            if footer.part.partname in seen:
                continue
            seen.add(footer.part.partname)
            candidates = [p for p in footer.paragraphs if re.search(
                r"\b(?:PAGE|NUMPAGES)\b", " ".join(d._xp(p._p, ".//w:instrText/text() | .//w:fldSimple/@w:instr")), re.I)
                or re.fullmatch(r"\s*第\s*\d+\s*页\s*[/／]?\s*共\s*\d+\s*页\s*", p.text)]
            if len(candidates) > 1:
                raise ValueError("页脚存在多组页码，不能自动替换不明确的页码区域。")
            paragraph = candidates[0] if candidates else footer.add_paragraph()
            if re.sub(r"[\s\d第页共/／]+", "", paragraph.text):
                raise ValueError("页码段落混有其他页脚文字，不能覆盖原内容。")
            paragraph.clear()
            paragraph.alignment = 1
            for label, code in (("第 ", "PAGE"), (" 页 共 ", "NUMPAGES")):
                paragraph.add_run(label)
                for kind in ("begin", "instruction", "separate", "result", "end"):
                    run = paragraph.add_run()._r
                    if kind == "instruction":
                        node = OxmlElement("w:instrText")
                        node.text = code
                    elif kind == "result":
                        node = OxmlElement("w:t")
                        node.text = "1"
                    else:
                        node = OxmlElement("w:fldChar")
                        node.set(qn("w:fldCharType"), kind)
                    run.append(node)
            paragraph.add_run(" 页")


def normalize(path, kind, project_ranges=(), *, result_format=False, split_result_tables=True):
    doc = Document(path)
    before = content_signature(doc)
    template = template_properties(kind)
    for section in doc.sections:
        for name in ("pgSz", "pgMar", "docGrid"):
            original = section._sectPr.find(qn("w:" + name))
            replacement = template.find(qn("w:" + name))
            if original is not None and replacement is not None:
                section._sectPr.replace(original, deepcopy(replacement))
        page = section._sectPr.find(qn("w:pgNumType"))
        if page is not None:
            page.attrib.pop(qn("w:start"), None)
    if kind == "result":
        # Each semantic project is laid out separately; no metadata is lost
        # when a user combines several complete projects in one output.
        blocks = list(doc.element.body)
        for start, end in reversed(project_ranges):
            selected = blocks[start:end]
            fragment = d._source_fragment(doc, selected)
            fragment.element.body[-1].getparent().replace(fragment.element.body[-1], deepcopy(doc.sections[-1]._sectPr))
            first_heading(fragment.element.body)
            paragraph_metadata(fragment)
            metadata = [t for t in d._tables(fragment) if all(
                label in d._text(t) for label in ("环境温度", "评价人员", "复核人员"))]
            if len(metadata) > 1:
                raise ValueError("同一项目含多组环境/人员信息，统一版式不能自动删除原内容。")
            if metadata:
                d._unpack_metadata_equipment(metadata[0])
            d._flow_project_layout(fragment)
            # A later row rebuild must own and clean up its generated separators.
            if split_result_tables:
                d._split_repeated_project_headers(fragment)
            replacement = [b for b in fragment.element.body if b.tag != qn("w:sectPr")]
            d._page_start(replacement[0], enabled=start > 0)
            position = list(doc.element.body).index(selected[0])
            for b in selected:
                doc.element.body.remove(b)
            for offset, b in enumerate(replacement):
                doc.element.body.insert(position + offset, b)
    first_heading(doc.element.body)
    if content_signature(doc) != before:
        raise ValueError("统一排版改变了正文、单元格内容/合并或图片，停止交付。")
    if kind == "result" and result_format:
        from separation_result_format import format_results

        format_results(doc)
    page_numbers(doc)
    doc.save(path)
    d._normalize_report_fonts(path)
    validate_layout(path, kind)
    return content_signature(Document(path))


def validate_layout(path, kind):
    doc = Document(path)
    template = template_properties(kind)
    for section in doc.sections:
        for name in ("pgSz", "pgMar"):
            actual, expected = section._sectPr.find(qn("w:" + name)), template.find(qn("w:" + name))
            if actual is None or expected is None or dict(actual.attrib) != dict(expected.attrib):
                raise ValueError("拆分文件纸张/页边距不符合报告模板。")
    heading = next((p for p in doc.paragraphs if p.text.strip()), None)
    if heading is None or heading.paragraph_format.space_before is None or heading.paragraph_format.space_before.pt != 0:
        raise ValueError("拆分文件首段段前间距未归零。")
    return doc


def verify_rendered(path, pdf_path, signature, kind, source_start, source_end):
    doc = validate_layout(path, kind)
    if content_signature(doc) != signature:
        raise ValueError("Word重分页后正文、单元格或图片内容发生变化。")
    with pymupdf.open(pdf_path) as pdf:
        if not len(pdf):
            raise ValueError("Word未生成可核对的页面。")
        visible = normalized_text("".join(page.get_text() for page in pdf))
        expected = normalized_text("".join(signature["paragraphs"]))
        missing = Counter(expected) - Counter(visible)
        if missing:
            detail = "、".join(f"{char}×{count}" for char, count in list(missing.items())[:12])
            raise ValueError("Word渲染存在缺字或内容未完整显示：" + detail)
        for index, page in enumerate(pdf):
            if not re.search(rf"第\s*{index + 1}\s*页\s*共\s*{len(pdf)}\s*页", page.get_text()):
                raise ValueError(f"输出P{index + 1}的连续页码/实际总页数未正确更新。")
            if "\ufffd" in page.get_text():
                raise ValueError(f"输出P{index + 1}存在无法显示的字符。")
            for word in page.get_text("words"):
                if word[0] < -1 or word[1] < -1 or word[2] > page.rect.width + 1 or word[3] > page.rect.height + 1:
                    raise ValueError(f"输出P{index + 1}有文字超出纸张边界。")
        return {"layout_policy": POLICY, "source_start_page": source_start, "source_end_page": source_end,
                "output_pages": len(pdf), "content_verified": True, "layout_verified": True,
                "review_note": "内容及自动版式检查通过；统一排版允许页数变化，仍需人工审版。"}
