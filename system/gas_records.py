"""Gas-meter XLSX records in the JDJ/E70T029.1 print-sheet layout."""

from copy import deepcopy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re
from statistics import stdev
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from lxml.etree import XMLSyntaxError

import documents
from curve_results import _manual_result, refresh_curve_results


PARSER = "gas-excel-v4"
PREFIX = "燃气表Excel："
FLOWS = ("qmax", "0.7qmax", "0.4qmax", "0.2qmax", "qt", "10qmin", "5qmin", "qmin")
INITIAL = (17, 23, 29, 35, 41, 47, 49, 51)
DURABLE = (258, 261, 264, 267, 270, 273, 275, 277)
DATE_ROWS = {7: 55, 66: 99, 111: 145, 156: 188, 201: 234, 246: 279, 290: 324}
PRINT_PROJECTS = {
    7: ("准确度等级和最大允许误差", "复现性", "误差曲线", "加权平均误差", "压力损失", "密封性"),
    66: ("重复性", "不同工作模式的误差偏差"),
    111: ("温度适应性",),
    156: ("环境适应性复测",),
    201: ("流体扰动",),
    246: ("耐久性",),
    290: ("机械环境复测",),
}
METRIC_LABELS = {
    "mean": "平均示值误差", "sd": "复现性（标准偏差）", "repeat": "重复性（极差）",
    "span": "误差曲线落差", "weighted": "加权平均误差", "pressure": "压力损失",
    "mode": "工作模式误差偏差", "reference": "参比条件误差", "elbow": "90°单弯头误差",
    "shift": "误差偏移", "reference-before": "初始参比温度", "reference-after": "恢复参比温度",
    "high": "最高温度", "low": "最低温度",
}


def compact(value):
    return re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")")


def decimal(value):
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("不是有效数值：" + str(value)) from exc
    if not result.is_finite():
        raise ValueError("不支持非有限数值：" + str(value))
    return result


def format_result(value, unit="%"):
    places = 0 if unit == "Pa" else 2
    rounded = decimal(value).quantize(Decimal("1") if places == 0 else Decimal("0.01"), rounding=ROUND_HALF_UP)
    if not rounded:
        rounded = abs(rounded)
    return format(rounded, f".{places}f") + unit


def parse_measurement_period(value):
    match = re.fullmatch(
        r"\s*(\d{4})[./-](\d{1,2})[./-](\d{1,2})\s+(\d{1,2}):(\d{2})"
        r"\s*[~～至—–]\s*"
        r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})\s+(\d{1,2}):(\d{2})\s*",
        str(value or ""),
    )
    if not match:
        raise ValueError("须按YYYY.MM.DD HH:MM~YYYY.MM.DD HH:MM填写")
    numbers = [int(item) for item in match.groups()]
    from datetime import datetime
    start = datetime(*numbers[:5])
    end = datetime(*numbers[5:])
    if start > end:
        raise ValueError("开始时间晚于结束时间")
    return start.strftime("%Y-%m-%d %H:%M"), end.strftime("%Y-%m-%d %H:%M")


def check_xlsx(path):
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)):
                raise ValueError("XLSX包中存在重复部件。")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise ValueError("不是有效的Excel XLSX工作簿。")
            if len(names) > 10000 or sum(i.file_size for i in archive.infolist()) > 160 * 1024 * 1024:
                raise ValueError("Excel解压后过大，无法安全处理。")
            if (any("vbaproject" in name.lower() or name.startswith("xl/externalLinks/") for name in names)
                    or b"macroEnabled" in archive.read("[Content_Types].xml")):
                raise ValueError("不支持含宏或外部工作簿链接的Excel，请另存为独立无宏XLSX并核对。")
            for name in names:
                if name.endswith((".xml", ".rels")):
                    data = archive.read(name)
                    if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
                        raise ValueError("不支持带外部实体定义的Excel。")
    except BadZipFile as exc:
        raise ValueError("不是有效XLSX，请勿仅修改扩展名。") from exc


class Workbook:
    def __init__(self, path):
        self.path = Path(path)
        self.values = load_workbook(path, data_only=True, read_only=False, keep_links=False)
        self.formulas = load_workbook(path, data_only=False, read_only=False, keep_links=False)
        self.date_issues = []

    def source(self, address, sheet="打印页"):
        return "%s · Excel/%s!%s" % (self.path.name, sheet, address)

    def text(self, address, sheet="打印页"):
        value = self.values[sheet][address].value
        return str(value).strip() if value is not None else ""

    def number(self, address, sheet="打印页", positive=False):
        cell = self.values[sheet][address]
        if cell.data_type == "e" or cell.value is None:
            raise ValueError(self.source(address, sheet) + "缺少有效结果或公式缓存（%s）" % cell.value)
        try:
            value = decimal(cell.value)
        except ValueError as exc:
            raise ValueError(self.source(address, sheet) + "：" + str(exc)) from exc
        if positive and value <= 0:
            raise ValueError(self.source(address, sheet) + "缺少有效的正数测量读数，不能采用模板零值。")
        formula = self.formulas[sheet][address].value
        if isinstance(formula, str) and formula.startswith("="):
            reference = re.fullmatch(r"='?([^'!]+)'?!([$A-Z]+\$?\d+)", formula)
            if reference:
                raw = self.values[reference[1]][reference[2].replace("$", "")]
                if raw.value is None or raw.data_type == "e":
                    raise ValueError(self.source(address, sheet) + "引用空白或错误原始单元格：" + formula)
                if abs(decimal(raw.value) - value) > Decimal("0.00000001"):
                    raise ValueError(self.source(address, sheet) + "与原始单元格不一致，请在Excel重算保存后上传。")
        return value

    def percent(self, address):
        value = self.number(address)
        return value * 100 if "%" in self.values["打印页"][address].number_format else value

    def verify(self, address, expected, percent=False):
        actual = self.percent(address) if percent else self.number(address)
        if abs(actual - expected) > Decimal("0.00000001"):
            raise ValueError(self.source(address) + "与测量数据计算不一致，请在Excel重算保存后上传。")
        return actual

    def anchor(self, address, expected):
        if compact(self.text(address)) != compact(expected):
            raise ValueError(self.source(address) + "表头或流量点与支持的记录布局不符，未按固定位置猜读。")

    def mean(self, row, count, flow):
        self.anchor("B%d" % row, flow)
        errors = []
        for index in range(row, row + count):
            meter = self.number("E%d" % index, positive=True)
            standard = self.number("G%d" % index, positive=True)
            errors.append(self.verify("I%d" % index, (meter - standard) / standard * 100))
        return self.verify("J%d" % row, sum(errors) / len(errors)), errors

    def context(self, row, sample, model):
        from atmospheric_pressure import extract_pressure

        if self.text("C%d" % row) != sample or compact(self.text("K%d" % (row - 1))) != compact(model):
            raise ValueError(self.source("C%d" % row) + "分项目样机编号或型号不一致，不能跨样机取值。")
        date_row = DATE_ROWS[row]
        if compact(self.text("J%d" % date_row)) == compact("测量日期："):
            date_address = "K%d" % date_row
        elif compact(self.text("K%d" % date_row)) == compact("测量日期："):
            date_address = "L%d" % date_row
        else:
            date_address = ""
        try:
            if not date_address or not self.text(date_address):
                raise ValueError("测量日期为空")
            start, end = parse_measurement_period(self.text(date_address))
        except ValueError as exc:
            start = end = ""
            source = self.source(date_address or ("J%d" % date_row))
            self.date_issues.append(PREFIX + source + "测量日期未采用：" + str(exc))
        return {
            "sample_id": sample, "model": model, "certificate": self.text("K4"),
            "source": self.source("C%d" % row), "start": start, "end": end,
            "start_source": self.source(date_address) if start else "",
            "end_source": self.source(date_address) if end else "",
            "temperature": self.text("K%d" % row).replace("℃", "").replace("°C", "").strip(),
            "temperature_unit": "℃",
            "humidity": self.text("M%d" % row).replace("%RH", "").replace("%", "").strip(),
            "humidity_unit": "%RH", "parser": PARSER,
            "atmospheric_pressure_candidates": extract_pressure(
                self.text("B%d" % (row + 1)) + self.text("C%d" % (row + 1)),
                self.source("C%d" % (row + 1))),
        }

    def metric(self, key, value, addresses, unit="%"):
        locations = [self.source(address) for address in addresses]
        parts = key.split(":")
        labels = [METRIC_LABELS[parts[0]]] + [
            {"high": "高流量区间", "low": "低流量区间"}.get(part, part) for part in parts[1:]]
        if len(parts) == 3:
            labels[-1] += "℃"
        return {"key": key, "label": " / ".join(labels), "value": format_result(value, unit),
                "raw_value": str(value), "source": "；".join(locations), "unit": unit,
                "cells": [{"address": address, "value": self.text(address),
                           "formula": self.formulas["打印页"][address].value
                           if self.formulas["打印页"][address].data_type == "f" else ""}
                          for address in addresses]}


def parse_upload(path):
    check_xlsx(path)
    try:
        book = Workbook(path)
    except (KeyError, BadZipFile, XMLSyntaxError, InvalidFileException) as exc:
        raise ValueError("Excel内部结构无法读取，请另存为有效XLSX后重新上传。") from exc
    try:
        required = {"打印页", "降流量", "升流量", "重复性", "用户模式", "温度适应性",
                    "环适复测", "流体扰动", "耐久性复测", "振动跌落后"}
        if not required.issubset(book.values.sheetnames):
            raise ValueError("不支持的燃气表Excel布局，缺少工作表：" + "、".join(sorted(required - set(book.values.sheetnames))))
        for address, label in (("B3", "燃气表测量记录单"), ("B7", "出厂编号："),
                               ("I6", "型号："), ("J13", "平均示值误差"), ("K11", "加权平均误差(±0.6%)：")):
            book.anchor(address, label)
        sample, model = book.text("C7"), book.text("K6")
        if not sample or not model or "燃气表" not in book.text("K5"):
            raise ValueError("燃气表Excel缺少明确的器具名称、样机编号或型号。")
        qmax, qmin, qt = (book.number(a, positive=True) for a in ("E12", "I12", "L12"))
        if not qmin < qt < qmax:
            raise ValueError("燃气表Excel流量参数须满足qmin < qt < qmax。")
        result = {"parser": PARSER, "report_type": "燃气表", "fields": {}, "records": [],
                  "warnings": [], "outline": {}}
        if book.text("K4"):
            result["fields"]["report_number"] = [{"value": book.text("K4"), "source": book.source("K4")}]
        else:
            result["warnings"].append("未填写记录证书编号，报告编号需补录，不推测报告后缀。")
        initial = {}

        def add(project, header, metrics):
            context = book.context(header, sample, model)
            context.update(project=project, rows=metrics, qmax=str(qmax), qmin=str(qmin), qt=str(qt))
            result["records"].append(context)

        def group(name, action):
            before = len(result["records"])
            try:
                action()
            except ValueError as exc:
                del result["records"][before:]
                if name == "初始误差/复现性/曲线/加权平均误差":
                    initial.clear()
                result["warnings"].append(PREFIX + name + "未采用：" + str(exc))

        def accuracy():
            means, deviations = [], []
            for index, (row, flow) in enumerate(zip(INITIAL, FLOWS)):
                value, errors = book.mean(row, 6 if index < 5 else 2, flow)
                initial[flow] = value
                means.append(book.metric("mean:" + flow, value, ["J%d" % row]))
                if index < 5:
                    sd = book.verify("K%d" % row, decimal(stdev(float(e) for e in errors)), percent=True)
                    deviations.append(book.metric("sd:" + flow, sd, ["K%d" % row]))
            high = [initial[flow] for flow in FLOWS[:5]]
            sign = "/（异号）" if min(high) < 0 < max(high) else format_result(max(abs(v) for v in high))
            means.append({"key": "same-sign", "label": "同号示值误差", "value": sign,
                          "source": "；".join(book.source("J%d" % r) for r in INITIAL[:5]) + "；同号取最大绝对误差，否则异号"})
            add("准确度等级和最大允许误差", 7, means)
            add("复现性", 7, deviations)
            curves = []
            for key, flows, address in (("high", FLOWS[:5], "M17"), ("low", FLOWS[5:], "M47")):
                values = [initial[f] for f in flows]
                value = book.verify(address, max(values) - min(values), percent=True)
                curves.append(book.metric("span:" + key, value, [address]))
            add("误差曲线", 7, curves)
            ratio = qmin / qmax
            weights = [Decimal(str(v)) for v in (0.4, 0.7, 0.4, 0.2, 0.1)] + [10 * ratio, 5 * ratio, ratio]
            weighted = sum(initial[f] * w for f, w in zip(FLOWS, weights)) / sum(weights)
            add("加权平均误差", 7, [book.metric("weighted", book.verify("M11", weighted, percent=True), ["M11"])])

        group("初始误差/复现性/曲线/加权平均误差", accuracy)

        def pressure(address, raw_sheet):
            for cell in ("O2", "O7", "O12"):
                if book.number(cell, raw_sheet) < 0:
                    raise ValueError(book.source(cell, raw_sheet) + "压损不能为负数。")
            match = re.fullmatch(r"([+\-]?\d+(?:\.\d+)?)\s*Pa", book.text(address), re.I)
            if not match:
                raise ValueError(book.source(address) + "压损结果缺失或单位不是Pa。")
            if decimal(match[1]) < 0:
                raise ValueError(book.source(address) + "压损结果不能为负数。")
            return book.metric("pressure", decimal(match[1]), [address], "Pa")

        group("压力损失", lambda: add("压力损失", 7, [pressure("L8", "降流量")]))
        if any(r["project"] == "准确度等级和最大允许误差" for r in result["records"]):
            seal = book.text("I11")
            if seal:
                add("密封性", 7, [{"key": "seal", "label": "密封性", "value": seal, "source": book.source("I11")}])

        def repeat():
            metrics = []
            for row, flow in zip((75, 78, 81), ("qmax", "qt", "qmin")):
                _, errors = book.mean(row, 3, flow)
                value = book.verify("K%d" % row, max(errors) - min(errors))
                metrics.append(book.metric("repeat:" + flow, value, ["K%d" % row]))
            add("重复性", 66, metrics)
        group("重复性", repeat)

        def mode():
            volume = book.number("D2", "用户模式", positive=True)
            standard = book.number("E2", "用户模式", positive=True) * book.number("I2", "用户模式", positive=True)
            user = book.number("D87")
            if abs(user - (volume - standard) / standard * 100) > Decimal("0.00000001"):
                raise ValueError("用户模式误差与原始体积/流量/时间计算不一致，请重算保存。")
            detector = book.number("H87")
            if "qmax" not in initial or abs(detector - initial["qmax"]) > Decimal("0.00000001"):
                raise ValueError("检测模式没有可核对的本样机初始qmax误差。")
            add("不同工作模式的误差偏差", 66, [
                book.metric("mode", book.verify("K87", user - detector), ["D87", "H87", "K87"])])
        group("不同工作模式的误差偏差", mode)

        def temperature():
            for address in ("B2", "B12"):
                if book.text(address, "温度适应性") != sample:
                    raise ValueError("高、低温原始工作表未填写本台样机编号，参比温度模板残留不作为实测。")
            metrics = []
            for index, row in enumerate(range(120, 143, 2)):
                flow = ("qmax", "0.7qmax", "0.2qmax")[index % 3]
                value, _ = book.mean(row, 2, flow)
                phase = ("reference-before", "high", "low", "reference-after")[index // 3]
                condition_row = (120, 126, 132, 138)[index // 3]
                condition = re.search(r"([+\-]?\d+(?:\.\d+)?)\s*℃", book.text("L%d" % condition_row))
                if not condition:
                    raise ValueError(book.source("L%d" % condition_row) + "缺少明确的温度试验条件。")
                metrics.append(book.metric(phase + ":" + flow + ":" + str(decimal(condition[1])),
                                           value, ["J%d" % row]))
            add("温度适应性", 111, metrics)
        group("温度适应性", temperature)

        def environment():
            metrics = []
            for row, flow in zip((165, 167, 169), ("qmax", "0.2qmax", "qmin")):
                value, _ = book.mean(row, 2, flow)
                metrics.append(book.metric("mean:" + flow, value, ["J%d" % row]))
            add("环境适应性复测", 156, metrics)
        group("环境适应性复测", environment)

        def disturbance():
            book.anchor("M210", "流体扰动：参比条件")
            book.anchor("M216", "流体扰动：90°单弯头")
            metrics, before = [], {}
            for index, row in enumerate(range(210, 221, 2)):
                flow = ("qmax", "0.4qmax", "0.25qmax")[index % 3]
                value, _ = book.mean(row, 2, flow)
                if index < 3:
                    before[flow] = value
                metrics.append(book.metric(("reference:" if index < 3 else "elbow:") + flow, value, ["J%d" % row]))
                if index >= 3:
                    shift = book.verify("L%d" % row, abs(value - before[flow]))
                    metrics.append(book.metric("shift:" + flow, shift, ["L%d" % row]))
            add("流体扰动", 201, metrics)
        group("流体扰动", disturbance)

        def durability():
            metrics, after = [], {}
            for index, (row, flow) in enumerate(zip(DURABLE, FLOWS)):
                value, _ = book.mean(row, 3 if index < 5 else 2, flow)
                after[flow] = value
                metrics.append(book.metric("mean:" + flow, value, ["J%d" % row]))
            if set(initial) != set(FLOWS):
                raise ValueError("初始流量点不完整，不能计算耐久性最大偏离。")
            for key, flows, rows in (("high", FLOWS[:5], DURABLE[:5]), ("low", FLOWS[5:], DURABLE[5:])):
                value = max(abs(after[f] - initial[f]) for f in flows)
                metrics.append(book.metric("shift:" + key, value,
                                           ["J%d" % r for r in rows] + ["J%d" % INITIAL[FLOWS.index(f)] for f in flows]))
                metrics[-1]["source"] += "；逐点与初始平均误差作差后取最大绝对偏离"
            metrics.append(pressure("L247", "耐久性复测"))
            if book.text("M248"):
                metrics.append({"key": "seal", "label": "密封性", "value": book.text("M248"), "source": book.source("M248")})
            add("耐久性", 246, metrics)
        group("耐久性", durability)

        def mechanical():
            metrics = []
            for before, after, flow in zip(range(299, 310, 2), range(311, 322, 2), FLOWS[:5] + ("qmin",)):
                old, _ = book.mean(before, 2, flow)
                new, _ = book.mean(after, 2, flow)
                value = book.verify("L%d" % after, abs(new - old))
                metrics.append(book.metric("shift:" + flow, value, ["J%d" % before, "J%d" % after, "L%d" % after]))
            add("机械环境复测", 290, metrics)
        group("机械环境复测", mechanical)
        result["warnings"].extend(dict.fromkeys(book.date_issues))
        result["warnings"].extend([
            "Excel百分数结果四舍五入保留两位小数，压力损失取整数Pa，负零归零；原始精度及单元格来源保留。",
            "Excel用于燃气表数据识别，符合条件的打印记录合并到配套Word；生成结果不再单独输出Excel，上传原件保留。",
            "测量日期从打印页对应记录区块读取；空白或格式无效时不推测。人员及未提供的定性项目仍须补录。",
            "压力损失沿用Excel已保存结果及公式；工作簿中的0.75系数不作为所有燃气表的通用规则。",
        ])
        if not result["records"]:
            raise ValueError("未找到可采用的燃气表实测数据：" + "；".join(result["warnings"]))
        return result
    finally:
        book.values.close()
        book.formulas.close()


def _project_parent(project):
    name = re.sub(r"-\d+\s*models?$", "", documents._project_name(project["name"]), flags=re.I)
    parent, separator, children = name.partition("—")
    if separator:
        requirements = " ".join(row.get("requirement", "") for row in project.get("rows", []))
        if not children or any(not child or compact(child) not in compact(requirements)
                               for child in children.split("、")):
            raise ValueError("燃气表素材名称中的子项目与技术要求不一致：" + name)
    return parent.strip()


def project_kind(project):
    name = _project_parent(project)
    requirements = " ".join(row.get("requirement", "") for row in project.get("rows", []))
    if name == "功能性要求":
        children = [s for s in ("密封性", "气流方向", "零流量读数") if s in requirements]
        return children[0] if len(children) == 1 else name
    if name == "机械环境" and "复测" in requirements:
        return "机械环境复测"
    return name


def fill_metadata(doc, project):
    values = {"环境温度": project.get("temperature", ""), "相对湿度": project.get("humidity", ""),
              "评价人员": "、".join(project.get("evaluators", [])), "复核人员": "、".join(project.get("reviewers", []))}
    for paragraph in documents._xp(doc.element.body, ".//w:p"):
        text = documents._text(paragraph)
        if sum(label in text for label in values) < 2:
            continue
        for label, value in values.items():
            if not str(value).strip():
                continue
            match = re.search(re.escape(label) + r"[：:](.*?)(?=环境温度[：:]|相对湿度[：:]|大气压力[：:]|评价人员[：:]|复核人员[：:]|$)", text)
            if not match:
                continue
            original = match[1]
            if label in ("环境温度", "相对湿度"):
                replacement = documents._format_result_environment(original, value, label, False)
            else:
                replacement = original[:len(original) - len(original.lstrip())] + value + original[len(original.rstrip()):]
            text = text[:match.start(1)] + replacement + text[match.end(1):]
        documents._replace_result_field_text(paragraph, text)


def order_projects(items, projects):
    positioned = []
    for project in projects:
        name = project_kind(project)
        segments = [["".join(p.split()) for p in row["label"].split(" / ")] for row in items]
        if name == "机械环境复测":
            matches = [i for i, parts in enumerate(segments) if parts[-2:] == ["机械环境", "复测"]]
        else:
            matches = [i for i, parts in enumerate(segments) if compact(name) in [compact(p) for p in parts]]
        parent = _project_parent(project)
        scoped = [i for i in matches if parent in segments[i]]
        if scoped:
            depth = min(segments[i].index(parent) for i in scoped)
            matches = [i for i in scoped if segments[i].index(parent) == depth]
        elif name != parent:
            matches = []
        if len(matches) > 1:
            requirements = " ".join(row["requirement"] for row in project.get("rows", []))
            legacy_durability = name == "耐久性" and len(matches) == 2 and {
                segments[i][-1] for i in matches} == {"断续流量试验", "连续流量试验"}
            if not legacy_durability:
                matches = [i for i in matches if segments[i][-1] in requirements]
            if not matches or matches != list(range(min(matches), max(matches) + 1)):
                raise ValueError("燃气表一览表项目匹配歧义，子项目不连续或无法确定：" + name)
        if not matches:
            raise ValueError("项目未在本次燃气表一览表中找到：" + name)
        positioned.append((min(matches), project))
    return [deepcopy(p) for _, p in sorted(positioned, key=lambda item: item[0])]


def _key(name, row, rows):
    requirement = compact(row["requirement"]).lower()
    flow = next((compact(p).lower() for p in reversed(row["requirement"].split("|"))
                 if compact(p).lower() in FLOWS + ("0.25qmax",)), "")
    low = "qmin≤q<qt" in requirement or "qmin<=q<qt" in requirement
    high = "qt≤q≤qmax" in requirement or "qt<=q<=qmax" in requirement
    band = "low" if low else "high" if high else ""
    if name == "准确度等级和最大允许误差":
        return "same-sign" if "同号" in requirement else "mean:" + flow
    if name == "误差曲线":
        return "span:" + band if band else ""
    if name == "加权平均误差":
        return "weighted"
    if name in ("复现性", "重复性"):
        return ("sd:" if name == "复现性" else "repeat:") + flow
    if name == "不同工作模式的误差偏差":
        return "mode"
    if name in ("密封性", "压力损失"):
        return "seal" if name == "密封性" else "pressure"
    if name == "环境适应性复测":
        return "mean:" + flow
    if name == "机械环境复测":
        return "shift:" + flow
    if name == "温度适应性":
        if "最高" in requirement:
            phase = "high"
        elif "最低" in requirement:
            phase = "low"
        elif "参比" in requirement:
            refs = [r for r in rows if r.get("sample_label") == row.get("sample_label")
                    and compact(r["requirement"]) == compact(row["requirement"])]
            if len(refs) != 2:
                return ""
            phase = "reference-before" if refs[0]["id"] == row["id"] else "reference-after"
        else:
            return ""
        condition = re.search(r"([+\-]?\d+(?:\.\d+)?)℃", requirement)
        return phase + ":" + flow + ":" + str(decimal(condition[1])) if condition else ""
    if name == "流体扰动":
        phase = "shift" if "变化" in requirement else "elbow" if "90°单弯头" in requirement else "reference" if "参比" in requirement else ""
        return phase + ":" + flow if phase else ""
    if name == "耐久性":
        if "密封性" in requirement:
            return "seal"
        if "压力损失" in requirement:
            return "pressure"
        return ("shift:" + band if band else "") if "最大偏离" in requirement else "mean:" + flow
    return ""


def apply_records(project, parsed, include_conditions=False):
    if not any(data.get("parser") == PARSER for data in parsed) and not any(
            "Excel/" in row.get("source", "") for row in project["rows"]) and not any(
                note.startswith(PREFIX) for note in project.get("notes", [])):
        return refresh_curve_results(project, parsed, include_conditions=include_conditions)
    result = deepcopy(project)
    name = project_kind(project)
    records = [record for data in parsed if data.get("parser") == PARSER
               for record in data["records"] if record["project"] == name]
    values, issues = documents._record_conditions(records) if records else ({}, [])
    result["record_summary"] = {"values": values, "issues": issues,
                                "records": [{key: r.get(key, "") for key in
                                             ("source", "sample_id", "model", "start", "end", "temperature", "humidity")}
                                            for r in records]}
    if include_conditions:
        for key in ("start", "end", "temperature", "humidity"):
            previous = project.get("record_summary", {}).get("values", {}).get(key)
            if key in values or previous == result.get(key):
                result[key] = values.get(key, "")
    result["source"] = "；".join(r["source"] for r in records)
    result["notes"] = [note for note in result.get("notes", []) if not note.startswith(PREFIX)]
    result["gas_curves"] = []
    if name in ("误差曲线", "耐久性"):
        record_name = "准确度等级和最大允许误差" if name == "误差曲线" else "耐久性"
        phase = "初始误差" if name == "误差曲线" else "耐久性复测"
        initial_records = [record for data in parsed if data.get("parser") == PARSER
                           for record in data["records"] if record["project"] == record_name]
        for sample in dict.fromkeys(row.get("sample_id", "") for row in result["rows"]):
            candidates = [r for r in initial_records if r["sample_id"] == sample]
            if sample and len(candidates) == 1:
                record = candidates[0]
                metrics = [[m for m in record["rows"] if m["key"] == "mean:" + flow]
                           for flow in reversed(FLOWS)]
                if all(len(matches) == 1 and matches[0].get("raw_value") is not None for matches in metrics):
                    points = [{"flow": flow, "error": matches[0]["raw_value"]}
                              for flow, matches in zip(reversed(FLOWS), metrics)]
                    result["gas_curves"].append({"sample_id": sample, "source": record["source"],
                                                 "model": record.get("model", ""), "points": points})
                    continue
            result["notes"].append(PREFIX + (sample or "未对应样机") + "缺少唯一完整的" + phase + "记录，不生成曲线图。")
    for row in result["rows"]:
        if _manual_result(row):
            if row.get("matched_sample_id") and row["matched_sample_id"] != row.get("sample_id"):
                row["match_issue"] = "样机对应已改变，保留的人工结果属于旧编号，请核对后重新保存"
            continue
        key = _key(name, row, result["rows"])
        sample = row.get("sample_id", "")
        candidates = [r for r in records if r["sample_id"] == sample]
        metrics = [m for r in candidates for m in r["rows"] if m["key"] == key]
        issue = ""
        if not sample:
            issue = "缺少本次样机编号对应关系，不能按模板旧编号猜配"
        elif not key:
            issue = "未支持或无法唯一确定该子试验条件，请手工核对"
        elif len(candidates) > 1:
            issue = "同一样机、项目存在多份记录/批次，不能自动选择"
        elif len(metrics) != 1:
            issue = "缺少本样机对应项目/流量点的有效Excel数据，空表和错误缓存不作为结果"
        if issue:
            row.update(value="", auto_value="", source="", matched_sample_id="", match_issue=issue)
            result["notes"].append(PREFIX + name + "/" + (sample or row["sample_label"]) + "：" + issue)
        else:
            metric = metrics[0]
            row.update(value=metric["value"], auto_value=metric["value"], source=metric["source"],
                       matched_sample_id=sample, match_issue="")
    result["notes"] = list(dict.fromkeys(result["notes"]))
    return result
